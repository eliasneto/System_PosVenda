"""Tela "Projeto > BDO" (pedido do usuário, 2026-10-14): faturamento
mensal das manutenções de RI e RE, a partir da planilha FR_XXXX do
Boletim de desempenho. Cada Boletim é independente: o mesmo município
aparece uma vez em cada Boletim, nunca somado entre Boletins. Grid por
INEP (os valores são do município naquele Boletim, repetidos em cada INEP
dele) ou agrupado por Município/UF dentro de cada Boletim. Fora do
Visualizador: as rotas não estão na lista liberada do
`VisualizadorAccessMiddleware`. A importação fica em "Administrador >
Relatório BDO" (`relatorio_bdo_view`), só do Administrador, mesmo critério
das demais importações (RN-004)."""

import io
import zipfile
from collections import defaultdict
from datetime import date
from decimal import Decimal
from urllib.parse import quote, urlencode

import openpyxl

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, DateField, OuterRef, Subquery, Sum
from django.db.models.functions import Upper
from django.http import HttpResponse, HttpResponseForbidden, QueryDict
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils.http import url_has_allowed_host_and_scheme
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.ri.models import KitPadrao, Ri

from .bdo import (
    COLUNAS_VALORES, PlanilhaBdoError, _normalizar, dias_no_mes_referencia, importar_planilha_bdo, importar_pro_rata,
    valor_proporcional,
)
from .leitura_eace_bdo import conferencia_eace_boletim, enfileirar_leitura_eace_bdo, mensagem_motivo_bdo
from .forms import MesReferenciaBdoForm, PlanilhaBdoUploadForm, PlanilhaProRataUploadForm
from .medicao_bdo import MedicaoBdoError, gerar_medicoes
from .models import BoletimBdo, Escola, InepBdo, LeituraEaceBdo, LogStatusDivergenciaBdo, MunicipioBdo
from .views import _parse_data_filtro

DIFERENCA_TODOS = "1"
AGRUPAMENTOS = (("inep", "Por INEP"), ("municipio", "Por Município"), ("uf", "Por UF"))
LINHAS_POR_PAGINA = 50
TOLERANCIA_VALOR = Decimal("0.01")


def _somar(atual, valor):
    """Soma que fica vazia (`None`) se algum valor do grupo é vazio - só
    acontece no Boletim importado antes dos valores de composição."""
    return None if atual is None or valor is None else atual + valor


def _linhas_agrupadas(municipios, ineps, agrupar):
    """Agrupa por Município (nome sem diferença de maiúsculas + UF) ou por
    UF dentro de cada Boletim, somando os valores de cada linha da
    planilha uma vez só."""
    quantidade_ineps = defaultdict(int)
    for municipio_id in ineps.values_list("municipio_bdo_id", flat=True):
        quantidade_ineps[municipio_id] += 1

    grupos = {}
    for municipio in municipios.select_related("boletim"):
        nome = municipio.municipio.upper()
        numero = municipio.boletim.numero
        chave = (numero, municipio.uf, nome) if agrupar == "municipio" else (numero, municipio.uf)
        grupo = grupos.setdefault(chave, {
            "boletim": numero, "tipo": municipio.boletim.tipo, "uf": municipio.uf, "municipio": nome, "municipios": set(),
            "quantidade_ineps": 0, "municipio_ids": [], **{campo: Decimal("0") for campo in COLUNAS_VALORES},
        })
        grupo["municipios"].add(nome)
        grupo["municipio_ids"].append(municipio.pk)
        grupo["quantidade_ineps"] += quantidade_ineps[municipio.pk]
        for campo in COLUNAS_VALORES:
            grupo[campo] = _somar(grupo[campo], getattr(municipio, campo))

    linhas = sorted(grupos.values(), key=lambda grupo: (grupo["uf"], grupo["municipio"]))
    linhas.sort(key=lambda grupo: grupo["boletim"], reverse=True)
    for grupo in linhas:
        grupo["quantidade_municipios"] = len(grupo.pop("municipios"))
        grupo["valores_em_ordem"] = [grupo[campo] for campo in MunicipioBdo.CAMPOS_VALORES]
    return linhas


def _anotar_data_ativacao(linhas):
    """Pedido do usuário (2026-10-14): Data de Ativação (`Ri.data_ativacao`)
    do RI atual de cada INEP da página - a mesma no Boletim de RE, que não
    tem data própria. INEP sem RI no sistema fica vazio. Busca em Python,
    não por subquery: no MySQL a tabela do BDO tem collation diferente da
    `Escola` e o `=` entre as duas colunas dá erro 1267."""
    datas = {}
    ris = Ri.objects.filter(escola__inep__in={linha.inep for linha in linhas}).order_by("escola__inep", "-criado_em")
    for inep, data_ativacao in ris.values_list("escola__inep", "data_ativacao"):
        datas.setdefault(inep, data_ativacao)
    for linha in linhas:
        linha.data_ativacao = datas.get(linha.inep)
        # Pedido do usuário (2026-10-15): confere com o "Ativado em" da
        # Pro rata do cliente; `None` quando falta uma das duas datas.
        linha.ativacao_confere = (
            None if linha.data_ativacao is None or linha.ativado_em_cliente is None
            else linha.data_ativacao == linha.ativado_em_cliente
        )
        linha.dias_proporcionais = dias_no_mes_referencia(linha.data_ativacao, linha.municipio_bdo.boletim.mes_referencia)
    _anotar_valor_assinatura(linhas)
    for linha in linhas:
        linha.valor_proporcional = valor_proporcional(linha.valor_assinatura, linha.dias_proporcionais)
        _anotar_divergencia(linha)


def _anotar_divergencia(linha):
    """Pedido do usuário (2026-10-15): o INEP diverge da Pro rata do cliente
    quando a Data de Ativação é diferente do "Ativado em" ou o valor pago
    pelo cliente (Valor assinatura − "Valor" da Pro rata, que é o desconto;
    decisão do usuário de 2026-10-15) é menor que o Valor proporcional do
    sistema em mais de R$ 0,01 (arredondamento por INEP); pago a mais
    confere (verde, pedido do usuário de 2026-10-15). Só compara quando os dois lados existem.
    Pedido do usuário (2026-10-15): INEP sem Data de Ativação no sistema
    ou sem "Ativado em" do cliente também recebe status."""
    linha.valor_pago_cliente = (
        None if linha.valor_assinatura is None or linha.valor_desconto_cliente is None
        else linha.valor_assinatura - linha.valor_desconto_cliente
    )
    linha.valor_desconto_sistema = (
        None if linha.valor_assinatura is None or linha.valor_proporcional is None
        else linha.valor_assinatura - linha.valor_proporcional
    )
    linha.valor_confere = (
        None if linha.valor_proporcional is None or linha.valor_pago_cliente is None
        else linha.valor_pago_cliente >= linha.valor_proporcional - TOLERANCIA_VALOR
    )
    linha.tem_divergencia = (
        linha.data_ativacao is None or linha.ativado_em_cliente is None
        or linha.ativacao_confere is False or linha.valor_confere is False
    )
    # Pedido do usuário (2026-10-15): parte do INEP na Dif. do Total MEGA -
    # valor do sistema menos o pago pelo cliente. INEP sem valor
    # proporcional fica fora do Total MEGA e entra na Dif. sem valor.
    linha.valor_recebido = _valor_pago_para_diferenca(linha)
    linha.diferenca_inep = (
        None if linha.valor_proporcional is None or linha.valor_recebido is None
        else linha.valor_proporcional - linha.valor_recebido
    )
    # Lado da Dif. (pedido do usuário, 2026-10-15): "mais" quando o cliente
    # paga mais que o sistema (ou INEP fora do Total MEGA que a planilha
    # cobra), "menos" quando paga menos.
    if linha.diferenca_inep is None or linha.diferenca_inep < -TOLERANCIA_VALOR:
        linha.lado_diferenca = "mais"
    elif linha.diferenca_inep > TOLERANCIA_VALOR:
        linha.lado_diferenca = "menos"
    else:
        linha.lado_diferenca = None
    linha.entra_na_diferenca = linha.lado_diferenca is not None
    linha.valor_lado_diferenca = None if linha.diferenca_inep is None else abs(linha.diferenca_inep)
    # Pedido do usuário (2026-10-15): no INEP Contestado o valor recebido
    # (o valor pago) fica e só o faltando (Valor proporcional − recebido)
    # é contestado para vir em outro Boletim - só quando o cliente pagou
    # menos.
    linha.contestado = linha.status_divergencia == InepBdo.STATUS_CONTESTADO
    linha.valor_faltando = (
        linha.diferenca_inep if linha.contestado and linha.lado_diferenca == "menos" else None
    )


def _valor_pago_para_diferenca(linha):
    """Pedido do usuário (2026-10-15): nunca supor o valor cheio quando o
    cliente informou a ativação. Com o "Valor" da Pro rata (o desconto), o
    pago é Valor assinatura − desconto. Sem ele, o valor pago
    pelo cliente sai do "Ativado em" dele pela mesma conta do sistema -
    data igual, diferença zero. Só o INEP fora da Pro rata conta pelo mês
    cheio (o cliente cobrou sem desconto)."""
    if linha.valor_desconto_cliente is not None:
        return linha.valor_pago_cliente
    if linha.ativado_em_cliente is None:
        return linha.valor_assinatura
    dias = dias_no_mes_referencia(linha.ativado_em_cliente, linha.municipio_bdo.boletim.mes_referencia)
    proporcional_cliente = valor_proporcional(linha.valor_assinatura, dias)
    return linha.valor_assinatura if proporcional_cliente is None else proporcional_cliente


def _anotar_total_mega(grupos, ineps):
    """Pedido do usuário (2026-10-14): "Total MEGA" (nome de 2026-10-15,
    antes "Total Howbe") de cada município da página no agrupamento por
    Município - soma do valor proporcional dos
    INEPs dele no Boletim (os mesmos contados na coluna INEPs). INEP sem
    valor proporcional fica fora da soma e é contado em
    `ineps_sem_valor_mega`; sem nenhum, fica vazio. A Dif. é Total MEGA −
    Valor total da planilha do cliente (pedido do usuário, 2026-10-16).

    Pedido do usuário (2026-10-15): do INEP Contestado só o valor recebido
    (Valor assinatura − desconto da Pro rata) fica no Total MEGA; o
    faltando vai para `total_contestado_mega`, a cobrar em outro Boletim, e
    fica fora da Dif.: o Valor total do cliente não o inclui."""
    por_municipio = {}
    for grupo in grupos:
        for municipio_id in grupo["municipio_ids"]:
            por_municipio[municipio_id] = grupo
        grupo["total_mega"] = None
        grupo["ineps_sem_valor_mega"] = 0
        grupo["ineps_contestados_mega"] = 0
        grupo["total_contestado_mega"] = Decimal("0")
    linhas = list(
        ineps.filter(municipio_bdo_id__in=por_municipio).select_related("municipio_bdo__boletim")
    )
    _anotar_data_ativacao(linhas)
    for linha in linhas:
        grupo = por_municipio[linha.municipio_bdo_id]
        if linha.valor_proporcional is None:
            grupo["ineps_sem_valor_mega"] += 1
        elif linha.contestado:
            grupo["ineps_contestados_mega"] += 1
            faltando = linha.valor_faltando or Decimal("0")
            grupo["total_contestado_mega"] += faltando
            grupo["total_mega"] = (grupo["total_mega"] or Decimal("0")) + linha.valor_proporcional - faltando
        else:
            grupo["total_mega"] = (grupo["total_mega"] or Decimal("0")) + linha.valor_proporcional
    for grupo in grupos:
        grupo["valor_conferencia_mega"] = grupo["valor_total"]
        grupo["diferenca_mega"] = (
            None if grupo["total_mega"] is None or grupo["valor_total"] is None
            else grupo["total_mega"] - grupo["valor_total"]
        )
        # Pedido do usuário (2026-10-16): o Valor total é o que o cliente
        # pagou, então o lado da Dif. sai dele, não da soma por INEP.
        diferenca = grupo["diferenca_mega"] or Decimal("0")
        grupo["lado_diferenca_mega"] = "mais" if diferenca < 0 else "menos" if diferenca > 0 else None
        grupo["valor_lado_diferenca_mega"] = abs(diferenca)


def _faturamento_por_municipio(chaves):
    """Pedido do usuário (2026-10-15): para cada `(uf, municipio)` - nome
    sem acento nem maiúsculas -, as escolas do município no sistema e os
    INEPs dele já faturados em qualquer Boletim de cada tipo (RE/RI), cada
    INEP uma vez. Em Python pelo mesmo motivo de collation de
    `_anotar_data_ativacao`. Devolve `(escolas, faturados)`: `escolas[(uf,
    municipio)]` e `faturados[(tipo, uf, municipio)]` são conjuntos de INEPs."""
    ufs = {uf for uf, _municipio in chaves}
    escolas = defaultdict(set)
    for inep, municipio, uf in Escola.objects.filter(estado__in=ufs).values_list("inep", "municipio", "estado"):
        escolas[(uf, _normalizar(municipio))].add(inep)
    faturados = defaultdict(set)
    for inep, municipio, uf, tipo in InepBdo.objects.filter(
        municipio_bdo__uf__in=ufs, municipio_bdo__boletim__tipo__in=dict(BoletimBdo.TIPO_CHOICES),
    ).values_list("inep", "municipio_bdo__municipio", "municipio_bdo__uf", "municipio_bdo__boletim__tipo"):
        faturados[(tipo, uf, _normalizar(municipio))].add(inep)
    return escolas, faturados


def _situacao_faturamento(faltam, tipo):
    """Pedido do usuário (2026-10-15): enquanto faltar uma escola do
    município no faturamento do tipo, o município é Proporcional; sem
    nenhuma faltando, Integral. Devolve `(situacao, titulo)`."""
    if faltam:
        return "Proporcional", f"Falta(m) {faltam} escola(s) do município no faturamento {tipo}"
    return "Integral", f"Todas as escolas do município no faturamento {tipo}"


def _anotar_ineps_faturados(grupos):
    """Pedido do usuário (2026-10-15): no agrupamento por Município, um
    quadro para RE e outro para RI com os INEPs do município já faturados
    e os que faltam entre todas as escolas do município no sistema, e a
    situação (Proporcional/Integral) no tipo do Boletim da linha."""
    escolas, faturados = _faturamento_por_municipio({(grupo["uf"], _normalizar(grupo["municipio"])) for grupo in grupos})
    for grupo in grupos:
        chave = (grupo["uf"], _normalizar(grupo["municipio"]))
        grupo["quantidade_escolas"] = len(escolas[chave])
        grupo["faturamento_por_tipo"] = [
            {
                "tipo": tipo,
                "faturados": len(faturados[(tipo, *chave)]),
                "faltam": len(escolas[chave] - faturados[(tipo, *chave)]),
            }
            for tipo, _rotulo in BoletimBdo.TIPO_CHOICES
        ]
        grupo["situacao_faturamento"], grupo["titulo_situacao"] = next(
            (_situacao_faturamento(item["faltam"], item["tipo"])
             for item in grupo["faturamento_por_tipo"] if item["tipo"] == grupo["tipo"]),
            (None, ""),
        )


def _anotar_situacao_faturamento(linhas):
    """Pedido do usuário (2026-10-15): no grid por INEP, a situação
    (Proporcional/Integral) do município do INEP no tipo do Boletim e
    quantas escolas faltam. Boletim sem tipo fica vazio."""
    chaves = {(linha.municipio_bdo.uf, _normalizar(linha.municipio_bdo.municipio)) for linha in linhas}
    escolas, faturados = _faturamento_por_municipio(chaves)
    for linha in linhas:
        chave = (linha.municipio_bdo.uf, _normalizar(linha.municipio_bdo.municipio))
        tipo = linha.municipio_bdo.boletim.tipo
        linha.ineps_faltando_municipio = len(escolas[chave] - faturados[(tipo, *chave)]) if tipo else None
        linha.situacao_faturamento, linha.titulo_situacao = (
            (None, "") if linha.ineps_faltando_municipio is None
            else _situacao_faturamento(linha.ineps_faltando_municipio, tipo)
        )


def _anotar_valor_assinatura(linhas):
    """Pedido do usuário (2026-10-14), só conferência no grid: valor mensal
    do INEP conforme o tipo do Boletim - no RE a assinatura da escola
    (`Escola.valor_assinatura_re`), no RI a "Manutenção de Rede Interna" do
    LPU (`KitPadrao`) no Lote da escola. Vazio sem escola, sem valor, sem
    Lote ou em Boletim sem tipo. Busca em Python pelo mesmo motivo de
    collation de `_anotar_data_ativacao`."""
    escolas = {
        inep: (lote, assinatura_re) for inep, lote, assinatura_re in Escola.objects.filter(
            inep__in={linha.inep for linha in linhas},
        ).values_list("inep", "lote", "valor_assinatura_re")
    }
    manutencao_ri = dict(
        KitPadrao.objects.filter(descricao__istartswith="Manutenção de Rede Interna")
        .values_list("lote", "valor_servico")
    )
    for linha in linhas:
        lote, assinatura_re = escolas.get(linha.inep, (None, None))
        tipo = linha.municipio_bdo.boletim.tipo
        if tipo == BoletimBdo.TIPO_RE:
            linha.valor_assinatura = assinatura_re
        elif tipo == BoletimBdo.TIPO_RI:
            linha.valor_assinatura = manutencao_ri.get(lote) if lote else None
        else:
            linha.valor_assinatura = None


def _valor_total_rateado(municipios, ineps):
    """Card "Valor total" com filtro de INEP ou de Data de Ativação (pedido
    do usuário, 2026-10-15): o Valor total da planilha é do município, então
    cada município entra só com a parte dos INEPs filtrados - Valor total ÷
    INEPs do município no Boletim × INEPs filtrados dele."""
    filtrados = dict(
        ineps.values("municipio_bdo_id").annotate(quantidade=Count("pk")).values_list("municipio_bdo_id", "quantidade")
    )
    total = Decimal("0")
    for municipio_id, valor, quantidade in municipios.annotate(
        quantidade=Count("ineps"),
    ).values_list("pk", "valor_total", "quantidade"):
        if quantidade:
            total += valor * filtrados.get(municipio_id, 0) / quantidade
    return total


def _ineps_ativados_no_periodo(data_inicial, data_final):
    """INEPs cuja Data de Ativação do RI atual está no período (pedido do
    usuário, 2026-10-14). Devolve a lista de INEPs, não uma subquery: no
    MySQL a tabela do BDO tem collation diferente da `Escola` e o `=`
    entre as duas colunas dá erro 1267. RI sem Data de Ativação fica de
    fora, mesmo critério do filtro do MIP (RN-075)."""
    ri_atual = Ri.objects.filter(escola=OuterRef("pk")).order_by("-criado_em")
    escolas = Escola.objects.annotate(
        data_ativacao_ri_atual=Subquery(ri_atual.values("data_ativacao")[:1], output_field=DateField()),
    )
    if data_inicial:
        escolas = escolas.filter(data_ativacao_ri_atual__gte=data_inicial)
    if data_final:
        escolas = escolas.filter(data_ativacao_ri_atual__lte=data_final)
    return list(escolas.values_list("inep", flat=True))


FILTROS_STATUS = (
    ("COM_STATUS", "Com status (todos)"),
    ("PENDENTE", "Pendente"),
    (InepBdo.STATUS_APROVADO, "Aprovado"),
    (InepBdo.STATUS_CONTESTADO, "Contestado"),
)


def _ineps_com_status(ineps, status_filtro):
    """Pedido do usuário (2026-10-15): filtro dos INEPs que têm o campo de
    Status (`tem_divergencia`), opcionalmente só os Pendentes, Aprovados ou
    Contestados. A divergência é calculada em Python, então anota todos os
    INEPs já filtrados e devolve os ids."""
    linhas = list(ineps.select_related("municipio_bdo__boletim"))
    _anotar_data_ativacao(linhas)
    return [
        linha.pk for linha in linhas
        if linha.tem_divergencia and (
            status_filtro == "COM_STATUS"
            or (status_filtro == "PENDENTE" and not linha.status_divergencia)
            or linha.status_divergencia == status_filtro
        )
    ]


def _ineps_da_diferenca(ineps, lado):
    """Pedido do usuário (2026-10-15): INEPs que formam a Dif. do Total MEGA
    (link na Dif. do agrupamento por Município) e a soma da parte de cada um.
    Só explica a Dif. por INEP quando o Valor serviços da planilha é a soma
    das assinaturas dos INEPs."""
    linhas = list(ineps.select_related("municipio_bdo__boletim"))
    _anotar_data_ativacao(linhas)
    sem_pro_rata = any(not linha.municipio_bdo.boletim.pro_rata_arquivo for linha in linhas)
    linhas = [
        linha for linha in linhas
        if linha.entra_na_diferenca and lado in (DIFERENCA_TODOS, linha.lado_diferenca)
    ]
    return {
        "ids": [linha.pk for linha in linhas],
        "soma": sum((linha.diferenca_inep for linha in linhas if linha.diferenca_inep is not None), Decimal("0")),
        "soma_lado": sum((linha.valor_lado_diferenca for linha in linhas if linha.diferenca_inep is not None), Decimal("0")),
        "sem_valor": sum(1 for linha in linhas if linha.diferenca_inep is None),
        "sem_pro_rata": sem_pro_rata,
    }


def _anotar_portal_eace(grupos):
    """Pedido do usuário (2026-10-16): coluna "Portal EACE" do grid Por
    Município - Status, "Valor total a ser emitido" e arquivo da NF do card
    do município no portal, gravados pela Leitura EACE do BDO. `eace_lido`
    falso = Boletim ainda não lido do portal; lido sem status = o portal
    não tem card do município."""
    ids = [pk for grupo in grupos for pk in grupo["municipio_ids"]]
    dados = {
        municipio["pk"]: municipio for municipio in MunicipioBdo.objects.filter(pk__in=ids).values(
            "pk", "status_eace", "valor_eace", "arquivo_nf_eace", "eace_lido_em",
        )
    }
    for grupo in grupos:
        itens = [dados[pk] for pk in grupo["municipio_ids"] if pk in dados]
        com_card = [item for item in itens if item["status_eace"]]
        grupo["eace_lido"] = any(item["eace_lido_em"] for item in itens)
        grupo["status_eace"] = ", ".join(sorted({item["status_eace"] for item in com_card}))
        valores = [item["valor_eace"] for item in com_card if item["valor_eace"] is not None]
        grupo["valor_eace"] = sum(valores, Decimal("0")) if valores else None
        grupo["arquivo_nf_eace"] = ", ".join(item["arquivo_nf_eace"] for item in com_card if item["arquivo_nf_eace"])
        grupo["diferenca_eace"] = (
            None if grupo["valor_eace"] is None or grupo["valor_total"] is None
            else grupo["valor_eace"] - grupo["valor_total"]
        )


def _anotar_linhas(linhas, agrupar, ineps):
    """Colunas calculadas do grid, iguais na página e no Excel."""
    if agrupar == "inep":
        _anotar_data_ativacao(linhas)
        _anotar_situacao_faturamento(linhas)
    elif agrupar == "municipio":
        _anotar_total_mega(linhas, ineps)
        _anotar_ineps_faturados(linhas)
        _anotar_portal_eace(linhas)
    return linhas


def _planilha_grid_bdo(linhas, agrupar):
    """Pedido do usuário (2026-10-16): botão "Excel" do grid do BDO - todas
    as linhas dos filtros e do agrupamento da tela, sem paginação, com as
    mesmas colunas do grid."""
    valores = [nome for _coluna, nome in COLUNAS_VALORES.values()]
    status = dict(InepBdo.STATUS_DIVERGENCIA_CHOICES)
    if agrupar == "inep":
        titulos = [
            "INEP", "Município", "UF", "Data de ativação", "Ativado em (cliente)", "Dias proporcionais",
            "Valor assinatura", "Valor proporcional", "Desconto (cliente)", "Valor pago (cliente)", "Status",
            "Valor faltando", *valores, "Boletim de desempenho", "Tipo", "Faturamento",
        ]
        dados = (
            [
                linha.inep, linha.municipio_bdo.municipio, linha.municipio_bdo.uf, linha.data_ativacao,
                linha.ativado_em_cliente, linha.dias_proporcionais, linha.valor_assinatura, linha.valor_proporcional,
                linha.valor_desconto_cliente, linha.valor_pago_cliente,
                status.get(linha.status_divergencia, "Pendente") if linha.tem_divergencia else None,
                linha.valor_faltando, *linha.municipio_bdo.valores_em_ordem, linha.municipio_bdo.boletim.numero,
                linha.municipio_bdo.boletim.tipo or None, linha.situacao_faturamento,
            ]
            for linha in linhas
        )
    elif agrupar == "municipio":
        tipos = [tipo for tipo, _rotulo in BoletimBdo.TIPO_CHOICES]
        titulos = [
            "Município", "UF", "INEPs",
            *(f"{tipo} {coluna}" for tipo in tipos for coluna in ("faturados", "faltam")),
            "Total MEGA", "Dif. (Total MEGA − Valor total)", "Status EACE", "Valor EACE (a ser emitido)",
            "Arquivo da NF (EACE)", *valores, "Boletim de desempenho", "Tipo", "Faturamento",
        ]
        dados = (
            [
                linha["municipio"], linha["uf"], linha["quantidade_ineps"],
                *(item[coluna] for item in linha["faturamento_por_tipo"] for coluna in ("faturados", "faltam")),
                linha["total_mega"], linha["diferenca_mega"],
                linha["status_eace"] or ("Sem card no portal" if linha["eace_lido"] else None),
                linha["valor_eace"], linha["arquivo_nf_eace"] or None, *linha["valores_em_ordem"], linha["boletim"],
                linha["tipo"] or None, linha["situacao_faturamento"],
            ]
            for linha in linhas
        )
    else:
        titulos = ["UF", "Municípios", "INEPs", *valores, "Boletim de desempenho", "Tipo"]
        dados = (
            [
                linha["uf"], linha["quantidade_municipios"], linha["quantidade_ineps"], *linha["valores_em_ordem"],
                linha["boletim"], linha["tipo"] or None,
            ]
            for linha in linhas
        )

    workbook = openpyxl.Workbook()
    aba = workbook.active
    aba.title = "BDO"
    aba.append(titulos)
    for celula in aba[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F2937")
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    colunas_moeda = {
        indice for indice, titulo in enumerate(titulos, start=1)
        if titulo.startswith(("Valor", "Desconto", "Total", "Dif."))
    }
    for linha in dados:
        aba.append(linha)
        for indice, celula in enumerate(aba[aba.max_row], start=1):
            if indice in colunas_moeda:
                celula.number_format = '"R$" #,##0.00'
            elif isinstance(celula.value, date):
                celula.number_format = "DD/MM/YYYY"
    for indice, titulo in enumerate(titulos, start=1):
        aba.column_dimensions[get_column_letter(indice)].width = max(12, min(len(titulo) + 4, 30))
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions
    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()


@login_required
def bdo_view(request):
    boletins = list(BoletimBdo.objects.all())
    numero_filtro = (request.GET.get("boletim") or "").strip()
    boletim = next((b for b in boletins if b.numero == numero_filtro), None)

    agrupar = request.GET.get("agrupar") or "inep"
    if agrupar not in dict(AGRUPAMENTOS):
        agrupar = "inep"
    inep_filtro = (request.GET.get("inep") or "").strip()
    uf_filtro = (request.GET.get("uf") or "").strip().upper()
    municipio_filtro = (request.GET.get("municipio") or "").strip()
    data_inicial = _parse_data_filtro(request.GET.get("data_inicial"))
    data_final = _parse_data_filtro(request.GET.get("data_final"))
    status_filtro = (request.GET.get("status") or "").strip().upper()
    if status_filtro not in dict(FILTROS_STATUS):
        status_filtro = ""
    diferenca_filtro = request.GET.get("diferenca") or ""
    if diferenca_filtro not in (DIFERENCA_TODOS, "mais", "menos"):
        diferenca_filtro = ""
    if data_inicial and data_final and data_inicial > data_final:
        messages.error(request, "A data inicial não pode ser depois da data final — filtro de data ignorado.")
        data_inicial = data_final = None
    contexto = {
        "boletins": boletins,
        "boletim": boletim,
        "agrupar": agrupar,
        "agrupamentos": AGRUPAMENTOS,
        "inep_filtro": inep_filtro,
        "uf_filtro": uf_filtro,
        "municipio_filtro": municipio_filtro,
        "data_inicial": data_inicial,
        "data_final": data_final,
        "status_filtro": status_filtro,
        "diferenca_filtro": diferenca_filtro,
    }
    if not boletins:
        return render(request, "escolas/bdo.html", contexto)
    contexto.update(_contexto_recebiveis(QueryDict(urlencode({
        "boletim": numero_filtro if boletim else "",
        "data_inicial": data_inicial.strftime("%d/%m/%Y") if data_inicial else "",
        "data_final": data_final.strftime("%d/%m/%Y") if data_final else "",
    }))))

    municipios = MunicipioBdo.objects.all()
    if boletim:
        municipios = municipios.filter(boletim=boletim)
    ufs = list(municipios.order_by("uf").values_list("uf", flat=True).distinct())
    if uf_filtro:
        municipios = municipios.filter(uf=uf_filtro)
    # Pedido do usuário (2026-10-14): lista só os municípios do grid no
    # Boletim e na UF escolhidos; nome sem diferença de maiúsculas.
    # Maiúsculas em Python: o UPPER do SQLite não converte acento.
    grafias = defaultdict(set)
    for nome in municipios.values_list("municipio", flat=True).distinct():
        grafias[nome.upper()].add(nome)
    nomes_municipios = sorted(grafias, key=_normalizar)
    municipio_filtro = municipio_filtro.upper()
    if municipio_filtro not in grafias:
        municipio_filtro = ""
    if municipio_filtro:
        municipios = municipios.filter(municipio__in=grafias[municipio_filtro])
    ineps = InepBdo.objects.filter(municipio_bdo__in=municipios)
    if inep_filtro:
        ineps = ineps.filter(inep__icontains=inep_filtro)
    if data_inicial or data_final:
        ineps = ineps.filter(inep__in=_ineps_ativados_no_periodo(data_inicial, data_final))
    if status_filtro:
        ineps = ineps.filter(pk__in=_ineps_com_status(ineps, status_filtro))
    if diferenca_filtro:
        contexto["diferenca"] = _ineps_da_diferenca(ineps, diferenca_filtro)
        ineps = ineps.filter(pk__in=contexto["diferenca"]["ids"])
    filtra_ineps = bool(inep_filtro or data_inicial or data_final or status_filtro or diferenca_filtro)
    if filtra_ineps:
        municipios = municipios.filter(pk__in=ineps.values("municipio_bdo_id"))

    if agrupar == "inep":
        linhas = ineps.select_related("municipio_bdo__boletim", "status_divergencia_por").order_by(
            "-municipio_bdo__boletim__numero", "municipio_bdo__uf", Upper("municipio_bdo__municipio"), "inep",
        )
    else:
        linhas = _linhas_agrupadas(municipios, ineps, agrupar)

    if request.GET.get("exportar") == "excel":
        linhas = _anotar_linhas(list(linhas), agrupar, ineps)
        nome = f"BDO - {boletim.numero if boletim else 'todos os Boletins'} - {dict(AGRUPAMENTOS)[agrupar]}.xlsx"
        return _anexo(_planilha_grid_bdo(linhas, agrupar), nome, MIME_XLSX)
    page_obj = Paginator(linhas, LINHAS_POR_PAGINA).get_page(request.GET.get("page"))
    page_obj.object_list = _anotar_linhas(list(page_obj.object_list), agrupar, ineps)
    filtros = {
        chave: valor for chave, valor in (
            ("boletim", numero_filtro if boletim else ""), ("inep", inep_filtro),
            ("uf", uf_filtro), ("municipio", municipio_filtro),
            ("data_inicial", data_inicial.strftime("%d/%m/%Y") if data_inicial else ""),
            ("data_final", data_final.strftime("%d/%m/%Y") if data_final else ""),
            ("status", status_filtro),
            ("diferenca", diferenca_filtro),
        ) if valor
    }
    if filtra_ineps:
        valor_total = _valor_total_rateado(municipios, ineps)
    else:
        valor_total = municipios.aggregate(total=Sum("valor_total"))["total"] or Decimal("0")
    contexto.update({
        "page_obj": page_obj,
        "ufs": ufs,
        "nomes_municipios": nomes_municipios,
        "municipio_filtro": municipio_filtro,
        # Card "INEPs": cada INEP uma vez, mesmo repetido em mais de um
        # Boletim (pedido do usuário, 2026-10-14) - o grid continua com
        # uma linha por INEP em cada Boletim.
        "total_ineps": ineps.values("inep").distinct().count(),
        "total_municipios": municipios.values("boletim_id", "uf", nome=Upper("municipio")).distinct().count(),
        "total_boletins": municipios.values("boletim_id").distinct().count(),
        "valor_total": valor_total.quantize(Decimal("0.01")),
        "filtros_ativos": bool(boletim or uf_filtro or municipio_filtro or filtra_ineps),
        "filtros_status": FILTROS_STATUS,
        "querystring_filtros": urlencode(filtros),
        "querystring_paginacao": urlencode({**filtros, "agrupar": agrupar}),
        "querystring_sem_diferenca": urlencode(
            {**{chave: valor for chave, valor in filtros.items() if chave != "diferenca"}, "agrupar": agrupar},
        ),
        "status_divergencia_choices": InepBdo.STATUS_DIVERGENCIA_CHOICES,
        "divergencias_na_pagina": agrupar == "inep" and any(linha.tem_divergencia for linha in page_obj),
        **(_contexto_leitura_eace(boletim) if boletim else {}),
    })
    return render(request, "escolas/bdo.html", contexto)


def _contexto_leitura_eace(boletim):
    leitura = boletim.leituras_eace.select_related("solicitado_por").first()
    if leitura and leitura.motivo_erro:
        leitura.motivo_legivel = mensagem_motivo_bdo(leitura.motivo_erro)
    return {
        "leitura_eace": leitura,
        "leitura_eace_em_andamento": bool(leitura and leitura.em_andamento),
        "total_municipios_boletim": boletim.municipios.count(),
        "conferencia_eace": conferencia_eace_boletim(boletim),
    }


def _voltar_para_bdo(request, boletim):
    destino = request.POST.get("next") or ""
    if destino and url_has_allowed_host_and_scheme(destino, allowed_hosts={request.get_host()}):
        return redirect(destino)
    return redirect(f"{reverse('bdo')}?{urlencode({'boletim': boletim.numero, 'agrupar': 'municipio'})}")


@login_required
@require_POST
def bdo_leitura_eace_view(request, pk):
    """Botão "Atualizar do portal EACE" (pedido do usuário, 2026-10-16):
    só coloca a leitura do Boletim na fila do RPA EACE."""
    boletim = get_object_or_404(BoletimBdo, pk=pk)
    if not boletim.numero.isdigit():
        messages.error(request, f"O Boletim {boletim.numero} não é um número de pedido do portal EACE.")
        return _voltar_para_bdo(request, boletim)
    _leitura, criada = enfileirar_leitura_eace_bdo(boletim, request.user)
    if criada:
        messages.success(
            request, f"Leitura do Boletim {boletim.numero} no portal EACE colocada na fila. "
            "A tela se atualiza sozinha quando terminar.",
        )
    else:
        messages.info(request, f"O Boletim {boletim.numero} já está sendo lido do portal EACE - aguarde terminar.")
    return _voltar_para_bdo(request, boletim)


@login_required
def bdo_leitura_eace_status_view(request, pk):
    """Fragmento do status da leitura (polling HTMX a cada 3 s enquanto
    estiver em andamento); quando termina, `HX-Refresh` recarrega a tela
    para o grid mostrar os dados novos."""
    boletim = get_object_or_404(BoletimBdo, pk=pk)
    contexto = _contexto_leitura_eace(boletim)
    if request.GET.get("acompanhando") and not contexto["leitura_eace_em_andamento"]:
        resposta = HttpResponse(status=204)
        resposta["HX-Refresh"] = "true"
        return resposta
    return render(request, "escolas/_bdo_leitura_eace.html", {
        **contexto, "boletim": boletim, "url_retorno": request.headers.get("HX-Current-URL", ""),
    })


@login_required
def relatorio_bdo_view(request):
    """Tela "Administrador > Relatório BDO" (pedido do usuário, 2026-10-14):
    importação da planilha FR_XXXX, mesmo layout das demais importações
    (`relatorio_eace_mip_view`). Lista os Boletins já importados - cada um
    independente, um número repetido é recusado."""
    if not request.user.is_administrador:
        return HttpResponseForbidden("Somente Administrador pode acessar esta tela.")

    if request.method == "POST":
        form = PlanilhaBdoUploadForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                gravados = importar_planilha_bdo(
                    form.cleaned_data["boletins"], form.cleaned_data["tipo"], form.cleaned_data["arquivo"].name,
                    request.user, form.cleaned_data["mes_referencia"],
                )
            except PlanilhaBdoError as erro:
                messages.error(request, f"Não foi possível importar: {erro}")
                return redirect("relatorio_bdo")
            for boletim in gravados:
                messages.success(
                    request,
                    f"Boletim {boletim.numero} ({boletim.tipo}) importado: {boletim.municipios.count()} município(s), "
                    f"{InepBdo.objects.filter(municipio_bdo__boletim=boletim).count()} INEP(s).",
                )
            return redirect("relatorio_bdo")
        mensagens_erro = [erro for erros in form.errors.values() for erro in erros]
        messages.error(request, "Não foi possível importar: " + " ".join(mensagens_erro))
    else:
        form = PlanilhaBdoUploadForm()

    tipo_filtro = (request.GET.get("tipo") or "").strip().upper()
    if tipo_filtro not in dict(BoletimBdo.TIPO_CHOICES):
        tipo_filtro = ""
    boletins = BoletimBdo.objects.select_related("importado_por").annotate(
        quantidade_municipios=Count("municipios", distinct=True),
        quantidade_ineps=Count("municipios__ineps"),
    )
    if tipo_filtro:
        boletins = boletins.filter(tipo=tipo_filtro)
    boletins = list(boletins)
    valores = dict(
        MunicipioBdo.objects.filter(boletim__in=boletins)
        .values("boletim_id").annotate(total=Sum("valor_total")).values_list("boletim_id", "total")
    )
    for boletim in boletins:
        boletim.valor_total = (valores.get(boletim.pk) or Decimal("0")).quantize(Decimal("0.01"))

    return render(request, "escolas/relatorio_bdo.html", {
        "form": form,
        "boletins": boletins,
        "tipo_filtro": tipo_filtro,
        "tipos": BoletimBdo.TIPO_CHOICES,
    })


@login_required
@require_POST
def bdo_mes_referencia_view(request, pk):
    """Grava o mês de referência de um Boletim já importado (pedido do
    usuário, 2026-10-14). Só Administrador, como a importação."""
    if not request.user.is_administrador:
        return HttpResponseForbidden("Somente Administrador pode alterar o Boletim.")
    boletim = get_object_or_404(BoletimBdo, pk=pk)
    form = MesReferenciaBdoForm(request.POST)
    if form.is_valid():
        boletim.mes_referencia = form.cleaned_data["mes_referencia"].replace(day=1)
        boletim.save(update_fields=["mes_referencia"])
        messages.success(request, f"Boletim {boletim.numero}: mês de referência {boletim.mes_referencia:%m/%Y}.")
    else:
        mensagens_erro = [erro for erros in form.errors.values() for erro in erros]
        messages.error(request, f"Boletim {boletim.numero}: " + " ".join(mensagens_erro))
    return redirect("relatorio_bdo")


@login_required
@require_POST
def bdo_pro_rata_view(request, pk):
    """Recebe a planilha Pro rata do cliente de um Boletim (pedido do
    usuário, 2026-10-15): grava o "Ativado em" de cada INEP para conferir
    com a Data de Ativação do sistema no grid por INEP. Só Administrador,
    como a importação; reenviar substitui a anterior."""
    if not request.user.is_administrador:
        return HttpResponseForbidden("Somente Administrador pode alterar o Boletim.")
    boletim = get_object_or_404(BoletimBdo, pk=pk)
    form = PlanilhaProRataUploadForm(request.POST, request.FILES)
    if not form.is_valid():
        mensagens_erro = [erro for erros in form.errors.values() for erro in erros]
        messages.error(request, f"Boletim {boletim.numero}: Pro rata não importada. " + " ".join(mensagens_erro))
        return redirect("relatorio_bdo")
    atualizados, fora = importar_pro_rata(
        boletim, form.cleaned_data["dados"], form.cleaned_data["arquivo"].name, request.user,
    )
    messages.success(request, f"Boletim {boletim.numero}: Pro rata importada, {atualizados} INEP(s) com Ativado em.")
    if fora:
        exemplos = ", ".join(fora[:10]) + (", ..." if len(fora) > 10 else "")
        messages.warning(request, f"Boletim {boletim.numero}: {len(fora)} INEP(s) da Pro rata não estão no Boletim: {exemplos}.")
    return redirect("relatorio_bdo")


STATUS_DIVERGENCIA = dict(InepBdo.STATUS_DIVERGENCIA_CHOICES)


def _alterar_status_divergencia(linhas, status, usuario, em_massa=False):
    """Pedido do usuário (2026-10-15): grava o status da divergência
    (Aprovado ou Contestado) dos INEPs e o log
    de cada troca em `LogStatusDivergenciaBdo`. Só INEP que diverge agora
    recebe status. Devolve
    `(alterados, sem_divergencia)`."""
    _anotar_data_ativacao(linhas)
    agora = timezone.now()
    alterados, sem_divergencia, logs = [], 0, []
    for linha in linhas:
        if not linha.tem_divergencia:
            sem_divergencia += 1
            continue
        if linha.status_divergencia == status:
            continue
        logs.append(LogStatusDivergenciaBdo(
            inep_bdo=linha, boletim=linha.municipio_bdo.boletim.numero, inep=linha.inep,
            acao=LogStatusDivergenciaBdo.ACAO_ALTERADO, status_anterior=linha.status_divergencia, status=status,
            em_massa=em_massa, data_ativacao=linha.data_ativacao, ativado_em_cliente=linha.ativado_em_cliente,
            valor_desconto_sistema=linha.valor_desconto_sistema, valor_desconto_cliente=linha.valor_desconto_cliente,
            usuario=usuario,
        ))
        linha.status_divergencia, linha.status_divergencia_em, linha.status_divergencia_por = status, agora, usuario
        _anotar_divergencia(linha)
        alterados.append(linha)
    with transaction.atomic():
        InepBdo.objects.bulk_update(
            alterados, ["status_divergencia", "status_divergencia_em", "status_divergencia_por"], batch_size=500,
        )
        LogStatusDivergenciaBdo.objects.bulk_create(logs)
    return len(alterados), sem_divergencia


def _status_do_post(request):
    status = (request.POST.get("status") or "").strip().upper()
    return status if status in STATUS_DIVERGENCIA else None


@login_required
@require_POST
def bdo_status_divergencia_view(request, pk):
    """Troca o status da divergência de um INEP pelo campo lista do grid
    (pedido do usuário, 2026-10-15). Administrador e Analista; Visualizador
    nem chega aqui (`VisualizadorAccessMiddleware`). HTMX recebe a célula,
    sem HTMX volta ao BDO."""
    if request.user.is_visualizador:
        return HttpResponseForbidden("Visualizador não pode alterar o status da divergência.")
    linha = get_object_or_404(
        InepBdo.objects.select_related("municipio_bdo__boletim", "status_divergencia_por"), pk=pk,
    )
    status = _status_do_post(request)
    if status is None:
        _anotar_data_ativacao([linha])
        messages.error(request, "Status inválido.")
    elif _alterar_status_divergencia([linha], status, request.user)[1]:
        messages.error(request, f"INEP {linha.inep}: não há divergência para mudar o status.")
    if request.headers.get("HX-Request"):
        return render(request, "escolas/_bdo_status_divergencia.html", {
            "linha": linha, "status_divergencia_choices": InepBdo.STATUS_DIVERGENCIA_CHOICES, "oob": True,
        })
    return redirect(f"{reverse('bdo')}?boletim={linha.municipio_bdo.boletim.numero}&inep={linha.inep}")


@login_required
@require_POST
def bdo_status_divergencia_massa_view(request):
    """Troca em massa o status da divergência dos INEPs marcados no grid
    (pedido do usuário, 2026-10-15), mesmas regras do campo de cada INEP.
    Volta para a mesma página do grid, com os mesmos filtros."""
    if request.user.is_visualizador:
        return HttpResponseForbidden("Visualizador não pode alterar o status da divergência.")
    destino = request.POST.get("next") or ""
    if not url_has_allowed_host_and_scheme(destino, allowed_hosts={request.get_host()}, require_https=request.is_secure()):
        destino = reverse("bdo")
    status = _status_do_post(request)
    ids = [valor for valor in request.POST.getlist("inep_bdo") if valor.isdigit()]
    if status is None:
        messages.error(request, "Escolha um status válido.")
        return redirect(destino)
    if not ids:
        messages.error(request, "Marque pelo menos um INEP.")
        return redirect(destino)
    linhas = list(
        InepBdo.objects.filter(pk__in=ids).select_related("municipio_bdo__boletim", "status_divergencia_por")
    )
    alterados, sem_divergencia = _alterar_status_divergencia(linhas, status, request.user, em_massa=True)
    messages.success(request, f"Status \"{STATUS_DIVERGENCIA[status]}\" aplicado em {alterados} INEP(s).")
    if sem_divergencia:
        messages.error(request, f"{sem_divergencia} INEP(s) sem divergência ficaram de fora.")
    return redirect(destino)


VISOES_RECEBIVEIS = (("linhas", "Linhas"), ("municipio", "Por Município"), ("boletim", "Por Boletim"))


def _ids(valores):
    """Ids de `MunicipioBdo` vindos do formulário, um por campo ou vários
    separados por vírgula (botão "Retirar" de um município agrupado)."""
    ids = set()
    for valor in valores:
        ids.update(int(parte) for parte in valor.split(",") if parte.strip().isdigit())
    return ids


def _contexto_recebiveis(dados):
    """Pedido do usuário (2026-10-14): modal "Recebíveis" da tela BDO, só
    montagem e conferência - nada é gravado. Escolhe um ou mais Boletins,
    traz as linhas da planilha deles e deixa retirar por município, nunca
    por INEP: os valores são do município. Cada linha da planilha é
    independente - município repetido no Boletim já vem rateado pelo
    cliente - e "Retirar" age só nela (ou em todas as do grupo, na visão
    "Por Município"). Os Boletins continuam independentes: o agrupamento
    por município é dentro de cada Boletim. O estado (Boletins, retirados,
    visão, período) vai e volta no próprio formulário do modal.

    Pedido do usuário (2026-10-15): abre com o Boletim e o período de Data
    de Ativação filtrados no grid, na visão "Por Município". Só entram os
    INEPs ativados no período; o município entra com a parte dos INEPs não
    Contestados (Valor total ÷ INEPs do município × INEPs não Contestados),
    mesmo rateio do card "Valor total". Do INEP Contestado entra só o valor
    recebido (Valor assinatura − desconto da Pro rata, pedido do usuário
    de 2026-10-15); o faltando fica para outro Boletim. Município sem nenhum INEP no período
    fica de fora.

    Pedido do usuário (2026-10-16): o valor do município é o "Valor total a
    ser emitido" do card dele no portal EACE (`MunicipioBdo.valor_eace`, da
    Leitura EACE do BDO), no lugar do rateio acima - que só fica na tela
    para o município ainda sem valor EACE (`sem_valor_eace`), e nesse caso
    a Medição não é gerada (`recebiveis_eace_confere`)."""
    boletins = list(BoletimBdo.objects.all())
    numeros = set(dados.getlist("boletim"))
    selecionados = [boletim for boletim in boletins if boletim.numero in numeros]
    visao = dados.get("visao") or "municipio"
    if visao not in dict(VISOES_RECEBIVEIS):
        visao = "municipio"
    data_inicial = _parse_data_filtro(dados.get("data_inicial"))
    data_final = _parse_data_filtro(dados.get("data_final"))
    if data_inicial and data_final and data_inicial > data_final:
        data_inicial = data_final = None

    retirados = _ids(dados.getlist("retirado")) | _ids(dados.getlist("retirar"))
    restaurar = dados.getlist("restaurar")
    retirados = set() if "todos" in restaurar else retirados - _ids(restaurar)

    municipios = list(
        MunicipioBdo.objects.filter(boletim__in=selecionados).select_related("boletim")
        .annotate(quantidade_ineps=Count("ineps")).order_by("-boletim__numero", "uf", Upper("municipio"), "pk")
    )
    ineps = defaultdict(list)
    recebido_contestados = defaultdict(Decimal)
    quantidade_contestados = defaultdict(int)
    entram = InepBdo.objects.filter(municipio_bdo__in=municipios)
    if data_inicial or data_final:
        entram = entram.filter(inep__in=_ineps_ativados_no_periodo(data_inicial, data_final))
    contestados = list(entram.filter(status_divergencia=InepBdo.STATUS_CONTESTADO).select_related("municipio_bdo__boletim"))
    _anotar_data_ativacao(contestados)
    for linha in contestados:
        recebido_contestados[linha.municipio_bdo_id] += linha.valor_recebido or Decimal("0")
        quantidade_contestados[linha.municipio_bdo_id] += 1
    for municipio_id, inep in entram.exclude(status_divergencia=InepBdo.STATUS_CONTESTADO).values_list(
        "municipio_bdo_id", "inep",
    ):
        ineps[municipio_id].append(inep)
    for linha in contestados:
        ineps[linha.municipio_bdo_id].append(linha.inep)
    for municipio in municipios:
        municipio.lista_ineps = ineps[municipio.pk]
        nao_contestados = len(municipio.lista_ineps) - quantidade_contestados[municipio.pk]
        municipio.valor_recebivel = (
            municipio.valor_total * nao_contestados / municipio.quantidade_ineps + recebido_contestados[municipio.pk]
        ).quantize(Decimal("0.01")) if municipio.quantidade_ineps else Decimal("0.00")
        municipio.sem_valor_eace = municipio.valor_eace is None
        if not municipio.sem_valor_eace:
            municipio.valor_recebivel = municipio.valor_eace
        municipio.quantidade_ineps = len(municipio.lista_ineps)
    municipios = [municipio for municipio in municipios if municipio.quantidade_ineps]
    retirados &= {municipio.pk for municipio in municipios}
    mantidos = [municipio for municipio in municipios if municipio.pk not in retirados]

    agrupados = {}
    for municipio in mantidos:
        chave = (municipio.boletim.numero, municipio.uf, municipio.municipio.upper())
        if visao == "boletim":
            chave = chave[:1]
        grupo = agrupados.setdefault(chave, {
            "boletim": municipio.boletim, "uf": municipio.uf, "municipio": municipio.municipio.upper(),
            "ids": [], "nomes": set(), "linhas": 0, "quantidade_ineps": 0, "valor_total": Decimal("0"),
        })
        grupo["ids"].append(str(municipio.pk))
        grupo["nomes"].add((municipio.uf, municipio.municipio.upper()))
        grupo["linhas"] += 1
        grupo["quantidade_ineps"] += municipio.quantidade_ineps
        grupo["valor_total"] += municipio.valor_recebivel
        grupo["sem_valor_eace"] = grupo.get("sem_valor_eace", False) or municipio.sem_valor_eace
    for grupo in agrupados.values():
        grupo["ids"] = ",".join(grupo.pop("ids"))
        grupo["quantidade_municipios"] = len(grupo.pop("nomes"))
    conferencias = [conferencia_eace_boletim(boletim) for boletim in selecionados]

    return {
        "recebiveis_boletins": boletins,
        "recebiveis_selecionados": {boletim.numero for boletim in selecionados},
        "recebiveis_visao": visao,
        "recebiveis_visoes": VISOES_RECEBIVEIS,
        "recebiveis_linhas": mantidos,
        "recebiveis_grupos": list(agrupados.values()),
        "recebiveis_retirados": [municipio for municipio in municipios if municipio.pk in retirados],
        "recebiveis_quantidade_ineps": sum(municipio.quantidade_ineps for municipio in mantidos),
        "recebiveis_quantidade_municipios": len({(m.boletim_id, m.uf, m.municipio.upper()) for m in mantidos}),
        "recebiveis_valor_total": sum((municipio.valor_recebivel for municipio in mantidos), Decimal("0.00")),
        "recebiveis_data_inicial": data_inicial,
        "recebiveis_data_final": data_final,
        "recebiveis_contestados": len(contestados),
        "recebiveis_recebido_contestados": sum(recebido_contestados.values(), Decimal("0.00")),
        "recebiveis_conferencias_eace": conferencias,
        "recebiveis_eace_confere": all(conferencia["confere"] for conferencia in conferencias),
    }


@login_required
def bdo_recebiveis_view(request):
    """Conteúdo do modal "Recebíveis" (HTMX). POST só para o estado do
    formulário (Boletins e retirados) não esbarrar no limite da URL - não
    grava nada. Fora do Visualizador: rota fora da lista liberada do
    `VisualizadorAccessMiddleware`."""
    dados = request.POST if request.method == "POST" else request.GET
    return render(request, "escolas/_bdo_recebiveis.html", _contexto_recebiveis(dados))


MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _anexo(conteudo, nome, content_type):
    """Download com nome acentuado: filename ASCII de reserva + filename*
    (RFC 5987), mesmo padrão das demais planilhas."""
    resposta = HttpResponse(conteudo, content_type=content_type)
    nome_ascii = nome.encode("ascii", "ignore").decode("ascii") or "medicao"
    resposta["Content-Disposition"] = f"attachment; filename=\"{nome_ascii}\"; filename*=UTF-8''{quote(nome)}"
    return resposta


@login_required
@require_POST
def bdo_medicao_view(request):
    """Gera a planilha de Medição de cada Boletim do modal "Recebíveis"
    (pedido do usuário, 2026-10-15) com o que ficou nele depois da
    conferência: uma .xlsx com um Boletim, um .zip com uma por Boletim.
    Grava o número da Medição e o 1º faturamento dos INEPs novos, por isso
    fica fora do Visualizador."""
    if request.user.is_visualizador:
        return HttpResponseForbidden("Visualizador não pode gerar a Medição.")
    contexto = _contexto_recebiveis(request.POST)
    linhas = contexto["recebiveis_linhas"]
    if linhas and not contexto["recebiveis_eace_confere"]:
        problemas = " ".join(c["mensagem"] for c in contexto["recebiveis_conferencias_eace"] if not c["confere"])
        messages.error(
            request, f"Medição não gerada: {problemas} Atualize do portal EACE na tela BDO (Boletim escolhido).",
        )
        return redirect("bdo")
    if not linhas:
        messages.error(request, "Medição não gerada: escolha um Boletim no Recebíveis e mantenha pelo menos um município.")
        return redirect("bdo")
    chaves = {(linha.uf, _normalizar(linha.municipio)) for linha in linhas}
    escolas, faturados = _faturamento_por_municipio(chaves)
    situacoes = {
        (tipo, *chave): _situacao_faturamento(len(escolas[chave] - faturados[(tipo, *chave)]), tipo)[0]
        for chave in chaves for tipo, _rotulo in BoletimBdo.TIPO_CHOICES
    }
    try:
        arquivos = gerar_medicoes(linhas, situacoes, request.user)
    except MedicaoBdoError as erro:
        messages.error(request, f"Medição não gerada: {erro}")
        return redirect("bdo")
    if len(arquivos) == 1:
        nome, conteudo = arquivos[0]
        return _anexo(conteudo, nome, MIME_XLSX)
    saida = io.BytesIO()
    with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as arquivo_zip:
        for nome, conteudo in arquivos:
            arquivo_zip.writestr(nome, conteudo)
    return _anexo(saida.getvalue(), "MEDIÇÕES - MANUTENÇÃO.zip", "application/zip")
