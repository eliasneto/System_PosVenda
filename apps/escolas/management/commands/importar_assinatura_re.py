from decimal import InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError

from apps.escolas.bdo import _normalizar, _texto, _valor_decimal
from apps.escolas.models import Escola

try:
    import openpyxl
except ImportError:
    openpyxl = None

COLUNA_INEP = "INEP"
COLUNA_ASSINATURA = "ASSINATURA MENSAL R.E"


class Command(BaseCommand):
    """Pedido do usuário (2026-10-14): grava `Escola.valor_assinatura_re` a
    partir da aba FATURAMENTO BDO do CONSOLIDADO EACE Atualizado.xlsx
    (colunas INEP e ASSINATURA MENSAL R.E), para conferência no grid do
    BDO. Repetir é seguro: o valor da planilha substitui o gravado. INEP
    sem escola no sistema é só contado, nunca cria escola."""

    help = "Grava a assinatura mensal do RE de cada escola (aba FATURAMENTO BDO do CONSOLIDADO EACE)."

    def add_arguments(self, parser):
        parser.add_argument("arquivo", type=str, help="Caminho para 'CONSOLIDADO EACE Atualizado.xlsx'.")
        parser.add_argument("--aba", default="FATURAMENTO BDO", help="Nome da aba (padrao: FATURAMENTO BDO).")

    def handle(self, *args, **options):
        if openpyxl is None:
            raise CommandError("Dependencia 'openpyxl' nao instalada.")
        caminho = Path(options["arquivo"])
        if not caminho.exists():
            raise CommandError(f"Arquivo nao encontrado: {caminho}")
        try:
            planilha = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        except Exception as erro:
            raise CommandError(f"Nao foi possivel abrir '{caminho}': {erro}")
        if options["aba"] not in planilha.sheetnames:
            raise CommandError(f"Aba '{options['aba']}' nao encontrada. Abas disponiveis: {planilha.sheetnames}")

        try:
            valores, invalidas = self._ler(planilha[options["aba"]].iter_rows(values_only=True))
        finally:
            planilha.close()

        escolas = list(Escola.objects.filter(inep__in=valores))
        for escola in escolas:
            escola.valor_assinatura_re = valores[escola.inep]
        Escola.objects.bulk_update(escolas, ["valor_assinatura_re"], batch_size=500)

        self.stdout.write(self.style.SUCCESS(
            f"Assinatura RE: {len(escolas)} escola(s) atualizada(s), "
            f"{len(valores) - len(escolas)} INEP(s) sem escola no sistema, "
            f"{invalidas} linha(s) com valor invalido ignorada(s)."
        ))

    @staticmethod
    def _ler(linhas):
        """O cabeçalho não é a 1ª linha (a 1ª é o título da aba): procura a
        linha que tem as duas colunas."""
        indice = None
        for linha in linhas:
            cabecalho = [_normalizar(valor) for valor in linha]
            if COLUNA_INEP in cabecalho and COLUNA_ASSINATURA in cabecalho:
                indice = (cabecalho.index(COLUNA_INEP), cabecalho.index(COLUNA_ASSINATURA))
                break
        if indice is None:
            raise CommandError(f"Colunas '{COLUNA_INEP}' e '{COLUNA_ASSINATURA}' nao encontradas na aba.")

        valores = {}
        invalidas = 0
        for linha in linhas:
            inep = _texto(linha[indice[0]]) if indice[0] < len(linha) else ""
            if not inep:
                continue
            try:
                valores[inep] = _valor_decimal(linha[indice[1]] if indice[1] < len(linha) else None)
            except (InvalidOperation, ValueError):
                invalidas += 1
        return valores, invalidas
