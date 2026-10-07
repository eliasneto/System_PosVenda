import csv
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.escolas.models import PlanilhaRelatorioEaceMip
from apps.escolas.services import (
    MOTIVO_IMPORTACAO_INCLUIDO,
    MOTIVO_IMPORTACAO_JA_EM_LOTE,
    MOTIVO_IMPORTACAO_NAO_CADASTRADO,
    MOTIVO_IMPORTACAO_STATUS,
    MOTIVO_IMPORTACAO_VALOR,
    _normalizar_cabecalho,
    importar_lotes_mip_em_massa,
)

COLUNA_VALOR_LIBERADO = "VALOR LIBERADO ACS"
COLUNAS_EXIGIDAS = set(PlanilhaRelatorioEaceMip.COLUNAS_OBRIGATORIAS) | {COLUNA_VALOR_LIBERADO}

ROTULOS_MOTIVO = {
    MOTIVO_IMPORTACAO_INCLUIDO: "Incluído em LOTE",
    MOTIVO_IMPORTACAO_JA_EM_LOTE: "Já estava em LOTE (não mexido)",
    MOTIVO_IMPORTACAO_NAO_CADASTRADO: "INEP não cadastrado no sistema",
    MOTIVO_IMPORTACAO_STATUS: 'Status (MIP) diferente de "Aguardando Validação EACE"',
    MOTIVO_IMPORTACAO_VALOR: "Valor da planilha, IXC e EACE não batem",
}


class _SimulacaoConcluida(Exception):
    """Desfaz a transação no fim da simulação."""


def _abas_da_planilha(caminho):
    """{nome da aba: lista de linhas (listas de valores)} — `.xlsb` pelo
    `pyxlsb` (formato da "BASE CONSOLIDADA MIP"), `.xlsx` pelo `openpyxl`."""
    if caminho.suffix.lower() == ".xlsb":
        from pyxlsb import open_workbook

        abas = {}
        with open_workbook(str(caminho)) as workbook:
            for nome in workbook.sheets:
                with workbook.get_sheet(nome) as aba:
                    abas[nome] = [[celula.v for celula in linha] for linha in aba.rows()]
        return abas
    workbook = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        return {nome: [list(linha) for linha in workbook[nome].iter_rows(values_only=True)] for nome in workbook.sheetnames}
    finally:
        workbook.close()


def _inep(valor):
    try:
        return str(int(valor)).zfill(8)
    except (TypeError, ValueError):
        return None


def _decimal(valor):
    try:
        return Decimal(str(valor)).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return Decimal("0.00")


def _linhas_por_inep(abas):
    """Agrupa por INEP (coluna "Projeto") a 1ª aba com as colunas do
    Relatório EACE (MIP) + "Valor Liberado ACS" — mesmas chaves de
    `_agrupar_linhas_relatorio_eace_mip_por_inep`."""
    for nome, linhas in abas.items():
        if not linhas:
            continue
        indice = {_normalizar_cabecalho(valor): posicao for posicao, valor in enumerate(linhas[0])}
        if not COLUNAS_EXIGIDAS <= set(indice):
            continue
        agrupado = {}
        for linha in linhas[1:]:
            linha = list(linha) + [None] * (len(indice) - len(linha))
            inep = _inep(linha[indice["PROJETO"]])
            if not inep:
                continue
            agrupado.setdefault(inep, []).append({
                "cod_fornecedor": linha[indice["COD FORNECEDOR"]],
                "descricao": linha[indice["DESCRIÇÃO DO ITEM"]],
                "qtde_produto": linha[indice["QTDE PRODUTO"]],
                "uf": linha[indice["UF"]],
                "cidade": linha[indice["CIDADE"]],
                "data_emissao_acs": linha[indice["DATA EMISSÃO ACS"]],
                "valor_liberado": _decimal(linha[indice[COLUNA_VALOR_LIBERADO]]),
            })
        return nome, agrupado
    raise CommandError(
        "Nenhuma aba tem as colunas exigidas: " + ", ".join(sorted(COLUNAS_EXIGIDAS)) + "."
    )


def _ineps_da_aba_faturamento(abas):
    """INEPs da aba "FATURAMENTO" (coluna "INEP"), quando existir — só
    para o relatório listar os que não estão na aba de itens."""
    for nome, linhas in abas.items():
        if _normalizar_cabecalho(nome) != "FATURAMENTO" or not linhas:
            continue
        cabecalho = [_normalizar_cabecalho(valor) for valor in linhas[0]]
        if "INEP" not in cabecalho:
            return set()
        posicao = cabecalho.index("INEP")
        return {inep for inep in (_inep(linha[posicao]) for linha in linhas[1:] if len(linha) > posicao) if inep}
    return set()


class Command(BaseCommand):
    """Pedido do usuário (2026-09-29): cria LOTEs do MIP a partir da
    planilha "BASE CONSOLIDADA MIP" (`apps.escolas.services.
    importar_lotes_mip_em_massa`) — 1 LOTE por Estado + Município, já em
    "Processo Concluído", marcado como importação em massa.

    Simulação por padrão (faz tudo e desfaz no fim, para o relatório ser
    exatamente o que seria gravado); `--aplicar` grava de verdade."""

    help = "Cria LOTEs do MIP a partir da BASE CONSOLIDADA MIP. Simulação por padrão; use --aplicar para gravar."

    def add_arguments(self, parser):
        parser.add_argument("arquivo", help='Caminho da planilha (ex.: "doc/BASE CONSOLIDADA MIP 2026.xlsb").')
        parser.add_argument("--usuario", required=True, help="Username gravado como autor dos LOTEs e dos históricos.")
        parser.add_argument("--relatorio", help="Caminho de um .csv com o resultado de cada INEP.")
        parser.add_argument("--aplicar", action="store_true", help="Grava de verdade (sem isto, só simula).")

    def handle(self, *args, **opcoes):
        caminho = Path(opcoes["arquivo"])
        if not caminho.is_file():
            raise CommandError(f"Arquivo não encontrado: {caminho}")
        try:
            usuario = get_user_model().objects.get(username=opcoes["usuario"])
        except get_user_model().DoesNotExist:
            raise CommandError(f"Usuário não encontrado: {opcoes['usuario']}")

        abas = _abas_da_planilha(caminho)
        nome_aba, linhas_por_inep = _linhas_por_inep(abas)
        ineps_so_faturamento = sorted(_ineps_da_aba_faturamento(abas) - set(linhas_por_inep))
        aplicar = opcoes["aplicar"]
        self.stdout.write(
            f"Arquivo: {caminho.name} | aba: {nome_aba} | INEPs: {len(linhas_por_inep)}\n"
            f"Modo: {'APLICAR (grava de verdade)' if aplicar else 'SIMULACAO (nada é gravado)'}"
        )

        resultado = None
        try:
            with transaction.atomic():
                resultado = importar_lotes_mip_em_massa(linhas_por_inep, usuario, nome_arquivo=caminho.name)
                resumo_lotes = [
                    (str(lote), lote.estado, lote.municipio, lote.escolas.count()) for lote in resultado["lotes"]
                ]
                if not aplicar:
                    raise _SimulacaoConcluida
        except _SimulacaoConcluida:
            pass

        ineps = resultado["ineps"]
        contagem = Counter(registro["motivo"] for registro in ineps)
        self.stdout.write("\nResumo por INEP:")
        for motivo, rotulo in ROTULOS_MOTIVO.items():
            self.stdout.write(f"  {rotulo}: {contagem.get(motivo, 0)}")
        preenchidos = sum(1 for registro in ineps if registro["lado3_preenchido"])
        self.stdout.write(f"  (dos incluídos, com Lado EACE (MIP) preenchido pela planilha: {preenchidos})")
        corrigidos = [registro for registro in ineps if registro.get("municipio_anterior")]
        self.stdout.write(f"  (Município corrigido para juntar grafias diferentes: {len(corrigidos)})")
        for grafia_anterior, grafia_nova in sorted(Counter(
            (registro["municipio_anterior"], registro["municipio"]) for registro in corrigidos
        )):
            self.stdout.write(f"     {grafia_anterior} -> {grafia_nova}")
        if ineps_so_faturamento:
            self.stdout.write(
                f"  Só na aba FATURAMENTO, sem itens (não importados): {len(ineps_so_faturamento)}"
            )

        titulo_lotes = "LOTEs criados" if aplicar else "LOTEs que seriam criados (numeração da simulação)"
        self.stdout.write(f"\n{titulo_lotes}: {len(resumo_lotes)}")
        for nome, estado, municipio, quantidade in resumo_lotes:
            self.stdout.write(f"  {nome} — {municipio}/{estado} — {quantidade} INEP(s)")

        fora = [registro for registro in ineps if registro["motivo"] not in (MOTIVO_IMPORTACAO_INCLUIDO, MOTIVO_IMPORTACAO_JA_EM_LOTE)]
        if fora:
            self.stdout.write("\nINEPs que ficaram de fora:")
            for registro in fora:
                self.stdout.write(
                    f"  {registro['inep']} — {ROTULOS_MOTIVO[registro['motivo']]}"
                    f" (planilha {registro['total_planilha']} | IXC {registro['total_ixc']} | EACE {registro['total_eace']}"
                    f"{' | ' + registro['status_mip'] if registro.get('status_mip') else ''})"
                )

        if opcoes["relatorio"]:
            with open(opcoes["relatorio"], "w", newline="", encoding="utf-8-sig") as arquivo_csv:
                escritor = csv.writer(arquivo_csv, delimiter=";")
                escritor.writerow([
                    "INEP", "UF", "Municipio", "Resultado", "LOTE", "Total planilha", "Total IXC", "Total EACE",
                    "Lado EACE preenchido pela planilha", "Status (MIP)", "Municipio anterior",
                ])
                for registro in ineps:
                    escritor.writerow([
                        registro["inep"], registro.get("estado", ""), registro.get("municipio", ""),
                        ROTULOS_MOTIVO[registro["motivo"]], registro["lote"], registro["total_planilha"],
                        registro["total_ixc"] or "", registro["total_eace"] or "",
                        "Sim" if registro["lado3_preenchido"] else "", registro.get("status_mip", ""),
                        registro.get("municipio_anterior", ""),
                    ])
                for inep in ineps_so_faturamento:
                    escritor.writerow([inep, "", "", "Só na aba FATURAMENTO (sem itens)", "", "", "", "", "", "", ""])
            self.stdout.write(f"\nRelatório: {opcoes['relatorio']}")
