import datetime
import io
import tempfile
import zipfile
from pathlib import Path
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse

from apps.escolas.bdo import (
    PlanilhaBdoError, dias_no_mes_referencia, dias_proporcionais, importar_planilha_bdo, importar_pro_rata, ler_planilha_bdo,
    ler_planilha_pro_rata, valor_proporcional,
)
from apps.escolas.models import (
    BoletimBdo, Escola, InepBdo, LeituraEaceBdo, LogStatusDivergenciaBdo, MedicaoBdo, MunicipioBdo,
    PrimeiroFaturamentoInep,
)
from apps.ri.models import KitPadrao, Ri

User = get_user_model()

CABECALHO = (
    "Fornecedor", "CNPJ", "Boletim de desempenho", "Município", "UF", "Valor serviços", "Valor créditos",
    "Valor descontos", "Valor descontos proporcionais", "Valor total", "Status", "Qntd escola", "INEPs",
)


def _linha(boletim, municipio, uf, valor_total, ineps, servicos="0,00", creditos="0,00", descontos="0,00",
           proporcionais="0,00"):
    return (
        "MEGA", "07.525.262/0004-61", boletim, municipio, uf, servicos, creditos, descontos, proporcionais,
        valor_total, "Pendente", "", ineps,
    )


def _planilha(linhas, cabecalho=CABECALHO, nome="FR_3658.xlsx"):
    workbook = openpyxl.Workbook()
    aba = workbook.active
    aba.title = "FR_3658"
    aba.append(cabecalho)
    for linha in linhas:
        aba.append(linha)
    conteudo = io.BytesIO()
    workbook.save(conteudo)
    return SimpleUploadedFile(
        nome, conteudo.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


LINHAS_3658 = [
    _linha("3658", "AMERICANA", "SP", "49,92", "35058592"),
    _linha("3658", "ATIBAIA", "SP", "277,89", "35011006"),
    _linha(
        "3658", "Campo Grande", "MS", "1.094,45", "50031287, 50005111, 50033433, 50028340",
        servicos="1.462,48", descontos="2,41", proporcionais="365,62",
    ),
    _linha("3658", "Brasília", "DF", "31.183,68", "53002814, 53002474"),
]


class LerPlanilhaBdoTests(TestCase):
    def test_le_valor_brasileiro_e_ineps_separados_por_virgula(self):
        boletins = ler_planilha_bdo(_planilha(LINHAS_3658))
        self.assertEqual(list(boletins), ["3658"])
        campo_grande = boletins["3658"][2]
        self.assertEqual(
            [campo_grande[campo] for campo in MunicipioBdo.CAMPOS_VALORES],
            [Decimal("1462.48"), Decimal("0.00"), Decimal("2.41"), Decimal("365.62"), Decimal("1094.45")],
        )
        self.assertEqual(campo_grande["ineps"], ["50031287", "50005111", "50033433", "50028340"])

    def test_aceita_valor_numerico_do_excel_e_cabecalho_sem_acento(self):
        cabecalho = tuple("Municipio" if coluna == "Município" else coluna for coluna in CABECALHO)
        boletins = ler_planilha_bdo(_planilha([_linha(3658, "AMERICANA", "sp", 49.92, 35058592)], cabecalho))
        linha = boletins["3658"][0]
        self.assertEqual((linha["uf"], linha["valor_total"], linha["ineps"]), ("SP", Decimal("49.92"), ["35058592"]))

    def test_aceita_colunas_a_mais_a_menos_e_fora_de_ordem(self):
        """Só as colunas usadas pelo BDO importam: as outras podem faltar,
        sobrar ou vir em qualquer ordem, e o nome não depende de
        maiúscula, acento ou espaço extra."""
        cabecalho = (
            "  ineps ", "Coluna Nova", "VALOR TOTAL", "uf", "Valor descontos proporcionais", "Observação",
            "Valor Descontos", "valor creditos", "Valor  Servicos", "MUNICIPIO", "Boletim de Desempenho",
        )
        linha = (
            "35058592, 35004330", "x", "49,92", "sp", "316,98", "qualquer coisa",
            "13,48", "0,00", "380,38", "AMERICANA", "3658",
        )
        boletins = ler_planilha_bdo(_planilha([linha], cabecalho))
        [americana] = boletins["3658"]
        self.assertEqual((americana["municipio"], americana["uf"]), ("AMERICANA", "SP"))
        self.assertEqual(americana["ineps"], ["35058592", "35004330"])
        self.assertEqual(
            [americana[campo] for campo in MunicipioBdo.CAMPOS_VALORES],
            [Decimal("380.38"), Decimal("0.00"), Decimal("13.48"), Decimal("316.98"), Decimal("49.92")],
        )

    def test_sem_coluna_obrigatoria_da_erro(self):
        for faltando in ("INEPs", "Valor créditos", "Valor descontos proporcionais"):
            with self.subTest(faltando=faltando):
                cabecalho = tuple(coluna for coluna in CABECALHO if coluna != faltando)
                with self.assertRaisesMessage(PlanilhaBdoError, "colunas obrigatórias"):
                    ler_planilha_bdo(_planilha([], cabecalho))

    def test_linha_sem_inep_ou_valor_invalido_da_erro_com_numero_da_linha(self):
        linhas = [LINHAS_3658[0], _linha("3658", "ATIBAIA", "SP", "abc", "", creditos="")]
        with self.assertRaisesMessage(PlanilhaBdoError, "linha 3 (Valor créditos, Valor total, INEPs)"):
            ler_planilha_bdo(_planilha(linhas))

    def test_planilha_sem_dados_da_erro(self):
        with self.assertRaisesMessage(PlanilhaBdoError, "nenhuma linha de dados"):
            ler_planilha_bdo(_planilha([]))


class ImportarPlanilhaBdoTests(TestCase):
    def test_boletim_ja_importado_e_recusado_sem_alterar_os_dados(self):
        importar_planilha_bdo(ler_planilha_bdo(_planilha(LINHAS_3658)), "RI", "FR_3658.xlsx", None)
        with self.assertRaisesMessage(PlanilhaBdoError, "Boletim de desempenho já importado: 3658"):
            importar_planilha_bdo(ler_planilha_bdo(_planilha(LINHAS_3658[:1])), "RI", "FR_3658.xlsx", None)
        self.assertEqual(BoletimBdo.objects.count(), 1)
        self.assertEqual(MunicipioBdo.objects.count(), 4)
        self.assertEqual(InepBdo.objects.count(), 8)

    def test_planilha_com_boletim_novo_e_repetido_nao_grava_nenhum(self):
        importar_planilha_bdo(ler_planilha_bdo(_planilha(LINHAS_3658)), "RI", "FR_3658.xlsx", None)
        linhas = [LINHAS_3658[0], _linha("3700", "AMERICANA", "SP", "10,00", "35058592")]
        with self.assertRaises(PlanilhaBdoError):
            importar_planilha_bdo(ler_planilha_bdo(_planilha(linhas)), "RI", "misto.xlsx", None)
        self.assertFalse(BoletimBdo.objects.filter(numero="3700").exists())

    def test_boletim_novo_e_gravado_sem_mexer_nos_anteriores(self):
        importar_planilha_bdo(ler_planilha_bdo(_planilha(LINHAS_3658)), "RI", "FR_3658.xlsx", None)
        importar_planilha_bdo(
            ler_planilha_bdo(_planilha([_linha("3700", "AMERICANA", "SP", "10,00", "35058592")])), "RI", "FR_3700.xlsx", None,
        )
        self.assertEqual(sorted(BoletimBdo.objects.values_list("numero", flat=True)), ["3658", "3700"])
        self.assertEqual(InepBdo.objects.count(), 9)


class DiasProporcionaisTests(TestCase):
    def test_conta_ate_o_dia_30_sem_o_dia_da_ativacao(self):
        self.assertEqual(dias_proporcionais(datetime.date(2026, 9, 15)), 15)
        self.assertEqual(dias_proporcionais(datetime.date(2026, 9, 1)), 29)
        self.assertEqual(dias_proporcionais(datetime.date(2026, 9, 30)), 0)

    def test_fevereiro_tambem_vai_ate_o_dia_30(self):
        self.assertEqual(dias_proporcionais(datetime.date(2026, 2, 20)), 10)

    def test_dia_31_vai_para_o_mes_seguinte_cheio(self):
        self.assertEqual(dias_proporcionais(datetime.date(2026, 8, 31)), 30)

    def test_sem_data_de_ativacao_fica_vazio(self):
        self.assertIsNone(dias_proporcionais(None))


class DiasNoMesReferenciaTests(TestCase):
    SETEMBRO = datetime.date(2026, 9, 1)

    def test_ativado_antes_do_mes_conta_cheio(self):
        self.assertEqual(dias_no_mes_referencia(datetime.date(2026, 7, 10), self.SETEMBRO), 30)
        self.assertEqual(dias_no_mes_referencia(datetime.date(2026, 8, 31), self.SETEMBRO), 30)
        self.assertEqual(dias_no_mes_referencia(datetime.date(2025, 12, 5), self.SETEMBRO), 30)

    def test_ativado_no_mes_conta_proporcional(self):
        self.assertEqual(dias_no_mes_referencia(datetime.date(2026, 9, 15), self.SETEMBRO), 15)
        self.assertEqual(dias_no_mes_referencia(datetime.date(2026, 10, 31), datetime.date(2026, 10, 1)), 0)

    def test_ativado_depois_do_mes_ou_sem_data_fica_vazio(self):
        self.assertIsNone(dias_no_mes_referencia(datetime.date(2026, 10, 1), self.SETEMBRO))
        self.assertIsNone(dias_no_mes_referencia(None, self.SETEMBRO))

    def test_boletim_sem_mes_de_referencia_conta_como_antes(self):
        self.assertEqual(dias_no_mes_referencia(datetime.date(2026, 7, 10), None), 20)


class ValorProporcionalTests(TestCase):
    def test_valor_mensal_vezes_dias_sobre_30_em_centavos(self):
        self.assertEqual(valor_proporcional(Decimal("380.38"), 15), Decimal("190.19"))
        self.assertEqual(valor_proporcional(Decimal("365.62"), 29), Decimal("353.43"))
        self.assertEqual(valor_proporcional(Decimal("760.00"), 30), Decimal("760.00"))
        self.assertEqual(valor_proporcional(Decimal("760.00"), 0), Decimal("0.00"))

    def test_sem_valor_ou_sem_dias_fica_vazio(self):
        self.assertIsNone(valor_proporcional(None, 15))
        self.assertIsNone(valor_proporcional(Decimal("760.00"), None))


class BdoViewTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="bdo-admin", password="senha-teste-123", perfil=User.PERFIL_ADMINISTRADOR,
        )
        self.analista = User.objects.create_user(
            username="bdo-analista", password="senha-teste-123", perfil=User.PERFIL_ANALISTA,
        )

    def _importar(self, linhas=LINHAS_3658, tipo="RI"):
        importar_planilha_bdo(ler_planilha_bdo(_planilha(linhas)), tipo, "FR_3658.xlsx", self.admin)

    def test_anonimo_vai_pro_login(self):
        resp = self.client.get(reverse("bdo"))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse("login"), resp["Location"])

    def test_sem_boletim_mostra_estado_vazio(self):
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"))
        self.assertContains(resp, "Nenhum Boletim importado")
        self.assertNotContains(resp, 'enctype="multipart/form-data"')
        self.assertNotContains(resp, f'href="{reverse("relatorio_bdo")}"')

    def test_tela_do_bdo_nao_tem_mais_upload_e_admin_ve_o_atalho(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("bdo"))
        self.assertNotContains(resp, 'enctype="multipart/form-data"')
        self.assertContains(resp, f'href="{reverse("relatorio_bdo")}"')

    def test_grid_por_inep_repete_o_valor_do_municipio_e_total_nao_duplica(self):
        self._importar()
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["page_obj"].paginator.count, 8)
        self.assertEqual(resp.context["total_ineps"], 8)
        self.assertEqual(resp.context["total_municipios"], 4)
        self.assertEqual(resp.context["valor_total"], Decimal("32605.94"))
        self.assertContains(resp, "53002814")
        self.assertContains(resp, "Boletim de desempenho")

    def test_grid_por_inep_mostra_a_data_de_ativacao_do_ri_atual_tambem_no_re(self):
        escola = Escola.objects.create(inep="35058592", nome="Escola Americana", estado="SP", municipio="Americana")
        Ri.objects.create(escola=escola, data_ativacao=datetime.date(2026, 8, 1))
        Ri.objects.create(escola=escola, data_ativacao=datetime.date(2026, 9, 5))
        self._importar(tipo="RI")
        importar_planilha_bdo(
            ler_planilha_bdo(_planilha([_linha("3700", "AMERICANA", "SP", "10,00", "35058592")])),
            "RE", "FR_3700.xlsx", self.admin,
        )
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"))
        datas = {
            (linha.municipio_bdo.boletim.tipo, linha.inep): linha.data_ativacao for linha in resp.context["page_obj"]
        }
        self.assertEqual(datas[("RI", "35058592")], datetime.date(2026, 9, 5))
        self.assertEqual(datas[("RE", "35058592")], datetime.date(2026, 9, 5))
        self.assertIsNone(datas[("RI", "53002814")])
        self.assertContains(resp, "Dias proporcionais")
        self.assertContains(resp, "Data de ativação")
        self.assertContains(resp, "05/09/2026", count=2)

    def test_valor_assinatura_conforme_o_tipo_do_boletim(self):
        """RE: assinatura da escola; RI: manutenção do LPU no Lote da escola."""
        KitPadrao.objects.create(
            descricao="Manutenção de Rede Interna (serviços, materiais e equipamentos)", lote=9,
            unidade="Escola/Mês", valor_servico=Decimal("380.38"),
        )
        KitPadrao.objects.create(
            descricao="Manutenção de Rede Interna (serviços, materiais e equipamentos)", lote=11,
            unidade="Escola/Mês", valor_servico=Decimal("365.62"),
        )
        Escola.objects.create(inep="35058592", nome="Americana", lote=9, valor_assinatura_re=Decimal("760.00"))
        Escola.objects.create(inep="35011006", nome="Atibaia", lote=11)
        Escola.objects.create(inep="50031287", nome="Campo Grande", valor_assinatura_re=Decimal("370.00"))
        self._importar(tipo="RI")
        importar_planilha_bdo(
            ler_planilha_bdo(_planilha([_linha("3700", "AMERICANA", "SP", "10,00", "35058592, 35011006")])),
            "RE", "FR_3700.xlsx", self.admin,
        )
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"))
        valores = {
            (linha.municipio_bdo.boletim.tipo, linha.inep): linha.valor_assinatura for linha in resp.context["page_obj"]
        }
        self.assertEqual(valores[("RI", "35058592")], Decimal("380.38"))
        self.assertEqual(valores[("RI", "35011006")], Decimal("365.62"))
        self.assertEqual(valores[("RE", "35058592")], Decimal("760.00"))
        self.assertIsNone(valores[("RE", "35011006")])  # escola sem assinatura RE
        self.assertIsNone(valores[("RI", "50031287")])  # escola sem Lote
        self.assertIsNone(valores[("RI", "53002814")])  # INEP sem escola
        self.assertContains(resp, "Valor assinatura")
        self.assertContains(resp, "R$ 760,00")

    def test_valor_assinatura_vazio_em_boletim_sem_tipo(self):
        Escola.objects.create(inep="35058592", nome="Americana", valor_assinatura_re=Decimal("760.00"))
        self._importar(tipo="")
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"inep": "35058592"})
        self.assertIsNone(resp.context["page_obj"][0].valor_assinatura)

    def test_filtro_de_data_de_ativacao_usa_o_ri_atual(self):
        americana = Escola.objects.create(inep="35058592", nome="Escola Americana", estado="SP", municipio="Americana")
        Ri.objects.create(escola=americana, data_ativacao=datetime.date(2026, 9, 15))
        atibaia = Escola.objects.create(inep="35011006", nome="Escola Atibaia", estado="SP", municipio="Atibaia")
        Ri.objects.create(escola=atibaia, data_ativacao=datetime.date(2026, 9, 20))
        Ri.objects.create(escola=atibaia, data_ativacao=datetime.date(2026, 8, 10))
        self._importar()
        self.client.force_login(self.analista)

        resp = self.client.get(reverse("bdo"), {"data_inicial": "01/09/2026", "data_final": "30/09/2026"})
        self.assertEqual([linha.inep for linha in resp.context["page_obj"]], ["35058592"])
        self.assertEqual(resp.context["total_municipios"], 1)
        self.assertContains(resp, "data_inicial=01%2F09%2F2026")

        resp = self.client.get(reverse("bdo"), {"data_inicial": "01/08/2026", "agrupar": "uf"})
        self.assertEqual(resp.context["total_ineps"], 2)

        resp = self.client.get(reverse("bdo"), {"data_inicial": "30/09/2026", "data_final": "01/09/2026"})
        self.assertEqual(resp.context["page_obj"].paginator.count, 8)
        self.assertContains(resp, "A data inicial não pode ser depois da data final")

    def test_agrupar_por_municipio_mostra_quantidade_de_ineps_e_valor(self):
        self._importar()
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"agrupar": "municipio"})
        linhas = {linha["municipio"]: (linha["quantidade_ineps"], linha["valor_total"]) for linha in resp.context["page_obj"]}
        self.assertEqual(linhas["CAMPO GRANDE"], (4, Decimal("1094.45")))
        self.assertContains(resp, "R$ 1.462,48")
        self.assertContains(resp, "R$ 365,62")
        self.assertEqual(linhas["BRASÍLIA"], (2, Decimal("31183.68")))

    def _escola_ri(self, inep, lote, data_ativacao, assinatura_re=None):
        escola = Escola.objects.create(inep=inep, nome=f"Escola {inep}", lote=lote, valor_assinatura_re=assinatura_re)
        Ri.objects.create(escola=escola, data_ativacao=data_ativacao)

    def test_valor_proporcional_por_inep_e_total_mega_por_municipio(self):
        KitPadrao.objects.create(
            descricao="Manutenção de Rede Interna (serviços, materiais e equipamentos)", lote=11,
            unidade="Escola/Mês", valor_servico=Decimal("365.62"),
        )
        self._escola_ri("50031287", 11, datetime.date(2026, 9, 15))   # 15 dias: 182,81
        self._escola_ri("50005111", 11, datetime.date(2026, 9, 1))    # 29 dias: 353,43
        self._escola_ri("50033433", 11, datetime.date(2026, 8, 31))   # 30 dias: 365,62
        self._escola_ri("53002814", 11, datetime.date(2026, 9, 10))   # 20 dias: 243,75
        self._escola_ri("53002474", 11, datetime.date(2026, 9, 10))
        self._importar(tipo="RI")
        self.client.force_login(self.analista)

        resp = self.client.get(reverse("bdo"), {"inep": "50031287"})
        self.assertEqual(resp.context["page_obj"][0].valor_proporcional, Decimal("182.81"))
        self.assertContains(resp, "Valor proporcional")
        self.assertContains(resp, "R$ 182,81")

        resp = self.client.get(reverse("bdo"), {"agrupar": "municipio"})
        linhas = {linha["municipio"]: linha for linha in resp.context["page_obj"]}
        # 50028340 sem escola: fica fora da soma e é sinalizado
        self.assertEqual(linhas["CAMPO GRANDE"]["total_mega"], Decimal("901.86"))
        self.assertEqual(linhas["CAMPO GRANDE"]["ineps_sem_valor_mega"], 1)
        self.assertEqual(linhas["BRASÍLIA"]["total_mega"], Decimal("487.50"))
        self.assertEqual(linhas["BRASÍLIA"]["ineps_sem_valor_mega"], 0)
        self.assertIsNone(linhas["AMERICANA"]["total_mega"])
        self.assertContains(resp, "Total MEGA")
        self.assertContains(resp, "1 INEP(s) sem valor")

    def test_total_mega_com_mes_de_referencia_cobra_cheio_quem_ativou_antes(self):
        self._escola_ri("35058592", None, datetime.date(2026, 7, 20), assinatura_re=Decimal("600.00"))  # cheio
        self._escola_ri("35011006", None, datetime.date(2026, 9, 25), assinatura_re=Decimal("300.00"))  # 5 dias
        self._escola_ri("35022222", None, datetime.date(2026, 10, 2), assinatura_re=Decimal("900.00"))  # depois
        importar_planilha_bdo(
            ler_planilha_bdo(_planilha([_linha("3700", "AMERICANA", "SP", "650,00", "35058592, 35011006, 35022222")])),
            "RE", "FR_3700.xlsx", self.admin, datetime.date(2026, 9, 18),
        )
        self.assertEqual(BoletimBdo.objects.get().mes_referencia, datetime.date(2026, 9, 1))
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"agrupar": "municipio"})
        linha = resp.context["page_obj"][0]
        self.assertEqual(linha["total_mega"], Decimal("650.00"))
        self.assertEqual(linha["ineps_sem_valor_mega"], 1)
        self.assertEqual(linha["diferenca_mega"], Decimal("0.00"))  # igual ao Valor total do cliente
        resp = self.client.get(reverse("bdo"))
        dias = {linha.inep: linha.dias_proporcionais for linha in resp.context["page_obj"]}
        self.assertEqual(dias, {"35058592": 30, "35011006": 5, "35022222": None})

    def test_dif_do_total_mega_e_com_o_valor_total_do_cliente(self):
        self._escola_ri("35058592", None, datetime.date(2026, 7, 20), assinatura_re=Decimal("600.00"))  # cheio
        self._escola_ri("35011006", None, datetime.date(2026, 9, 25), assinatura_re=Decimal("300.00"))  # 50,00
        importar_planilha_bdo(
            ler_planilha_bdo(_planilha([_linha(
                "3700", "AMERICANA", "SP", "645,00", "35058592, 35011006",
                servicos="900,00", descontos="5,00", proporcionais="250,00",
            )])),
            "RE", "FR_3700.xlsx", self.admin, datetime.date(2026, 9, 1),
        )
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"agrupar": "municipio"})
        linha = resp.context["page_obj"][0]
        self.assertEqual(linha["total_mega"], Decimal("650.00"))
        self.assertEqual(linha["valor_conferencia_mega"], Decimal("645.00"))
        self.assertEqual(linha["diferenca_mega"], Decimal("5.00"))  # o Valor descontos do cliente
        self.assertContains(resp, "Dif. R$ 5,00")
        self.assertContains(resp, "Cliente pagou a menos: R$ 5,00")
        MunicipioBdo.objects.update(valor_total=Decimal("660.00"))
        resp = self.client.get(reverse("bdo"), {"agrupar": "municipio"})
        self.assertContains(resp, "Cliente pagou a mais: R$ 10,00")
        self.assertNotContains(resp, "Cliente pagou a menos")
        MunicipioBdo.objects.update(valor_total=Decimal("650.00"))
        resp = self.client.get(reverse("bdo"), {"agrupar": "municipio"})
        self.assertEqual(resp.context["page_obj"][0]["diferenca_mega"], Decimal("0.00"))
        self.assertNotContains(resp, "Dif. R$")
        self.assertNotContains(resp, "Cliente pagou a")

    def test_inep_contestado_deixa_o_recebido_no_total_mega_e_a_diferenca_segue_a_planilha(self):
        self._escola_ri("35058592", None, datetime.date(2026, 7, 20), assinatura_re=Decimal("600.00"))  # cheio
        self._escola_ri("35011006", None, datetime.date(2026, 9, 25), assinatura_re=Decimal("300.00"))  # 50,00
        importar_planilha_bdo(
            ler_planilha_bdo(_planilha([_linha(
                "3700", "AMERICANA", "SP", "630,00", "35058592, 35011006",
                servicos="900,00", descontos="5,00", proporcionais="250,00",
            )])),
            "RE", "FR_3700.xlsx", self.admin, datetime.date(2026, 9, 1),
        )
        # Pro rata: desconto de 270,00 - o cliente pagou 30,00 dos 50,00 de 35011006.
        InepBdo.objects.filter(inep="35011006").update(
            status_divergencia=InepBdo.STATUS_CONTESTADO, valor_desconto_cliente=Decimal("270.00"),
            ativado_em_cliente=datetime.date(2026, 9, 25),
        )
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"agrupar": "municipio"})
        linha = resp.context["page_obj"][0]
        self.assertEqual(linha["total_mega"], Decimal("630.00"))  # 600,00 + 30,00 recebidos
        self.assertEqual(linha["ineps_contestados_mega"], 1)
        self.assertEqual(linha["total_contestado_mega"], Decimal("20.00"))
        self.assertEqual(linha["diferenca_mega"], Decimal("0.00"))  # o cliente pagou o recebido
        self.assertContains(resp, "1 contestado(s): faltando R$ 20,00")

        resp = self.client.get(reverse("bdo"), {"inep": "35011006"})
        contestado = resp.context["page_obj"][0]
        self.assertEqual((contestado.valor_recebido, contestado.valor_faltando), (Decimal("30.00"), Decimal("20.00")))
        self.assertContains(resp, 'id="valor-faltando-')
        self.assertContains(resp, "R$ 20,00")

        InepBdo.objects.update(status_divergencia=InepBdo.STATUS_CONTESTADO)
        linha = self.client.get(reverse("bdo"), {"agrupar": "municipio"}).context["page_obj"][0]
        # 35058592 sem Pro rata: o cliente cobrou o mês cheio, nada faltando
        self.assertEqual((linha["total_mega"], linha["total_contestado_mega"]), (Decimal("630.00"), Decimal("20.00")))
        self.assertEqual(linha["diferenca_mega"], Decimal("0.00"))

    def test_total_mega_usa_o_tipo_de_cada_boletim_e_segue_o_filtro_de_inep(self):
        self._escola_ri("35058592", None, datetime.date(2026, 9, 20), assinatura_re=Decimal("600.00"))
        self._escola_ri("35011006", None, datetime.date(2026, 9, 25), assinatura_re=Decimal("300.00"))
        importar_planilha_bdo(
            ler_planilha_bdo(_planilha([_linha("3700", "AMERICANA", "SP", "250,00", "35058592, 35011006")])),
            "RE", "FR_3700.xlsx", self.admin,
        )
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"agrupar": "municipio"})
        linha = resp.context["page_obj"][0]
        self.assertEqual(linha["total_mega"], Decimal("250.00"))  # 200,00 + 50,00
        resp = self.client.get(reverse("bdo"), {"agrupar": "municipio", "inep": "35011006"})
        self.assertEqual(resp.context["page_obj"][0]["total_mega"], Decimal("50.00"))

    def test_municipio_mostra_ineps_faturados_e_que_faltam_por_tipo(self):
        for inep, municipio in (("1", "São Paulo"), ("2", "SAO PAULO"), ("3", "são paulo"), ("4", "São Paulo"),
                                ("5", "Campinas")):
            Escola.objects.create(inep=inep, nome=f"Escola {inep}", estado="SP", municipio=municipio)
        Escola.objects.create(inep="6", nome="Escola 6", estado="RJ", municipio="São Paulo")  # outra UF
        importar_planilha_bdo(ler_planilha_bdo(_planilha([_linha("3700", "SÃO PAULO", "SP", "10,00", "1, 2")])),
                              "RE", "FR_3700.xlsx", self.admin)
        importar_planilha_bdo(ler_planilha_bdo(_planilha([_linha("3701", "Sao Paulo", "SP", "10,00", "2, 3")])),
                              "RE", "FR_3701.xlsx", self.admin)
        importar_planilha_bdo(ler_planilha_bdo(_planilha([_linha("3702", "SÃO PAULO", "SP", "10,00", "1")])),
                              "RI", "FR_3702.xlsx", self.admin)
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"agrupar": "municipio", "boletim": "3702"})
        linha = resp.context["page_obj"][0]
        self.assertEqual(linha["quantidade_escolas"], 4)
        self.assertEqual(linha["faturamento_por_tipo"], [
            {"tipo": "RE", "faturados": 3, "faltam": 1},  # 1, 2 e 3 em Boletins RE diferentes, cada um uma vez
            {"tipo": "RI", "faturados": 1, "faltam": 3},
        ])
        self.assertEqual(linha["situacao_faturamento"], "Proporcional")  # Boletim RI: faltam 3
        self.assertContains(resp, "Faturados / Faltam")
        self.assertContains(resp, "3 faturados")
        self.assertContains(resp, "Faltam 3")
        self.assertContains(resp, "Falta 1")

        # Grid por INEP: situação do município no tipo do Boletim de cada linha
        Escola.objects.filter(inep="4").delete()
        resp = self.client.get(reverse("bdo"), {"municipio": "SÃO PAULO"})
        situacoes = {(linha.municipio_bdo.boletim.numero, linha.inep): linha.situacao_faturamento for linha in resp.context["page_obj"]}
        self.assertEqual(situacoes, {
            ("3700", "1"): "Integral", ("3700", "2"): "Integral",  # RE: 1, 2 e 3 faturados
            ("3701", "2"): "Integral", ("3701", "3"): "Integral",
            ("3702", "1"): "Proporcional",  # RI: faltam 2 e 3
        })
        self.assertContains(resp, "Falta(m) 2 escola(s) do município no faturamento RI")
        self.assertContains(resp, "Todas as escolas do município no faturamento RE")

    def test_agrupar_por_uf_soma_cada_municipio_uma_vez(self):
        self._importar()
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"agrupar": "uf"})
        linhas = {linha["uf"]: linha for linha in resp.context["page_obj"]}
        self.assertEqual(linhas["SP"]["quantidade_municipios"], 2)
        self.assertEqual(linhas["SP"]["quantidade_ineps"], 2)
        self.assertEqual(linhas["SP"]["valor_total"], Decimal("327.81"))
        self.assertEqual(linhas["DF"]["valor_total"], Decimal("31183.68"))

    def test_valor_total_com_filtro_de_inep_ou_data_e_rateado_pelos_ineps_filtrados(self):
        """Pedido do usuário (2026-10-15): o valor é do município, então com
        filtro de INEP ou de data entra só a parte dos INEPs filtrados."""
        escola = Escola.objects.create(inep="50031287", nome="Campo Grande", estado="MS", municipio="Campo Grande")
        Ri.objects.create(escola=escola, data_ativacao=datetime.date(2026, 9, 10))
        self._importar()
        self.client.force_login(self.analista)
        self.assertEqual(self.client.get(reverse("bdo")).context["valor_total"], Decimal("32605.94"))
        resp = self.client.get(reverse("bdo"), {"inep": "50031287"})
        self.assertEqual(resp.context["valor_total"], Decimal("273.61"))  # 1.094,45 ÷ 4 INEPs
        resp = self.client.get(reverse("bdo"), {"inep": "5300"})
        self.assertEqual(resp.context["valor_total"], Decimal("31183.68"))  # os 2 INEPs de Brasília
        resp = self.client.get(reverse("bdo"), {"data_inicial": "01/09/2026", "data_final": "30/09/2026"})
        self.assertEqual(resp.context["valor_total"], Decimal("273.61"))
        resp = self.client.get(reverse("bdo"), {"uf": "SP"})
        self.assertEqual(resp.context["valor_total"], Decimal("327.81"))  # sem filtro de INEP: município inteiro

    def test_filtros_de_uf_municipio_e_inep(self):
        self._importar()
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"uf": "sp"})
        self.assertEqual(resp.context["total_ineps"], 2)
        resp = self.client.get(reverse("bdo"), {"municipio": "campo grande"})
        self.assertEqual(resp.context["total_ineps"], 4)
        resp = self.client.get(reverse("bdo"), {"inep": "53002814"})
        self.assertEqual(resp.context["total_ineps"], 1)
        self.assertEqual(resp.context["valor_total"], Decimal("15591.84"))  # rateio: 1 dos 2 INEPs de Brasília

    def test_lista_de_municipios_segue_a_uf_e_so_tem_os_do_grid(self):
        self._importar()
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"))
        self.assertEqual(resp.context["nomes_municipios"], ["AMERICANA", "ATIBAIA", "BRASÍLIA", "CAMPO GRANDE"])
        resp = self.client.get(reverse("bdo"), {"uf": "SP"})
        self.assertEqual(resp.context["nomes_municipios"], ["AMERICANA", "ATIBAIA"])

    def test_municipio_fora_da_uf_escolhida_e_ignorado(self):
        self._importar()
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"uf": "SP", "municipio": "Campo Grande"})
        self.assertEqual(resp.context["municipio_filtro"], "")
        self.assertEqual(resp.context["total_ineps"], 2)

    def _importar_dois_boletins(self):
        self._importar()
        self._importar([
            _linha("3700", "Americana", "SP", "10,00", "35058592"),
            _linha("3700", "ITU", "SP", "5,00", "35099999"),
        ])

    def test_sem_filtro_mostra_todos_os_boletins_cada_um_independente(self):
        self._importar_dois_boletins()
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"))
        self.assertIsNone(resp.context["boletim"])
        self.assertEqual(resp.context["page_obj"].paginator.count, 10)
        self.assertEqual(resp.context["total_boletins"], 2)
        self.assertEqual(resp.context["total_ineps"], 9)  # 35058592 está nos 2 Boletins: conta uma vez
        self.assertEqual(resp.context["total_municipios"], 6)
        self.assertEqual(resp.context["valor_total"], Decimal("32620.94"))

    def test_agrupar_por_municipio_mostra_o_mesmo_municipio_em_cada_boletim(self):
        self._importar_dois_boletins()
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"agrupar": "municipio"})
        americana = {
            linha["boletim"]: (linha["quantidade_ineps"], linha["valor_total"])
            for linha in resp.context["page_obj"] if linha["municipio"] == "AMERICANA"
        }
        self.assertEqual(americana, {"3658": (1, Decimal("49.92")), "3700": (1, Decimal("10.00"))})
        self.assertEqual(resp.context["page_obj"][0]["boletim"], "3700")

    def test_agrupar_por_uf_separa_os_boletins(self):
        self._importar_dois_boletins()
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"agrupar": "uf"})
        sp = {linha["boletim"]: linha for linha in resp.context["page_obj"] if linha["uf"] == "SP"}
        self.assertEqual(sp["3658"]["quantidade_municipios"], 2)
        self.assertEqual(sp["3658"]["valor_total"], Decimal("327.81"))
        self.assertEqual(sp["3700"]["quantidade_municipios"], 2)
        self.assertEqual(sp["3700"]["quantidade_ineps"], 2)
        self.assertEqual(sp["3700"]["valor_total"], Decimal("15.00"))

    def test_mesmo_municipio_repetido_no_boletim_soma_no_agrupamento(self):
        self._importar([
            _linha("3800", "ATIBAIA", "SP", "10,00", "35000001", servicos="12,00", descontos="2,00"),
            _linha("3800", "Atibaia", "SP", "5,00", "35000002", servicos="5,00"),
        ])
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"agrupar": "municipio"})
        [atibaia] = resp.context["page_obj"]
        self.assertEqual(atibaia["quantidade_ineps"], 2)
        self.assertEqual(atibaia["valores_em_ordem"], [
            Decimal("17.00"), Decimal("0.00"), Decimal("2.00"), Decimal("0.00"), Decimal("15.00"),
        ])

    def test_boletim_sem_os_valores_de_composicao_mostra_traco(self):
        self._importar()
        MunicipioBdo.objects.update(valor_servicos=None)
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"agrupar": "uf"})
        self.assertIsNone(resp.context["page_obj"][0]["valor_servicos"])
        self.assertContains(resp, '">—</span>')

    def test_filtro_de_boletim_mostra_so_aquele_boletim(self):
        self._importar_dois_boletins()
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"boletim": "3700"})
        self.assertEqual(resp.context["boletim"].numero, "3700")
        self.assertEqual(resp.context["total_ineps"], 2)
        self.assertEqual(resp.context["valor_total"], Decimal("15.00"))

    def test_relatorio_bdo_lista_os_boletins_importados(self):
        self._importar()
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("relatorio_bdo"))
        self.assertEqual(resp.status_code, 200)
        [boletim] = resp.context["boletins"]
        self.assertEqual((boletim.quantidade_municipios, boletim.quantidade_ineps), (4, 8))
        self.assertEqual(boletim.valor_total, Decimal("32605.94"))
        self.assertContains(resp, '<th class="px-4 py-2">Boletim de desempenho</th>', html=False)
        self.assertContains(resp, ">3658</span>")
        self.assertContains(resp, f'href="{reverse("bdo")}?boletim=3658"')
        self.assertContains(resp, 'enctype="multipart/form-data"')

    def test_relatorio_bdo_aparece_no_menu_administrador(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("relatorio_bdo"))
        self.assertContains(resp, f'href="{reverse("relatorio_bdo")}"')

    def test_administrador_importa_pela_tela(self):
        self.client.force_login(self.admin)
        resp = self.client.post(reverse("relatorio_bdo"), {"tipo": "RI", "mes_referencia": "2026-09", "arquivo": _planilha(LINHAS_3658)}, follow=True)
        self.assertRedirects(resp, reverse("relatorio_bdo"))
        self.assertContains(resp, "Boletim 3658 (RI) importado: 4 município(s), 8 INEP(s).")
        self.assertEqual(InepBdo.objects.count(), 8)

    def test_importacao_grava_o_mes_de_referencia_e_exige_o_campo(self):
        self.client.force_login(self.admin)
        resp = self.client.post(reverse("relatorio_bdo"), {"tipo": "RI", "arquivo": _planilha(LINHAS_3658)}, follow=True)
        self.assertContains(resp, "Informe o mês de referência da planilha.")
        self.assertFalse(BoletimBdo.objects.exists())
        self.client.post(
            reverse("relatorio_bdo"), {"tipo": "RI", "mes_referencia": "2026-09", "arquivo": _planilha(LINHAS_3658)},
        )
        self.assertEqual(BoletimBdo.objects.get().mes_referencia, datetime.date(2026, 9, 1))
        resp = self.client.get(reverse("relatorio_bdo"))
        self.assertContains(resp, 'value="2026-09"')

    def test_administrador_altera_o_mes_de_referencia_de_boletim_ja_importado(self):
        self._importar()
        boletim = BoletimBdo.objects.get()
        self.assertIsNone(boletim.mes_referencia)
        url = reverse("bdo_mes_referencia", args=[boletim.pk])
        self.client.force_login(self.admin)
        resp = self.client.post(url, {"mes_referencia": "2026-08"}, follow=True)
        self.assertRedirects(resp, reverse("relatorio_bdo"))
        self.assertContains(resp, "Boletim 3658: mês de referência 08/2026.")
        boletim.refresh_from_db()
        self.assertEqual(boletim.mes_referencia, datetime.date(2026, 8, 1))
        resp = self.client.post(url, {"mes_referencia": "agosto"}, follow=True)
        self.assertContains(resp, "Informe o mês de referência no formato mês/ano.")
        boletim.refresh_from_db()
        self.assertEqual(boletim.mes_referencia, datetime.date(2026, 8, 1))

    def test_so_administrador_altera_o_mes_de_referencia(self):
        self._importar()
        url = reverse("bdo_mes_referencia", args=[BoletimBdo.objects.get().pk])
        self.client.force_login(self.analista)
        self.assertEqual(self.client.post(url, {"mes_referencia": "2026-08"}).status_code, 403)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertIsNone(BoletimBdo.objects.get().mes_referencia)

    def test_importacao_grava_o_tipo_escolhido(self):
        self.client.force_login(self.admin)
        self.client.post(reverse("relatorio_bdo"), {"tipo": "RE", "mes_referencia": "2026-09", "arquivo": _planilha(LINHAS_3658)})
        self.assertEqual(BoletimBdo.objects.get().tipo, "RE")
        resp = self.client.get(reverse("relatorio_bdo"))
        self.assertContains(resp, ">RE</span>")

    def test_relatorio_bdo_filtra_os_boletins_por_tipo(self):
        self._importar(tipo="RI")
        self._importar([_linha("3700", "ITU", "SP", "5,00", "35099999")], tipo="RE")
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("relatorio_bdo"), {"tipo": "re"})
        self.assertEqual([boletim.numero for boletim in resp.context["boletins"]], ["3700"])
        resp = self.client.get(reverse("relatorio_bdo"), {"tipo": "XX"})
        self.assertEqual(len(resp.context["boletins"]), 2)
        BoletimBdo.objects.filter(tipo="RI").delete()
        resp = self.client.get(reverse("relatorio_bdo"), {"tipo": "RI"})
        self.assertContains(resp, "Nenhum Boletim RI importado.")

    def test_importacao_sem_tipo_ou_com_tipo_invalido_e_recusada(self):
        self.client.force_login(self.admin)
        for dados in ({"mes_referencia": "2026-09"}, {"tipo": "XX", "mes_referencia": "2026-09"}):
            with self.subTest(dados=dados):
                resp = self.client.post(
                    reverse("relatorio_bdo"), {**dados, "arquivo": _planilha(LINHAS_3658)}, follow=True,
                )
                self.assertContains(resp, "Não foi possível importar")
        self.assertFalse(BoletimBdo.objects.exists())

    def test_tipo_aparece_no_filtro_de_boletim_e_na_coluna_sem_filtro_proprio(self):
        self._importar(tipo="RI")
        self._importar([_linha("3700", "ITU", "SP", "5,00", "35099999")], tipo="RE")
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"agrupar": "uf"})
        self.assertContains(resp, "3700 — RE")
        self.assertContains(resp, "3658 — RI")
        self.assertEqual({linha["tipo"] for linha in resp.context["page_obj"]}, {"RE", "RI"})
        self.assertNotContains(resp, 'name="tipo"')

    def test_importacao_invalida_mostra_erro_e_nao_grava(self):
        self.client.force_login(self.admin)
        planilha = _planilha([_linha("3658", "ATIBAIA", "SP", "277,89", "")])
        resp = self.client.post(reverse("relatorio_bdo"), {"tipo": "RI", "mes_referencia": "2026-09", "arquivo": planilha}, follow=True)
        self.assertContains(resp, "Não foi possível importar")
        self.assertFalse(BoletimBdo.objects.exists())

    def test_reenviar_boletim_ja_importado_mostra_erro_pela_tela(self):
        self._importar()
        self.client.force_login(self.admin)
        resp = self.client.post(reverse("relatorio_bdo"), {"tipo": "RI", "mes_referencia": "2026-09", "arquivo": _planilha(LINHAS_3658[:1])}, follow=True)
        self.assertContains(resp, "Boletim de desempenho já importado: 3658")
        self.assertEqual(InepBdo.objects.count(), 8)

    def test_analista_ve_o_grid_mas_nao_acessa_o_relatorio_bdo(self):
        self.client.force_login(self.analista)
        self.assertEqual(self.client.get(reverse("relatorio_bdo")).status_code, 403)
        resp = self.client.post(reverse("relatorio_bdo"), {"tipo": "RI", "mes_referencia": "2026-09", "arquivo": _planilha(LINHAS_3658)})
        self.assertEqual(resp.status_code, 403)
        self.assertFalse(BoletimBdo.objects.exists())
        self.assertNotContains(self.client.get(reverse("bdo")), f'href="{reverse("relatorio_bdo")}"')

    def test_visualizador_e_bloqueado_e_nao_ve_o_link(self):
        visualizador = User.objects.create_user(
            username="bdo-visualizador", password="senha-teste-123", perfil=User.PERFIL_VISUALIZADOR,
        )
        self.client.force_login(visualizador)
        self.assertRedirects(self.client.get(reverse("bdo")), reverse("grid_inep"))
        resp = self.client.post(reverse("relatorio_bdo"), {"tipo": "RI", "mes_referencia": "2026-09", "arquivo": _planilha(LINHAS_3658)})
        self.assertRedirects(resp, reverse("grid_inep"))
        self.assertFalse(BoletimBdo.objects.exists())
        self.assertNotContains(self.client.get(reverse("grid_inep")), f'href="{reverse("bdo")}"')


class RecebiveisBdoTests(TestCase):
    """Modal "Recebíveis" da tela BDO (pedido do usuário, 2026-10-14): só
    montagem, nada é gravado; retirar é por município, nunca por INEP."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="receb-admin", password="senha-teste-123", perfil=User.PERFIL_ADMINISTRADOR,
        )
        self.analista = User.objects.create_user(
            username="receb-analista", password="senha-teste-123", perfil=User.PERFIL_ANALISTA,
        )
        importar_planilha_bdo(ler_planilha_bdo(_planilha(LINHAS_3658)), "RI", "FR_3658.xlsx", self.admin)
        # Campo Grande repetido no Boletim 3700: já vem rateado pelo cliente.
        importar_planilha_bdo(ler_planilha_bdo(_planilha([
            _linha("3700", "Campo Grande", "MS", "100,00", "50031287, 50005111"),
            _linha("3700", "CAMPO GRANDE", "MS", "50,00", "50033433"),
            _linha("3700", "ITU", "SP", "5,00", "35099999"),
        ])), "RE", "FR_3700.xlsx", self.admin)
        self.client.force_login(self.analista)

    def _post(self, **dados):
        return self.client.post(reverse("bdo_recebiveis"), dados)

    def _id(self, boletim, municipio, valor_total):
        return MunicipioBdo.objects.get(
            boletim__numero=boletim, municipio=municipio, valor_total=Decimal(valor_total),
        ).pk

    def test_tela_bdo_tem_o_botao_e_o_modal_sem_boletim_escolhido(self):
        resp = self.client.get(reverse("bdo"))
        self.assertContains(resp, 'data-abrir-modal-email="modal-recebiveis"')
        self.assertContains(resp, "Escolha um ou mais Boletins")
        self.assertEqual(resp.context["recebiveis_valor_total"], Decimal("0"))

    def test_escolher_mais_de_um_boletim_traz_todas_as_linhas(self):
        resp = self._post(boletim=["3658", "3700"], visao="linhas")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(resp.context["recebiveis_linhas"]), 7)
        self.assertEqual(resp.context["recebiveis_quantidade_ineps"], 12)
        self.assertEqual(resp.context["recebiveis_valor_total"], Decimal("32760.94"))
        self.assertContains(resp, "50005111, 50028340, 50031287, 50033433")

    def test_retirar_municipio_tira_o_valor_e_os_ineps_e_voltar_devolve(self):
        brasilia = self._id("3658", "Brasília", "31183.68")
        resp = self._post(boletim="3658", retirar=str(brasilia))
        self.assertEqual(resp.context["recebiveis_valor_total"], Decimal("1422.26"))
        self.assertEqual(resp.context["recebiveis_quantidade_ineps"], 6)
        self.assertContains(resp, f'name="retirado" value="{brasilia}"')
        # O estado volta no formulário: o retirado continua fora na próxima mudança.
        resp = self._post(boletim="3658", retirado=str(brasilia), visao="boletim")
        self.assertEqual(resp.context["recebiveis_valor_total"], Decimal("1422.26"))
        resp = self._post(boletim="3658", retirado=str(brasilia), restaurar=str(brasilia))
        self.assertEqual(resp.context["recebiveis_valor_total"], Decimal("32605.94"))
        resp = self._post(boletim="3658", retirado=str(brasilia), restaurar="todos")
        self.assertEqual(resp.context["recebiveis_retirados"], [])

    def test_linha_repetida_do_municipio_e_independente(self):
        rateio = self._id("3700", "CAMPO GRANDE", "50.00")
        resp = self._post(boletim="3700", retirar=str(rateio))
        self.assertEqual(resp.context["recebiveis_valor_total"], Decimal("105.00"))
        self.assertEqual(resp.context["recebiveis_quantidade_ineps"], 3)

    def test_agrupar_por_municipio_soma_as_linhas_repetidas_dentro_de_cada_boletim(self):
        resp = self._post(boletim=["3658", "3700"], visao="municipio")
        grupos = {(g["boletim"].numero, g["municipio"]): g for g in resp.context["recebiveis_grupos"]}
        self.assertEqual(grupos[("3700", "CAMPO GRANDE")]["valor_total"], Decimal("150.00"))
        self.assertEqual(grupos[("3700", "CAMPO GRANDE")]["quantidade_ineps"], 3)
        self.assertEqual(grupos[("3658", "CAMPO GRANDE")]["valor_total"], Decimal("1094.45"))
        # "Retirar" no agrupado tira todas as linhas do município naquele Boletim.
        resp = self._post(boletim="3700", visao="municipio", retirar=grupos[("3700", "CAMPO GRANDE")]["ids"])
        self.assertEqual(resp.context["recebiveis_valor_total"], Decimal("5.00"))
        self.assertEqual(len(resp.context["recebiveis_retirados"]), 2)

    def test_por_boletim_mostra_o_total_de_cada_boletim(self):
        resp = self._post(boletim=["3658", "3700"], visao="boletim")
        totais = {g["boletim"].numero: (g["valor_total"], g["quantidade_municipios"]) for g in resp.context["recebiveis_grupos"]}
        self.assertEqual(totais, {"3658": (Decimal("32605.94"), 4), "3700": (Decimal("155.00"), 2)})

    def test_nao_tem_retirar_por_inep_nem_adicionar_inep(self):
        resp = self._post(boletim="3658")
        self.assertNotContains(resp, 'name="inep"')
        self.assertNotContains(resp, 'name="retirar_inep"')

    def test_retirado_de_boletim_desmarcado_e_id_invalido_sao_ignorados(self):
        itu = self._id("3700", "ITU", "5.00")
        resp = self._post(boletim="3658", retirado=[str(itu), "abc", "999999"], visao="xx")
        self.assertEqual(resp.context["recebiveis_retirados"], [])
        self.assertEqual(resp.context["recebiveis_visao"], "municipio")  # padrão: Por Município (pedido do usuário, 2026-10-15)
        self.assertEqual(resp.context["recebiveis_valor_total"], Decimal("32605.94"))

    def test_abre_com_o_boletim_e_o_periodo_do_grid_por_municipio(self):
        """Pedido do usuário (2026-10-15): só o Boletim e o período de Data de
        Ativação filtrados no grid; o município entra com a parte dos INEPs
        do período e município sem nenhum fica de fora."""
        escola = Escola.objects.create(inep="50031287", nome="Campo Grande", estado="MS", municipio="Campo Grande")
        Ri.objects.create(escola=escola, data_ativacao=datetime.date(2026, 9, 10))
        resp = self.client.get(reverse("bdo"), {"boletim": "3658", "data_inicial": "01/09/2026", "data_final": "30/09/2026"})
        self.assertEqual(resp.context["recebiveis_selecionados"], {"3658"})
        self.assertEqual(resp.context["recebiveis_visao"], "municipio")
        [grupo] = resp.context["recebiveis_grupos"]
        self.assertEqual((grupo["municipio"], grupo["quantidade_ineps"]), ("CAMPO GRANDE", 1))
        self.assertEqual(resp.context["recebiveis_valor_total"], Decimal("273.61"))  # 1.094,45 ÷ 4 INEPs
        self.assertContains(resp, '<input type="hidden" name="data_inicial" value="01/09/2026">', html=False)
        self.assertContains(resp, "Ativação de 01/09/2026 até 30/09/2026")
        # o período segue no formulário do modal
        resp = self._post(boletim="3658", data_inicial="01/09/2026", data_final="30/09/2026")
        self.assertEqual(resp.context["recebiveis_valor_total"], Decimal("273.61"))

    def test_sem_filtro_no_grid_abre_sem_boletim(self):
        resp = self.client.get(reverse("bdo"))
        self.assertEqual(resp.context["recebiveis_selecionados"], set())

    def test_do_inep_contestado_entra_so_o_recebido_e_o_municipio_e_rateado(self):
        # Manutenção 300,00 no Lote 9, desconto 200,00: recebeu 100,00.
        KitPadrao.objects.create(
            descricao="Manutenção de Rede Interna (serviços, materiais e equipamentos)", lote=9,
            unidade="Escola/Mês", valor_servico=Decimal("300.00"),
        )
        Escola.objects.create(inep="50031287", nome="Escola 50031287", estado="MS", municipio="X", lote=9)
        InepBdo.objects.filter(inep="50031287", municipio_bdo__boletim__numero="3658").update(
            status_divergencia=InepBdo.STATUS_CONTESTADO, valor_desconto_cliente=Decimal("200.00"),
        )
        InepBdo.objects.filter(inep="35099999").update(status_divergencia=InepBdo.STATUS_CONTESTADO)
        InepBdo.objects.filter(inep="50005111").update(status_divergencia=InepBdo.STATUS_APROVADO)
        resp = self._post(boletim=["3658", "3700"], visao="linhas")
        linhas = {(linha.boletim.numero, linha.municipio): linha for linha in resp.context["recebiveis_linhas"]}
        campo_grande = linhas[("3658", "Campo Grande")]
        self.assertEqual(campo_grande.lista_ineps, ["50005111", "50028340", "50033433", "50031287"])
        self.assertEqual(campo_grande.valor_recebivel, Decimal("920.84"))  # 1.094,45 × 3 ÷ 4 + 100,00 recebidos
        self.assertEqual(linhas[("3700", "ITU")].valor_recebivel, Decimal("0.00"))  # contestado sem valor recebido
        self.assertEqual(resp.context["recebiveis_quantidade_ineps"], 12)
        self.assertEqual(resp.context["recebiveis_valor_total"], Decimal("32582.33"))
        self.assertEqual(resp.context["recebiveis_contestados"], 2)
        self.assertContains(resp, "Sem valor EACE")  # sem leitura EACE: rateio só para conferência

    def test_nada_e_gravado(self):
        antes = (BoletimBdo.objects.count(), MunicipioBdo.objects.count(), InepBdo.objects.count())
        self._post(boletim=["3658", "3700"], retirar=str(self._id("3700", "ITU", "5.00")))
        self.assertEqual((BoletimBdo.objects.count(), MunicipioBdo.objects.count(), InepBdo.objects.count()), antes)

    def test_anonimo_vai_pro_login_e_visualizador_e_bloqueado(self):
        self.client.logout()
        resp = self._post(boletim="3658")
        self.assertIn(reverse("login"), resp["Location"])
        visualizador = User.objects.create_user(
            username="receb-visualizador", password="senha-teste-123", perfil=User.PERFIL_VISUALIZADOR,
        )
        self.client.force_login(visualizador)
        self.assertRedirects(self._post(boletim="3658"), reverse("grid_inep"))


class ImportarAssinaturaReTests(TestCase):
    def _planilha(self, linhas):
        workbook = openpyxl.Workbook()
        aba = workbook.active
        aba.title = "FATURAMENTO BDO"
        aba.append(("EACE - APRENDER CONECTADO",))
        aba.append(("LOTE", "UF ", "MUNICIPIO", "INEP", "ASSINATURA MENSAL R.E", "MANUTENÇÃO R.I MENSAL"))
        for linha in linhas:
            aba.append(linha)
        pasta = tempfile.mkdtemp()
        caminho = Path(pasta) / "CONSOLIDADO EACE Atualizado.xlsx"
        workbook.save(caminho)
        return str(caminho)

    def test_grava_a_assinatura_e_substitui_a_anterior(self):
        americana = Escola.objects.create(inep="35058592", nome="Americana", valor_assinatura_re=Decimal("1.00"))
        atibaia = Escola.objects.create(inep="35011006", nome="Atibaia")
        saida = io.StringIO()
        call_command("importar_assinatura_re", self._planilha([
            (9, "SP", "Americana", 35058592, 760, 380.38),
            (11, "TO", "Atibaia", "35011006", "370,50", 365.62),
            (9, "SP", "Lins", 35005434, 360, 380.38),
            (9, "SP", "Itu", 35099999, None, 380.38),
        ]), stdout=saida)
        americana.refresh_from_db()
        atibaia.refresh_from_db()
        self.assertEqual(americana.valor_assinatura_re, Decimal("760.00"))
        self.assertEqual(atibaia.valor_assinatura_re, Decimal("370.50"))
        self.assertFalse(Escola.objects.filter(inep="35005434").exists())
        self.assertIn("2 escola(s) atualizada(s), 1 INEP(s) sem escola no sistema, 1 linha(s)", saida.getvalue())


def _pro_rata(linhas, cabecalho=("INEP", "Ativado em", "Dias Descontados", "Valor"), nome="Pro_rata3658.xlsx"):
    return _planilha(linhas, cabecalho, nome)


def _item(ativado_em, dias, valor):
    return {"ativado_em": ativado_em, "dias_descontados": dias, "valor_desconto": Decimal(valor)}


class ProRataBdoTests(TestCase):
    """Planilha Pro rata do cliente (pedido do usuário, 2026-10-15): o
    "Ativado em", os "Dias Descontados" e o "Valor" (o desconto) de cada
    INEP do Boletim, para conferir com a Data de Ativação e, pelo valor pago
    (Valor assinatura − desconto), com o Valor proporcional do sistema, e a
    aprovação das divergências com log."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="pro-rata-admin", password="senha-teste-123", perfil=User.PERFIL_ADMINISTRADOR,
        )
        self.analista = User.objects.create_user(
            username="pro-rata-analista", password="senha-teste-123", perfil=User.PERFIL_ANALISTA,
        )
        [self.boletim] = importar_planilha_bdo(
            ler_planilha_bdo(_planilha(LINHAS_3658)), "RI", "FR_3658.xlsx", self.admin, datetime.date(2026, 9, 1),
        )
        KitPadrao.objects.create(
            descricao="Manutenção de Rede Interna (serviços, materiais e equipamentos)", lote=9,
            unidade="Escola/Mês", valor_servico=Decimal("300.00"),
        )
        # 35058592: 21/09 -> 9 dias, proporcional 90,00.
        # 35011006: 05/09 -> 25 dias, proporcional 250,00.
        for inep, data in (("35058592", datetime.date(2026, 9, 21)), ("35011006", datetime.date(2026, 9, 5))):
            escola = Escola.objects.create(inep=inep, nome=f"Escola {inep}", estado="SP", municipio="X", lote=9)
            Ri.objects.create(escola=escola, data_ativacao=data)

    def _inep(self, inep):
        return InepBdo.objects.get(municipio_bdo__boletim=self.boletim, inep=inep)

    def _enviar(self, arquivo, usuario=None):
        self.client.force_login(usuario or self.admin)
        return self.client.post(reverse("bdo_pro_rata", args=[self.boletim.pk]), {"arquivo": arquivo}, follow=True)

    def _linhas(self):
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"))
        return resp, {linha.inep: linha for linha in resp.context["page_obj"]}

    def _status(self, inep, status, usuario=None, **extra):
        self.client.force_login(usuario or self.analista)
        return self.client.post(
            reverse("bdo_status_divergencia", args=[self._inep(inep).pk]), {"status": status}, **extra,
        )

    def test_le_inep_numerico_data_em_texto_ou_do_excel_dias_e_valor(self):
        dados = ler_planilha_pro_rata(_pro_rata([
            (35058592, "21/09/2026", 21, 322), ("35011006", datetime.datetime(2026, 9, 8), 8, 444.666666666667),
        ]))
        self.assertEqual(dados, {
            "35058592": _item(datetime.date(2026, 9, 21), 21, "322.00"),
            "35011006": _item(datetime.date(2026, 9, 8), 8, "444.67"),
        })

    def test_sem_coluna_dado_invalido_ou_inep_repetido_com_outros_dados_da_erro(self):
        with self.assertRaisesMessage(PlanilhaBdoError, "INEP, Ativado em, Dias Descontados, Valor"):
            ler_planilha_pro_rata(_pro_rata([(35058592, "21/09/2026")], cabecalho=("INEP", "Ativado em")))
        with self.assertRaisesMessage(PlanilhaBdoError, "linha 2 (Ativado em, Valor)"):
            ler_planilha_pro_rata(_pro_rata([(35058592, "31/31/2026", 0, "")]))
        with self.assertRaisesMessage(PlanilhaBdoError, "linha 2 (Dias Descontados)"):
            ler_planilha_pro_rata(_pro_rata([(35058592, "21/09/2026", "x", 10)]))
        with self.assertRaisesMessage(PlanilhaBdoError, "linha 3 (INEP 35058592 repetido com outros dados)"):
            ler_planilha_pro_rata(_pro_rata([(35058592, "21/09/2026", 21, 10), (35058592, "21/09/2026", 21, 11)]))

    def test_grava_so_nos_ineps_do_boletim_e_reenviar_substitui(self):
        outro = importar_planilha_bdo(
            ler_planilha_bdo(_planilha([_linha("3700", "AMERICANA", "SP", "10,00", "35058592")])),
            "RE", "FR_3700.xlsx", self.admin,
        )[0]
        atualizados, fora = importar_pro_rata(self.boletim, {
            "35058592": _item(datetime.date(2026, 9, 21), 21, "210"),
            "35011006": _item(datetime.date(2026, 9, 8), 8, "50"),
            "99999999": _item(datetime.date(2026, 9, 1), 1, "10"),
        }, "Pro_rata3658.xlsx")
        self.assertEqual((atualizados, fora), (2, ["99999999"]))
        inep = self._inep("35058592")
        self.assertEqual(
            (inep.ativado_em_cliente, inep.dias_descontados_cliente, inep.valor_desconto_cliente),
            (datetime.date(2026, 9, 21), 21, Decimal("210.00")),
        )
        self.assertIsNone(InepBdo.objects.get(municipio_bdo__boletim=outro).ativado_em_cliente)
        importar_pro_rata(self.boletim, {"35058592": _item(datetime.date(2026, 9, 22), 22, "220")}, "Pro_rata3658b.xlsx")
        self.assertEqual(self._inep("35058592").ativado_em_cliente, datetime.date(2026, 9, 22))
        self.assertIsNone(self._inep("35011006").valor_desconto_cliente)
        self.boletim.refresh_from_db()
        self.assertEqual(self.boletim.pro_rata_arquivo, "Pro_rata3658b.xlsx")

    def test_administrador_envia_pela_tela_e_ve_os_ineps_fora_do_boletim(self):
        resp = self._enviar(_pro_rata([(35058592, "21/09/2026", 21, 322), (99999999, "01/09/2026", 1, 10)]))
        self.assertRedirects(resp, reverse("relatorio_bdo"))
        self.assertContains(resp, "Boletim 3658: Pro rata importada, 1 INEP(s) com Ativado em.")
        self.assertContains(resp, "1 INEP(s) da Pro rata não estão no Boletim: 99999999.")
        self.assertContains(resp, "Pro_rata3658.xlsx")
        self.assertEqual(self._inep("35058592").valor_desconto_cliente, Decimal("322.00"))

    def test_planilha_invalida_mostra_erro_e_nao_grava(self):
        resp = self._enviar(_pro_rata([(35058592, "", 0, 0)]))
        self.assertContains(resp, "Boletim 3658: Pro rata não importada.")
        self.assertFalse(InepBdo.objects.filter(ativado_em_cliente__isnull=False).exists())
        self.assertEqual(BoletimBdo.objects.get().pro_rata_arquivo, "")

    def test_so_administrador_envia_a_pro_rata(self):
        resp = self._enviar(_pro_rata([(35058592, "21/09/2026", 21, 322)]), usuario=self.analista)
        self.assertEqual(resp.status_code, 403)
        self.assertFalse(InepBdo.objects.filter(ativado_em_cliente__isnull=False).exists())

    def test_grid_mostra_divergencia_de_data_e_de_valor_pago(self):
        importar_pro_rata(self.boletim, {
            "35058592": _item(datetime.date(2026, 9, 21), 21, "210.01"),  # tudo confere (R$ 0,01 de tolerância)
            "35011006": _item(datetime.date(2026, 9, 8), 22, "65"),  # data e valor pago diferentes
            "53002814": _item(datetime.date(2026, 9, 1), 29, "10"),  # sem escola: sem Data de Ativação
        }, "Pro_rata3658.xlsx")
        resp, linhas = self._linhas()
        americana = linhas["35058592"]
        self.assertEqual((americana.ativacao_confere, americana.valor_confere, americana.tem_divergencia), (True, True, False))
        atibaia = linhas["35011006"]
        self.assertEqual(atibaia.valor_proporcional, Decimal("250.00"))
        self.assertEqual((atibaia.ativacao_confere, atibaia.valor_confere, atibaia.tem_divergencia), (False, False, True))
        sem_data = linhas["53002814"]
        self.assertEqual((sem_data.ativacao_confere, sem_data.tem_divergencia), (None, True))
        self.assertContains(resp, "Ativado em (cliente)")
        self.assertContains(resp, "Desconto (cliente)")
        self.assertContains(resp, "Valor pago (cliente)")
        self.assertContains(resp, "Valor faltando")
        self.assertContains(resp, "(diferente da Data de ativação)", count=1)
        self.assertContains(resp, "(menor que o Valor proporcional do sistema, R$ 250,00)", count=1)
        # campo de status e caixa da seleção em massa só no INEP divergente
        # ou sem Data de Ativação: todos menos 35058592, que confere
        self.assertContains(resp, "Status da divergência do INEP", count=7)
        self.assertContains(resp, "Status da divergência do INEP 35011006")
        self.assertContains(resp, "Status da divergência do INEP 53002814")
        self.assertNotContains(resp, "Status da divergência do INEP 35058592")
        self.assertContains(resp, 'name="inep_bdo"', count=7)
        self.assertContains(resp, 'id="form-status-massa"')

    def test_sem_ativado_em_do_cliente_tambem_recebe_status(self):
        _resp, linhas = self._linhas()  # sem Pro rata importada
        self.assertIsNotNone(linhas["35058592"].data_ativacao)
        self.assertTrue(linhas["35058592"].tem_divergencia)
        self._status("35058592", "APROVADO")
        self.assertEqual(self._inep("35058592").status_divergencia, "APROVADO")

    def test_filtro_de_status_mostra_os_ineps_com_o_campo(self):
        importar_pro_rata(self.boletim, {
            "35058592": _item(datetime.date(2026, 9, 21), 21, "210"),  # confere: sem campo de status
            "35011006": _item(datetime.date(2026, 9, 8), 22, "65"),
        }, "P.xlsx")
        self._status("35011006", "CONTESTADO")
        self._status("53002814", "APROVADO")  # sem Data de Ativação

        def ineps(status):
            resp = self.client.get(reverse("bdo"), {"status": status})
            return {linha.inep for linha in resp.context["page_obj"]}

        todos = {"35058592", "35011006", "50031287", "50005111", "50033433", "50028340", "53002814", "53002474"}
        self.assertEqual(ineps("COM_STATUS"), todos - {"35058592"})
        self.assertEqual(ineps("CONTESTADO"), {"35011006"})
        self.assertEqual(ineps("APROVADO"), {"53002814"})
        self.assertEqual(ineps("PENDENTE"), todos - {"35058592", "35011006", "53002814"})
        self.assertEqual(ineps("QUALQUER"), todos)
        resp = self.client.get(reverse("bdo"), {"status": "CONTESTADO", "agrupar": "municipio"})
        self.assertEqual([linha["municipio"] for linha in resp.context["page_obj"]], ["ATIBAIA"])
        self.assertContains(resp, 'value="CONTESTADO" selected')

    def test_dif_do_total_mega_leva_aos_ineps_que_formam_a_diferenca(self):
        importar_pro_rata(self.boletim, {
            "35058592": _item(datetime.date(2026, 9, 21), 21, "210"),  # confere: fora da lista
            "35011006": _item(datetime.date(2026, 9, 5), 4, "65"),  # cliente paga 235,00, 15,00 a menos
        }, "P.xlsx")
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"agrupar": "municipio"})
        self.assertContains(resp, "&agrupar=inep&diferenca=1")

        resp = self.client.get(reverse("bdo"), {"diferenca": "1"})
        linhas = {linha.inep: linha for linha in resp.context["page_obj"]}
        todos = {"35058592", "35011006", "50031287", "50005111", "50033433", "50028340", "53002814", "53002474"}
        self.assertEqual(set(linhas), todos - {"35058592"})  # os sem escola ficam sem valor no sistema
        self.assertEqual(linhas["35011006"].diferenca_inep, Decimal("15.00"))
        self.assertIsNone(linhas["53002814"].diferenca_inep)
        self.assertEqual(resp.context["diferenca"]["soma"], Decimal("15.00"))
        self.assertEqual(resp.context["diferenca"]["sem_valor"], 6)
        self.assertFalse(resp.context["diferenca"]["sem_pro_rata"])
        self.assertContains(resp, "Cliente pagou R$ 15,00 a menos")
        self.assertContains(resp, 'name="diferenca" value="1"')
        self.assertNotContains(resp, "Boletim sem Pro rata do cliente importada")

        resp = self.client.get(reverse("bdo"), {"diferenca": "1", "agrupar": "municipio"})
        self.assertEqual(sum(linha["quantidade_ineps"] for linha in resp.context["page_obj"]), 7)

        # Cliente pagou a menos em 35011006 (15,00 a menos); os sem
        # valor no sistema a planilha cobra: cliente pagou a mais.
        resp = self.client.get(reverse("bdo"), {"diferenca": "menos"})
        self.assertEqual({linha.inep for linha in resp.context["page_obj"]}, {"35011006"})
        self.assertContains(resp, "Cliente pagou R$ 15,00 a menos")
        self.assertEqual(resp.context["diferenca"]["soma_lado"], Decimal("15.00"))
        resp = self.client.get(reverse("bdo"), {"diferenca": "mais"})
        self.assertEqual({linha.inep for linha in resp.context["page_obj"]}, todos - {"35058592", "35011006"})
        self.assertContains(resp, 'name="diferenca" value="mais"')
        self.assertContains(resp, "Total pago a mais")
        self.assertEqual(
            {linha.inep for linha in self.client.get(reverse("bdo"), {"diferenca": "x"}).context["page_obj"]}, todos,
        )

    def test_dif_sem_valor_do_cliente_usa_o_ativado_em_dele_e_nao_o_valor_cheio(self):
        importar_pro_rata(self.boletim, {
            "35058592": {**_item(datetime.date(2026, 9, 21), 21, "0"), "valor_desconto": None},  # mesma data
            "35011006": {**_item(datetime.date(2026, 9, 8), 8, "0"), "valor_desconto": None},  # 3 dias depois
        }, "P.xlsx")
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"diferenca": "1", "inep": "350"})
        linhas = {linha.inep: linha for linha in resp.context["page_obj"]}
        self.assertEqual(set(linhas), {"35011006"})
        self.assertEqual(linhas["35011006"].diferenca_inep, Decimal("30.00"))  # 250,00 do sistema - 220,00 pagos

    def test_dif_sem_pro_rata_conta_o_mes_cheio_como_pago(self):
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("bdo"), {"diferenca": "1", "inep": "350"})
        self.assertEqual({linha.inep for linha in resp.context["page_obj"]}, {"35058592", "35011006"})
        self.assertEqual(resp.context["diferenca"]["soma"], Decimal("-260.00"))
        self.assertContains(resp, "Boletim sem Pro rata do cliente importada")

    def test_so_valor_pago_diferente_tambem_e_divergencia(self):
        importar_pro_rata(self.boletim, {"35058592": _item(datetime.date(2026, 9, 21), 21, "215")}, "P.xlsx")
        _resp, linhas = self._linhas()
        self.assertTrue(linhas["35058592"].ativacao_confere)
        self.assertTrue(linhas["35058592"].tem_divergencia)

    def test_trocar_status_grava_log_e_devolve_o_campo(self):
        importar_pro_rata(self.boletim, {"35011006": _item(datetime.date(2026, 9, 8), 22, "65")}, "P.xlsx")
        resp = self._status("35011006", "APROVADO", HTTP_HX_REQUEST="true")
        self.assertContains(resp, '<option value="APROVADO" selected>Aprovado</option>', html=False)
        self.assertContains(resp, 'id="valor-faltando-')
        inep = self._inep("35011006")
        self.assertEqual((inep.status_divergencia, inep.status_divergencia_por), ("APROVADO", self.analista))
        self.assertIsNotNone(inep.status_divergencia_em)
        log = LogStatusDivergenciaBdo.objects.get()
        self.assertEqual(
            (log.boletim, log.inep, log.acao, log.status_anterior, log.status, log.em_massa, log.usuario,
             log.data_ativacao, log.ativado_em_cliente, log.valor_desconto_sistema, log.valor_desconto_cliente),
            ("3658", "35011006", "ALTERADO", "", "APROVADO", False, self.analista, datetime.date(2026, 9, 5),
             datetime.date(2026, 9, 8), Decimal("50.00"), Decimal("65.00")),
        )
        resp, _linhas = self._linhas()
        self.assertContains(resp, "Alterado por pro-rata-analista")

        resp = self._status("35011006", "CONTESTADO", usuario=self.admin)
        self.assertRedirects(resp, f"{reverse('bdo')}?boletim=3658&inep=35011006")
        self._status("35011006", "CONTESTADO")  # mesmo status: sem log novo
        self._status("35011006", "")  # pendente não é opção do usuário
        self._status("35011006", "REPROVADO")  # opção que não existe mais
        self._status("35011006", "APROVADO")
        self.assertEqual(self._inep("35011006").status_divergencia, "APROVADO")
        self.assertEqual(
            list(LogStatusDivergenciaBdo.objects.order_by("pk").values_list("status_anterior", "status", "usuario__username")),
            [("", "APROVADO", "pro-rata-analista"), ("APROVADO", "CONTESTADO", "pro-rata-admin"),
             ("CONTESTADO", "APROVADO", "pro-rata-analista")],
        )

    def test_contestar_devolve_o_valor_faltando_e_o_recebido_fica(self):
        importar_pro_rata(self.boletim, {"35011006": _item(datetime.date(2026, 9, 5), 25, "65")}, "P.xlsx")
        resp = self._status("35011006", "APROVADO", HTTP_HX_REQUEST="true")
        self.assertContains(resp, 'hx-swap-oob="true"')
        self.assertNotContains(resp, "R$ 15,00")
        resp = self._status("35011006", "CONTESTADO", HTTP_HX_REQUEST="true")
        self.assertContains(resp, "R$ 15,00")  # faltando: 250,00 − 235,00 recebidos
        self.assertContains(resp, "Recebido R$ 235,00 (fica)")
        self.assertEqual(self._inep("35011006").valor_desconto_cliente, Decimal("65.00"))

    def test_cliente_pagou_mais_que_o_proporcional_confere_e_nao_tem_faltando(self):
        importar_pro_rata(self.boletim, {
            "35058592": _item(datetime.date(2026, 9, 21), 21, "180"),  # pagou 120,00, 30,00 a mais: confere
            "35011006": _item(datetime.date(2026, 9, 8), 22, "40"),  # data diferente, pagou 260,00, 10,00 a mais
        }, "P.xlsx")
        _resp, linhas = self._linhas()
        self.assertEqual((linhas["35058592"].valor_confere, linhas["35058592"].tem_divergencia), (True, False))
        self.assertEqual((linhas["35011006"].valor_confere, linhas["35011006"].tem_divergencia), (True, True))
        resp = self._status("35011006", "CONTESTADO", HTTP_HX_REQUEST="true")
        self.assertNotContains(resp, "Recebido R$")
        linha = self.client.get(reverse("bdo"), {"agrupar": "municipio", "inep": "35011006"}).context["page_obj"][0]
        self.assertEqual((linha["total_mega"], linha["total_contestado_mega"]), (Decimal("250.00"), Decimal("0")))

    def test_status_invalido_ou_inep_sem_divergencia_nao_grava(self):
        importar_pro_rata(self.boletim, {
            "35058592": _item(datetime.date(2026, 9, 21), 21, "210"),
            "35011006": _item(datetime.date(2026, 9, 8), 22, "65"),
        }, "P.xlsx")
        self._status("35058592", "APROVADO")
        self._status("35011006", "QUALQUER")
        self.assertFalse(InepBdo.objects.exclude(status_divergencia="").exists())
        self.assertFalse(LogStatusDivergenciaBdo.objects.exists())

    def test_status_em_massa_so_nos_ineps_divergentes_e_volta_para_a_mesma_pagina(self):
        importar_pro_rata(self.boletim, {
            "35058592": _item(datetime.date(2026, 9, 21), 21, "210"),  # sem divergência
            "35011006": _item(datetime.date(2026, 9, 8), 22, "65"),  # data e valor pago diferentes
            "53002814": _item(datetime.date(2026, 9, 1), 29, "10"),  # sem Data de Ativação no sistema
        }, "P.xlsx")
        ids = [self._inep(inep).pk for inep in ("35058592", "35011006", "53002814")]
        self.client.force_login(self.analista)
        voltar = f"{reverse('bdo')}?boletim=3658&page=1"
        resp = self.client.post(
            reverse("bdo_status_divergencia_massa"),
            {"inep_bdo": ids, "status": "CONTESTADO", "next": voltar}, follow=True,
        )
        self.assertRedirects(resp, voltar)
        self.assertContains(resp, "Status &quot;Contestado&quot; aplicado em 2 INEP(s).")
        self.assertContains(resp, "1 INEP(s) sem divergência ficaram de fora.")
        self.assertEqual(
            dict(InepBdo.objects.filter(pk__in=ids).values_list("inep", "status_divergencia")),
            {"35058592": "", "35011006": "CONTESTADO", "53002814": "CONTESTADO"},
        )
        self.assertEqual(LogStatusDivergenciaBdo.objects.filter(em_massa=True, status="CONTESTADO").count(), 2)

    def test_status_em_massa_sem_inep_marcado_ou_destino_externo(self):
        self.client.force_login(self.analista)
        resp = self.client.post(
            reverse("bdo_status_divergencia_massa"), {"status": "APROVADO", "next": "https://exemplo.com/"}, follow=True,
        )
        self.assertRedirects(resp, reverse("bdo"))
        self.assertContains(resp, "Marque pelo menos um INEP.")

    def test_reenviar_pro_rata_com_outro_dado_volta_o_status_para_pendente_com_log(self):
        item = _item(datetime.date(2026, 9, 8), 22, "65")
        importar_pro_rata(self.boletim, {"35011006": item}, "P.xlsx")
        self._status("35011006", "CONTESTADO")
        importar_pro_rata(self.boletim, {"35011006": item}, "P2.xlsx", self.admin)
        self.assertEqual(self._inep("35011006").status_divergencia, "CONTESTADO")  # mesmo dado: continua
        importar_pro_rata(self.boletim, {"35011006": _item(datetime.date(2026, 9, 9), 21, "65")}, "P3.xlsx", self.admin)
        self.assertEqual(self._inep("35011006").status_divergencia, "")
        log = LogStatusDivergenciaBdo.objects.order_by("-pk").first()
        self.assertEqual(
            (log.acao, log.status_anterior, log.status, log.usuario, log.ativado_em_cliente),
            ("RETIRADO_PRO_RATA", "CONTESTADO", "", self.admin, datetime.date(2026, 9, 8)),
        )

    def test_visualizador_e_anonimo_nao_alteram_status(self):
        importar_pro_rata(self.boletim, {"35011006": _item(datetime.date(2026, 9, 8), 22, "65")}, "P.xlsx")
        visualizador = User.objects.create_user(
            username="pro-rata-vis", password="senha-teste-123", perfil=User.PERFIL_VISUALIZADOR,
        )
        resp = self._status("35011006", "APROVADO", usuario=visualizador)
        self.assertNotEqual(resp.status_code, 200)
        resp = self.client.post(
            reverse("bdo_status_divergencia_massa"), {"inep_bdo": [self._inep("35011006").pk], "status": "APROVADO"},
        )
        self.assertNotEqual(resp.status_code, 200)
        self.client.logout()
        resp = self.client.post(reverse("bdo_status_divergencia", args=[self._inep("35011006").pk]), {"status": "APROVADO"})
        self.assertIn(reverse("login"), resp["Location"])
        self.assertEqual(self._inep("35011006").status_divergencia, "")
        self.assertFalse(LogStatusDivergenciaBdo.objects.exists())


class MedicaoBdoTests(TestCase):
    """Planilha de Medição gerada pelo modal Recebíveis (pedido do usuário,
    2026-10-15): uma linha por linha do Boletim que ficou no modal, 1º
    faturamento por INEP, número sequencial por tipo."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="medicao-admin", password="senha-teste-123", perfil=User.PERFIL_ADMINISTRADOR,
        )
        self.analista = User.objects.create_user(
            username="medicao-analista", password="senha-teste-123", perfil=User.PERFIL_ANALISTA,
        )
        importar_planilha_bdo(ler_planilha_bdo(_planilha([
            _linha("4000", "São Paulo", "SP", "1.000,00", "99000001, 99000002"),
            _linha("4000", "SÃO PAULO", "SP", "50,00", "99000003"),
            _linha("4000", "Itu", "SP", "30,00", "99000004"),
        ])), "RE", "FR_4000.xlsx", self.admin)
        # São Paulo inteiro no faturamento RE: Integral; Itu tem escola faltando.
        for inep, municipio in (("99000001", "São Paulo"), ("99000002", "São Paulo"), ("99000003", "São Paulo"),
                                ("99000004", "Itu"), ("99000005", "Itu")):
            Escola.objects.create(inep=inep, nome=f"Escola {inep}", estado="SP", municipio=municipio)
        PrimeiroFaturamentoInep.objects.create(tipo="RE", inep="99000002", data=datetime.date(2026, 6, 15))
        PrimeiroFaturamentoInep.objects.create(tipo="RE", inep="99000001", data=datetime.date(2026, 7, 9))
        self.client.force_login(self.analista)

    @staticmethod
    def _conferir_eace(boletins=None, valores=None):
        """Leitura EACE com Sucesso que confere (1 card por município), com
        o valor do card = Valor total da planilha, salvo `valores[municipio]`."""
        for boletim in boletins or BoletimBdo.objects.all():
            for municipio in boletim.municipios.all():
                municipio.valor_eace = (valores or {}).get(municipio.municipio, municipio.valor_total)
                municipio.status_eace = "Aprovado"
                municipio.save()
            LeituraEaceBdo.objects.create(
                boletim=boletim, status=LeituraEaceBdo.SUCESSO, cards_lidos=boletim.municipios.count(),
                municipios_atualizados=boletim.municipios.count(),
            )

    def _gerar(self, **dados):
        if not LeituraEaceBdo.objects.exists():
            self._conferir_eace()
        return self.client.post(reverse("bdo_medicao"), dados)

    def test_valor_da_medicao_e_o_do_card_do_eace(self):
        self._conferir_eace(valores={"Itu": Decimal("31.50")})
        _aba, linhas = self._linhas(self._gerar(boletim="4000").content)
        self.assertEqual(linhas[0][:4][3], 31.5)

    def test_sem_leitura_eace_nao_gera_a_medicao(self):
        resp = self.client.post(reverse("bdo_medicao"), {"boletim": "4000"})
        self.assertRedirects(resp, reverse("bdo"), fetch_redirect_response=False)
        self.assertFalse(MedicaoBdo.objects.exists())
        mensagens = [str(m) for m in resp.wsgi_request._messages]
        self.assertIn("Boletim 4000 ainda não lido do portal EACE.", mensagens[0])

    def test_quantidade_de_municipios_diferente_da_de_cards_nao_gera(self):
        self._conferir_eace()
        LeituraEaceBdo.objects.update(cards_lidos=2)
        resp = self.client.post(reverse("bdo_medicao"), {"boletim": "4000"})
        self.assertFalse(MedicaoBdo.objects.exists())
        mensagens = [str(m) for m in resp.wsgi_request._messages]
        self.assertIn("3 município(s) na planilha importada e 2 card(s) no portal EACE", mensagens[0])

    def _linhas(self, conteudo):
        aba = openpyxl.load_workbook(io.BytesIO(conteudo)).worksheets[0]
        return aba, [
            tuple(aba.cell(linha, coluna).value for coluna in range(4, 10))
            for linha in range(8, aba.max_row + 1) if aba.cell(linha, 4).value
        ]

    def test_gera_a_planilha_do_boletim_no_formato_da_medicao(self):
        resp = self._gerar(boletim="4000")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("filename*=UTF-8''MEDI%C3%87%C3%83O%2005%20-%20MANUTEN%C3%87%C3%83O%204000.xlsx", resp["Content-Disposition"])
        workbook = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertEqual(workbook.sheetnames, ["MEDIÇÃO 05", "NOTAS FISCAIS"])
        aba, linhas = self._linhas(resp.content)
        hoje = datetime.datetime.combine(datetime.date.today(), datetime.time())
        self.assertEqual(linhas, [
            ("SP", "ITU", hoje, 30, "PROPORCIONAL", "=G8"),
            ("SP", "SÃO PAULO", datetime.datetime(2026, 6, 15), 1000, "INTEGRAL", "=G9"),  # a mais antiga dos INEPs
            ("SP", "SÃO PAULO", hoje, 50, "INTEGRAL", "=G10"),
        ])
        self.assertEqual(aba["I11"].value, "=SUM(I8:I10)")
        self.assertEqual(aba["A15"].value, "OBS NA NF: CÓDIGO DE LIBERAÇÃO 4000, FATURAR EM NOTAS DISTINTAS POR MUNÍCIPIO")
        self.assertEqual(aba["A16"].value, "VENCIMENTO PARA 30 DIAS, INSERIR DADOS BANCÁRIOS")
        self.assertEqual(len(aba._images), 3)
        notas = workbook["NOTAS FISCAIS"]
        self.assertEqual((notas["J7"].value, notas["E8"].value, notas["J8"].value), ("NOTA FISCAL", "ITU", None))
        # INEP novo fica com a data de hoje nos meses seguintes
        self.assertEqual(
            PrimeiroFaturamentoInep.objects.get(tipo="RE", inep="99000003").data, datetime.date.today(),
        )
        self.assertEqual(MedicaoBdo.objects.get().numero, 5)  # 05 do RE é a primeira do sistema

    def test_numero_sequencial_por_tipo_e_gerar_de_novo_mantem(self):
        importar_planilha_bdo(ler_planilha_bdo(_planilha([_linha("4001", "Itu", "SP", "10,00", "99000004")])),
                              "RE", "FR_4001.xlsx", self.admin)
        importar_planilha_bdo(ler_planilha_bdo(_planilha([_linha("4002", "Itu", "SP", "10,00", "99000004")])),
                              "RI", "FR_4002.xlsx", self.admin)
        self._gerar(boletim="4000")
        self._gerar(boletim="4001")
        self._gerar(boletim="4000")
        self._gerar(boletim="4002")
        self.assertEqual(
            dict(MedicaoBdo.objects.values_list("boletim__numero", "numero")), {"4000": 5, "4001": 6, "4002": 1},
        )
        self.assertEqual(MedicaoBdo.objects.get(boletim__numero="4000").gerada_por, self.analista)

    def test_retirado_no_modal_nao_entra_e_varios_boletins_viram_zip(self):
        importar_planilha_bdo(ler_planilha_bdo(_planilha([_linha("4001", "Itu", "SP", "10,00", "99000004")])),
                              "RE", "FR_4001.xlsx", self.admin)
        itu = MunicipioBdo.objects.get(boletim__numero="4000", municipio="Itu").pk
        resp = self._gerar(boletim=["4000", "4001"], retirado=str(itu))
        self.assertEqual(resp["Content-Type"], "application/zip")
        with zipfile.ZipFile(io.BytesIO(resp.content)) as arquivo_zip:
            self.assertEqual(
                sorted(arquivo_zip.namelist()),
                ["MEDIÇÃO 05 - MANUTENÇÃO 4000.xlsx", "MEDIÇÃO 06 - MANUTENÇÃO 4001.xlsx"],
            )
            _aba, linhas = self._linhas(arquivo_zip.read("MEDIÇÃO 05 - MANUTENÇÃO 4000.xlsx"))
        self.assertEqual([linha[1] for linha in linhas], ["SÃO PAULO", "SÃO PAULO"])

    def test_sem_boletim_ou_boletim_sem_tipo_nao_gera(self):
        resp = self.client.post(reverse("bdo_medicao"), {}, follow=True)
        self.assertContains(resp, "Medição não gerada: escolha um Boletim")
        BoletimBdo.objects.filter(numero="4000").update(tipo="")
        self._conferir_eace()
        resp = self.client.post(reverse("bdo_medicao"), {"boletim": "4000"}, follow=True)
        self.assertContains(resp, "Boletim sem tipo (RE/RI): 4000")
        self.assertFalse(MedicaoBdo.objects.exists())

    def test_visualizador_e_anonimo_nao_geram(self):
        visualizador = User.objects.create_user(
            username="medicao-vis", password="senha-teste-123", perfil=User.PERFIL_VISUALIZADOR,
        )
        self.client.force_login(visualizador)
        self.assertNotEqual(self._gerar(boletim="4000").status_code, 200)
        self.client.logout()
        self.assertIn(reverse("login"), self._gerar(boletim="4000")["Location"])
        self.assertFalse(MedicaoBdo.objects.exists())

    def test_modal_tem_o_botao_gerar_medicao(self):
        resp = self.client.post(reverse("bdo_recebiveis"), {"boletim": "4000"})
        self.assertContains(resp, f'data-gerar-medicao="{reverse("bdo_medicao")}"')
        self.assertContains(resp, "Gerar Medição (.xlsx)")



class BdoExportarExcelTests(TestCase):
    """Pedido do usuário (2026-10-16): botão "Excel" do grid do BDO com os
    filtros e o agrupamento da tela, sem paginação."""

    def setUp(self):
        admin = User.objects.create_user(
            username="bdo-excel-admin", password="senha-teste-123", perfil=User.PERFIL_ADMINISTRADOR,
        )
        importar_planilha_bdo(ler_planilha_bdo(_planilha(LINHAS_3658)), "RI", "FR_3658.xlsx", admin)
        self.client.force_login(User.objects.create_user(
            username="bdo-excel-analista", password="senha-teste-123", perfil=User.PERFIL_ANALISTA,
        ))

    def _exportar(self, **filtros):
        resp = self.client.get(reverse("bdo"), {**filtros, "exportar": "excel"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment;", resp["Content-Disposition"])
        aba = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        linhas = list(aba.iter_rows(values_only=True))
        return linhas[0], linhas[1:]

    def test_tela_tem_o_botao_com_os_filtros(self):
        resp = self.client.get(reverse("bdo"), {"uf": "SP", "agrupar": "municipio"})
        self.assertContains(resp, "?uf=SP&amp;agrupar=municipio&exportar=excel")

    def test_por_inep_exporta_todas_as_linhas_dos_filtros(self):
        titulos, linhas = self._exportar()
        self.assertEqual(titulos[0], "INEP")
        self.assertIn("Valor proporcional", titulos)
        self.assertEqual(len(linhas), 8)
        titulos, linhas = self._exportar(uf="SP")
        valor_total = {linha[0]: linha[titulos.index("Valor total")] for linha in linhas}
        self.assertEqual(valor_total, {"35058592": 49.92, "35011006": 277.89})

    def test_exporta_alem_da_primeira_pagina(self):
        ineps = ", ".join(str(50100000 + numero) for numero in range(60))
        admin = User.objects.get(username="bdo-excel-admin")
        importar_planilha_bdo(
            ler_planilha_bdo(_planilha([_linha("3700", "DOURADOS", "MS", "100,00", ineps)])), "RE", "FR_3700.xlsx", admin,
        )
        _titulos, linhas = self._exportar(boletim="3700")
        self.assertEqual(len(linhas), 60)

    def test_por_municipio_e_por_uf(self):
        titulos, linhas = self._exportar(agrupar="municipio")
        self.assertEqual(titulos[:3], ("Município", "UF", "INEPs"))
        self.assertIn("Total MEGA", titulos)
        self.assertEqual(len(linhas), 4)
        titulos, linhas = self._exportar(agrupar="uf")
        self.assertEqual(titulos[:3], ("UF", "Municípios", "INEPs"))
        self.assertEqual(sorted(linha[0] for linha in linhas), ["DF", "MS", "SP"])

    def test_anonimo_nao_exporta(self):
        self.client.logout()
        resp = self.client.get(reverse("bdo"), {"exportar": "excel"})
        self.assertIn(reverse("login"), resp["Location"])
