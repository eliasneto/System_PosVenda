"""Planilha de Medição da manutenção recorrente de um Boletim (pedido do
usuário, 2026-10-15), gerada pelo modal "Recebíveis" da tela BDO no mesmo
formato da "MEDIÇÃO 05 - MANUTENÇÃO 3657" do cliente: uma linha por linha
do Boletim que ficou no modal, com UF, município, 1º faturamento, valor,
Proporcional/Integral e o total; a aba NOTAS FISCAIS repete as linhas com a
coluna da nota em branco. O modelo (cabeçalho e logos) fica em
`planilhas/medicao_manutencao_modelo.xlsx`."""

import datetime
import io
from copy import copy
from pathlib import Path

import openpyxl
from django.db import transaction
from django.db.models import Max
from django.utils import timezone

from .bdo import _normalizar
from .models import BoletimBdo, MedicaoBdo, PrimeiroFaturamentoInep

MODELO = Path(__file__).resolve().parent / "planilhas" / "medicao_manutencao_modelo.xlsx"
# Linhas do modelo: dados a partir da 8, total na 9, observações na 13 e 14.
LINHA_DADOS, LINHA_TOTAL, LINHA_OBSERVACAO = 8, 9, 13
COLUNAS = "ABCDEFGHIJ"
# Medições feitas fora do sistema antes dele (pedido do usuário, 2026-10-15):
# a 05 do RE é a do Boletim 3657, então a próxima do RE é a 05.
ULTIMA_MEDICAO_FORA_DO_SISTEMA = {BoletimBdo.TIPO_RE: 4}


class MedicaoBdoError(Exception):
    pass


def numero_da_medicao(boletim, usuario):
    """Número sequencial por tipo; o Boletim que já tem Medição mantém o dele."""
    medicao = MedicaoBdo.objects.filter(boletim=boletim).first()
    if medicao is None:
        ultima = MedicaoBdo.objects.filter(tipo=boletim.tipo).aggregate(ultima=Max("numero"))["ultima"] or 0
        medicao = MedicaoBdo(
            boletim=boletim, tipo=boletim.tipo,
            numero=max(ultima, ULTIMA_MEDICAO_FORA_DO_SISTEMA.get(boletim.tipo, 0)) + 1,
        )
    medicao.gerada_por = usuario
    medicao.save()
    return medicao.numero


def primeiro_faturamento(tipo, ineps, hoje):
    """Data do 1º faturamento de cada INEP no tipo; o INEP que ainda não tem
    recebe `hoje` e fica com ela nos meses seguintes."""
    datas = dict(PrimeiroFaturamentoInep.objects.filter(tipo=tipo, inep__in=ineps).values_list("inep", "data"))
    novos = [PrimeiroFaturamentoInep(tipo=tipo, inep=inep, data=hoje) for inep in set(ineps) - set(datas)]
    PrimeiroFaturamentoInep.objects.bulk_create(novos, ignore_conflicts=True)
    datas.update({novo.inep: hoje for novo in novos})
    return datas


@transaction.atomic
def gerar_medicoes(linhas, situacoes, usuario):
    """`linhas`: as linhas do Boletim que ficaram no modal (`MunicipioBdo` com
    `lista_ineps` e `valor_recebivel`); `situacoes[(tipo, uf, municipio sem acento)]`:
    "Proporcional" ou "Integral". Devolve `[(nome_do_arquivo, bytes)]`, um por
    Boletim. A data de cada linha é a mais antiga dos INEPs dela."""
    por_boletim = {}
    for linha in linhas:
        por_boletim.setdefault(linha.boletim, []).append(linha)
    sem_tipo = sorted(boletim.numero for boletim in por_boletim if not boletim.tipo)
    if sem_tipo:
        raise MedicaoBdoError(f"Boletim sem tipo (RE/RI): {', '.join(sem_tipo)}. Informe o tipo antes de gerar a Medição.")

    hoje = timezone.localdate()
    arquivos = []
    for boletim, linhas_boletim in sorted(por_boletim.items(), key=lambda item: item[0].numero):
        datas = primeiro_faturamento(
            boletim.tipo, [inep for linha in linhas_boletim for inep in linha.lista_ineps], hoje,
        )
        numero = numero_da_medicao(boletim, usuario)
        dados = [
            (
                linha.uf, linha.municipio.upper(),
                min(datas[inep] for inep in linha.lista_ineps),
                linha.valor_recebivel,
                (situacoes.get((boletim.tipo, linha.uf, _normalizar(linha.municipio))) or "").upper(),
            )
            for linha in sorted(linhas_boletim, key=lambda linha: (linha.uf, linha.municipio.upper(), linha.pk))
        ]
        nome = f"MEDIÇÃO {numero:02d} - MANUTENÇÃO {boletim.numero}"
        arquivos.append((f"{nome}.xlsx", _planilha(dados, numero, boletim.numero)))
    return arquivos


def _planilha(dados, numero, numero_boletim):
    workbook = openpyxl.load_workbook(MODELO)
    for aba in workbook.worksheets:
        _preencher(aba, dados, numero_boletim)
    workbook.worksheets[0].title = f"MEDIÇÃO {numero:02d}"
    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()


def _preencher(aba, dados, numero_boletim):
    def estilos(linha):
        return {coluna: copy(aba[f"{coluna}{linha}"]._style) for coluna in COLUNAS}

    estilo_dados, estilo_total = estilos(LINHA_DADOS), estilos(LINHA_TOTAL)
    estilo_observacao = [estilos(LINHA_OBSERVACAO), estilos(LINHA_OBSERVACAO + 1)]
    for intervalo in [r for r in aba.merged_cells.ranges if r.min_row >= LINHA_DADOS]:
        aba.unmerge_cells(str(intervalo))
    aba.delete_rows(LINHA_DADOS, aba.max_row)

    for indice, (uf, municipio, data, valor, situacao) in enumerate(dados):
        linha = LINHA_DADOS + indice
        for coluna in COLUNAS:
            aba[f"{coluna}{linha}"]._style = copy(estilo_dados[coluna])
        aba[f"D{linha}"] = uf
        aba[f"E{linha}"] = municipio
        aba[f"F{linha}"] = datetime.datetime.combine(data, datetime.time())
        aba[f"G{linha}"] = valor
        aba[f"H{linha}"] = situacao
        aba[f"I{linha}"] = f"=G{linha}"

    ultima = LINHA_DADOS + len(dados) - 1
    total = ultima + 1
    for coluna in COLUNAS:
        aba[f"{coluna}{total}"]._style = copy(estilo_total[coluna])
    aba[f"I{total}"] = f"=SUM(I{LINHA_DADOS}:I{ultima})"

    textos = (
        f"OBS NA NF: CÓDIGO DE LIBERAÇÃO {numero_boletim}, FATURAR EM NOTAS DISTINTAS POR MUNÍCIPIO",
        "VENCIMENTO PARA 30 DIAS, INSERIR DADOS BANCÁRIOS",
    )
    for deslocamento, (texto, estilo) in enumerate(zip(textos, estilo_observacao)):
        linha = total + 4 + deslocamento
        for coluna in COLUNAS:
            aba[f"{coluna}{linha}"]._style = copy(estilo[coluna])
        aba[f"A{linha}"] = texto
        aba.merge_cells(f"A{linha}:E{linha}")
