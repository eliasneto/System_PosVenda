import io
import shutil
import tempfile
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

import openpyxl
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase, override_settings
from django.urls import reverse

from apps.escolas.models import Escola, EscolaItemRelatorioEaceMip, Lote
from apps.escolas.services import _normalizar_texto_cidade, criar_lote_mip, valor_total_lotes_mip_por_status
from apps.escolas.tests import _criar_escola_elegivel_lote
from apps.ri.models import KitPadrao, Ri, RiHistorico, RiItemIxc

User = get_user_model()

# `--aplicar` faz o backup de segurança do banco (mysqldump) antes de gravar
# — simulado em todos os testes deste módulo.
_PATCH_BACKUP = patch(
    "apps.escolas.management.commands.importar_lotes_mip_em_massa.criar_backup_seguranca",
    return_value="/backups/seguranca_teste.sql.gz",
)


def setUpModule():
    _PATCH_BACKUP.start()


def tearDownModule():
    _PATCH_BACKUP.stop()

CABECALHO = [
    "Projeto", "Cod Fornecedor", "Fornecedor", "Num Obra", "Cod Produto", "Descrição do Item", "Qtde Produto",
    "Qtde Produto UR", "Valor Unit UR", "Valor Produto", "Num ACS", "Valor Liberado ACS", "Data Emissão ACS",
    "Prod Serv", "UF", "CIDADE",
]
DESCRICAO_KIT = "Kit Cobertura Wi-Fi - 2 Access Points - Serv - MEGA - SE"


def _linha(inep, valor="300.00", cidade="Abadiânia", uf="GO"):
    return [
        int(inep), 19001, "MEGA INFRA", "4-IMPLANTAÇÃO_DE_REDE_INTERNA", 50876, DESCRICAO_KIT, 1, 1,
        float(valor), float(valor), 1, float(valor), "22/09/2026", "Serviço", uf, cidade,
    ]


class ImportarLotesMipEmMassaTests(TestCase):
    """Pedido do usuário (2026-09-29): comando `importar_lotes_mip_em_massa`
    (planilha "BASE CONSOLIDADA MIP")."""

    def setUp(self):
        self.usuario = User.objects.create_user(username="analista-importacao", password="senha-teste-123")
        KitPadrao.objects.create(
            descricao="Kit Cobertura Wi-Fi - 2 Access Points", lote=9,
            valor_equipamento="1000.00", valor_servico="300.00",
        )
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def _planilha(self, linhas, faturamento=None):
        workbook = openpyxl.Workbook()
        aba = workbook.active
        aba.title = "CONSOLIDADA"
        aba.append(CABECALHO)
        for linha in linhas:
            aba.append(linha)
        if faturamento is not None:
            aba_fat = workbook.create_sheet("FATURAMENTO")
            aba_fat.append(["INEP", "FATURADOS ", "DATA FATURAMENTO "])
            for inep in faturamento:
                aba_fat.append([int(inep), 1, None])
        caminho = self.tmp / "BASE CONSOLIDADA MIP.xlsx"
        workbook.save(caminho)
        return str(caminho)

    def _rodar(self, caminho, *extra):
        saida = io.StringIO()
        call_command("importar_lotes_mip_em_massa", caminho, "--usuario", self.usuario.username, *extra, stdout=saida)
        return saida.getvalue()

    def _escola_sem_lado3(self, inep, municipio="Abadiânia"):
        escola, ri = _criar_escola_elegivel_lote(inep, municipio=municipio)
        escola.itens_relatorio_eace_mip.all().delete()
        return escola, ri

    def test_simulacao_nao_grava_nada(self):
        _criar_escola_elegivel_lote("20000001")
        self._escola_sem_lado3("20000002")
        saida = self._rodar(self._planilha([_linha("20000001"), _linha("20000002")]))
        self.assertIn("SIMULACAO", saida)
        self.assertIn("Incluído em LOTE: 2", saida)
        self.assertEqual(Lote.objects.count(), 0)
        self.assertFalse(EscolaItemRelatorioEaceMip.objects.filter(escola__inep="20000002").exists())
        self.assertEqual(Escola.objects.get(inep="20000001").status_mip, Escola.AGUARDANDO_VALIDACAO_EACE)

    def test_aplicar_cria_lote_concluido_e_marcado_como_importacao_em_massa(self):
        escola, ri = _criar_escola_elegivel_lote("20000010")
        self._rodar(self._planilha([_linha("20000010")]), "--aplicar")

        lote = Lote.objects.get()
        self.assertEqual(lote.status, Lote.FATURAMENTO_CONCLUIDO)
        self.assertTrue(lote.importado_em_massa)
        self.assertEqual(lote.arquivo_importacao, "BASE CONSOLIDADA MIP.xlsx")
        self.assertEqual(lote.criado_por, self.usuario)
        self.assertEqual((lote.estado, lote.municipio), ("GO", "Abadiânia"))
        self.assertEqual(list(lote.escolas.all()), [escola])
        escola.refresh_from_db()
        self.assertEqual(escola.status_mip, Escola.FATURAMENTO_CONCLUIDO)
        ri.refresh_from_db()
        self.assertEqual(ri.status, Ri.FATURAMENTO_RI_CONCLUIDO)
        historico = RiHistorico.objects.filter(ri=ri)
        self.assertTrue(historico.filter(campo="LOTE", valor_novo=str(lote)).exists())
        self.assertTrue(historico.filter(campo="Status (MIP)", valor_novo="Processo Concluído").exists())

    def test_lado_eace_vazio_com_ixc_igual_a_planilha_e_preenchido_e_incluido(self):
        escola, ri = self._escola_sem_lado3("20000020")
        self._rodar(self._planilha([_linha("20000020")]), "--aplicar")

        item = EscolaItemRelatorioEaceMip.objects.get(escola=escola)
        self.assertEqual(item.valor_servico, 300)
        self.assertEqual(item.cidade, "Abadiânia")
        self.assertTrue(Lote.objects.get().escolas.filter(pk=escola.pk).exists())
        self.assertTrue(RiHistorico.objects.filter(ri=ri, campo__startswith="Relatório EACE (MIP)").exists())

    def test_valor_divergente_fica_de_fora_sem_gravar_lado_eace(self):
        escola, _ri = self._escola_sem_lado3("20000030")
        _criar_escola_elegivel_lote("20000031")
        saida = self._rodar(self._planilha([_linha("20000030", "999.00"), _linha("20000031", "999.00")]), "--aplicar")

        self.assertEqual(Lote.objects.count(), 0)
        self.assertFalse(EscolaItemRelatorioEaceMip.objects.filter(escola=escola).exists())
        self.assertIn("Valor da planilha, IXC e EACE não batem: 2", saida)

    def test_inep_ja_em_lote_nao_e_mexido(self):
        escola, _ri = _criar_escola_elegivel_lote("20000040")
        lote_existente = criar_lote_mip("GO", "Abadiânia", None, None, self.usuario)
        saida = self._rodar(self._planilha([_linha("20000040")]), "--aplicar")

        self.assertEqual(list(Lote.objects.all()), [lote_existente])
        lote_existente.refresh_from_db()
        self.assertEqual(lote_existente.status, Lote.EM_ANDAMENTO)
        self.assertFalse(lote_existente.importado_em_massa)
        escola.refresh_from_db()
        self.assertEqual(escola.status_mip, Escola.EM_ANDAMENTO)
        self.assertIn("Já estava em LOTE (não mexido): 1", saida)

    def test_status_mip_nao_elegivel_e_inep_nao_cadastrado_ficam_de_fora(self):
        escola, _ri = _criar_escola_elegivel_lote("20000050")
        Escola.objects.filter(pk=escola.pk).update(status_mip=Escola.EM_ANDAMENTO)
        saida = self._rodar(self._planilha([_linha("20000050"), _linha("29999999")]), "--aplicar")

        self.assertEqual(Lote.objects.count(), 0)
        self.assertIn('Status (MIP) diferente de "Aguardando Validação EACE": 1', saida)
        self.assertIn("INEP não cadastrado no sistema: 1", saida)

    def test_um_lote_por_municipio(self):
        _criar_escola_elegivel_lote("20000060")
        _criar_escola_elegivel_lote("20000061")
        _criar_escola_elegivel_lote("20000062", municipio="Anápolis")
        self._rodar(self._planilha([
            _linha("20000060"), _linha("20000061"), _linha("20000062", cidade="Anápolis"),
        ]), "--aplicar")

        contagem = {lote.municipio: lote.escolas.count() for lote in Lote.objects.all()}
        self.assertEqual(contagem, {"Abadiânia": 2, "Anápolis": 1})

    def test_grafias_diferentes_do_municipio_viram_um_lote_so(self):
        _criar_escola_elegivel_lote("20000065", estado="SP", municipio="São Paulo")
        _criar_escola_elegivel_lote("20000066", estado="SP", municipio="São Paulo")
        maiuscula, ri = _criar_escola_elegivel_lote("20000067", estado="SP", municipio="SAO PAULO")
        saida = self._rodar(self._planilha([
            _linha("20000065", cidade="São Paulo", uf="SP"), _linha("20000066", cidade="São Paulo", uf="SP"),
            _linha("20000067", cidade="SAO PAULO", uf="SP"),
        ]), "--aplicar")

        lote = Lote.objects.get()
        self.assertEqual((lote.municipio, lote.escolas.count()), ("São Paulo", 3))
        maiuscula.refresh_from_db()
        self.assertEqual(maiuscula.municipio, "São Paulo")
        self.assertTrue(RiHistorico.objects.filter(
            ri=ri, campo="Município", valor_anterior="SAO PAULO", valor_novo="São Paulo",
        ).exists())
        self.assertIn("SAO PAULO -> São Paulo", saida)

    def test_simulacao_nao_corrige_municipio(self):
        _criar_escola_elegivel_lote("20000068", estado="SP", municipio="São Paulo")
        _criar_escola_elegivel_lote("20000069", estado="SP", municipio="SAO PAULO")
        self._rodar(self._planilha([_linha("20000068", uf="SP"), _linha("20000069", uf="SP")]))
        self.assertEqual(Escola.objects.get(inep="20000069").municipio, "SAO PAULO")

    def test_inep_elegivel_fora_da_planilha_nao_entra_no_lote(self):
        _criar_escola_elegivel_lote("20000070")
        fora = _criar_escola_elegivel_lote("20000071")[0]
        self._rodar(self._planilha([_linha("20000070")]), "--aplicar")

        self.assertFalse(Lote.objects.get().escolas.filter(pk=fora.pk).exists())
        fora.refresh_from_db()
        self.assertEqual(fora.status_mip, Escola.AGUARDANDO_VALIDACAO_EACE)

    def test_relatorio_csv_lista_ineps_so_da_aba_faturamento(self):
        _criar_escola_elegivel_lote("20000080")
        relatorio = self.tmp / "relatorio.csv"
        saida = self._rodar(
            self._planilha([_linha("20000080")], faturamento=["20000080", "20000081"]), "--relatorio", str(relatorio),
        )
        self.assertIn("Só na aba FATURAMENTO, sem itens (não importados): 1", saida)
        conteudo = relatorio.read_text(encoding="utf-8-sig")
        self.assertIn("20000080;GO;Abadiânia;Incluído em LOTE", conteudo)
        self.assertIn("20000081;;;Só na aba FATURAMENTO", conteudo)

    def test_planilha_sem_colunas_ou_usuario_inexistente_levanta_command_error(self):
        workbook = openpyxl.Workbook()
        workbook.active.append(["INEP", "Outra"])
        caminho = self.tmp / "invalida.xlsx"
        workbook.save(caminho)
        with self.assertRaises(CommandError):
            self._rodar(str(caminho))
        with self.assertRaises(CommandError):
            call_command("importar_lotes_mip_em_massa", self._planilha([]), "--usuario", "nao-existe")

    def test_tela_de_lotes_mostra_selo_de_importacao_em_massa(self):
        _criar_escola_elegivel_lote("20000090")
        self._rodar(self._planilha([_linha("20000090")]), "--aplicar")
        self.client.force_login(self.usuario)
        resposta = self.client.get(reverse("mip_lote_inep"))
        self.assertContains(resposta, "Importação em massa")


class ImportarLotesMipIxcComItemSemValorTests(TestCase):
    """IXC com item sem Valor de serviço no catálogo = total incompleto —
    nunca entra, mesmo com o Lado EACE vazio (não preenche)."""

    def test_ixc_incompleto_nao_preenche_nem_inclui(self):
        usuario = User.objects.create_user(username="analista-importacao-2", password="senha-teste-123")
        escola = Escola.objects.create(
            inep="21000001", nome="Escola", estado="GO", municipio="Abadiânia", lote=9,
            status_mip=Escola.AGUARDANDO_VALIDACAO_EACE,
        )
        ri = Ri.objects.create(escola=escola, status=Ri.AGUARDANDO_VALIDACAO_EACE)
        RiItemIxc.objects.create(ri=ri, descricao_item="Produto sem catálogo", quantidade=1, valor_unitario="0.00")
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
        workbook = openpyxl.Workbook()
        workbook.active.append(CABECALHO)
        workbook.active.append(_linha("21000001"))
        caminho = tmp / "base.xlsx"
        workbook.save(caminho)

        call_command("importar_lotes_mip_em_massa", str(caminho), "--usuario", usuario.username, "--aplicar",
                     stdout=io.StringIO())

        self.assertEqual(Lote.objects.count(), 0)
        self.assertFalse(EscolaItemRelatorioEaceMip.objects.filter(escola=escola).exists())


class ImportarLotesMipIncluirDivergentesTests(TestCase):
    """Pedido do usuário (2026-10-07): `--incluir-divergentes` — o total de
    "Processo Concluído" tem que bater com a planilha; quem diverge entra
    valendo o faturado na planilha, sem mexer no Lado IXC."""

    setUp = ImportarLotesMipEmMassaTests.setUp
    _planilha = ImportarLotesMipEmMassaTests._planilha
    _rodar = ImportarLotesMipEmMassaTests._rodar

    def test_sem_a_opcao_divergente_continua_de_fora(self):
        _criar_escola_elegivel_lote("20100001")
        saida = self._rodar(self._planilha([_linha("20100001", "999.00")]), "--aplicar")
        self.assertIn("Valor da planilha, IXC e EACE não batem: 1", saida)
        self.assertEqual(Lote.objects.count(), 0)

    def test_divergentes_entram_com_valor_da_planilha_e_total_bate(self):
        _criar_escola_elegivel_lote("20100010")  # bate: 300
        divergente, ri_divergente = _criar_escola_elegivel_lote("20100011")  # IXC 300, planilha 999
        em_andamento, _ = _criar_escola_elegivel_lote("20100012")  # Status (MIP) fora da regra
        em_andamento.status_mip = Escola.EM_ANDAMENTO
        em_andamento.save()
        saida = self._rodar(self._planilha([
            _linha("20100010"), _linha("20100011", "999.00"), _linha("20100012", "450.00"),
        ]), "--aplicar", "--incluir-divergentes")

        lote = Lote.objects.get()
        self.assertEqual(lote.status, Lote.FATURAMENTO_CONCLUIDO)
        self.assertEqual(lote.escolas.count(), 3)
        self.assertEqual(
            lote.valores_faturados_planilha, {"20100010": "300.00", "20100011": "999.00", "20100012": "450.00"}
        )
        self.assertEqual(lote.valor_faturado_planilha(divergente), Decimal("999.00"))
        self.assertEqual(valor_total_lotes_mip_por_status()[Lote.FATURAMENTO_CONCLUIDO], Decimal("1749.00"))
        self.assertIn("Diferença: R$ 0.00", saida)
        self.assertIn("Incluído em LOTE com o valor da planilha (IXC diverge ou Status (MIP) fora da regra): 2", saida)

        # Lado IXC intacto; histórico explica o valor.
        self.assertEqual(ri_divergente.itens_ixc.get().descricao_item, "Kit Cobertura Wi-Fi - 2 Access Points")
        mensagem = RiHistorico.objects.get(ri=ri_divergente, tipo=RiHistorico.IMPORTACAO_MASSA).mensagem
        self.assertIn("R$ 999.00 faturado na planilha (Valor Total IXC: R$ 300.00 — divergente", mensagem)
        em_andamento.refresh_from_db()
        self.assertEqual(em_andamento.status_mip, Escola.FATURAMENTO_CONCLUIDO)

    def test_cidade_so_com_divergentes_tambem_ganha_lote(self):
        _criar_escola_elegivel_lote("20100020")
        self._rodar(self._planilha([_linha("20100020", "999.00")]), "--aplicar", "--incluir-divergentes")
        lote = Lote.objects.get()
        self.assertTrue(lote.importado_em_massa)
        self.assertEqual(valor_total_lotes_mip_por_status()[Lote.FATURAMENTO_CONCLUIDO], Decimal("999.00"))

    def test_tela_de_lotes_usa_valor_da_planilha_no_total_por_status(self):
        _criar_escola_elegivel_lote("20100030")
        self._rodar(self._planilha([_linha("20100030", "999.00")]), "--aplicar", "--incluir-divergentes")
        self.usuario.is_superuser = True
        self.usuario.save()
        self.client.force_login(self.usuario)
        resposta = self.client.get(reverse("mip_lote_inep"))
        concluido = next(r for r in resposta.context["resumo_por_status"] if r["status"] == Lote.FATURAMENTO_CONCLUIDO)
        self.assertEqual(concluido["valor_total"], Decimal("999.00"))
        self.assertContains(resposta, "No LOTE: R$ 999")

    def test_lote_criado_pela_tela_continua_usando_ixc(self):
        escola, _ = _criar_escola_elegivel_lote("20100040")
        lote = criar_lote_mip("GO", "Abadiânia", None, None, self.usuario)
        self.assertIsNone(lote.valor_faturado_planilha(escola))
        self.assertEqual(valor_total_lotes_mip_por_status()[Lote.EM_ANDAMENTO], Decimal("300.00"))


class NormalizarTextoCidadeTests(TestCase):
    """Cidade do rateio × Município do LOTE: acento e caixa não importam."""

    def test_ignora_acento_e_caixa(self):
        self.assertEqual(_normalizar_texto_cidade(" SAO PAULO "), _normalizar_texto_cidade("São Paulo"))
        self.assertNotEqual(_normalizar_texto_cidade("São Paulo2"), _normalizar_texto_cidade("São Paulo"))


class SeloImportacaoEmMassaEmEquipamentosTests(TestCase):
    """Pedido do usuário (2026-09-30): INEP de LOTE importado em massa
    mostra "Importação em massa" no grid de Equipamentos e no RI."""

    def setUp(self):
        self.usuario = User.objects.create_user(username="analista-selo", password="senha-teste-123")
        KitPadrao.objects.create(
            descricao="Kit Cobertura Wi-Fi - 2 Access Points", lote=9,
            valor_equipamento="1000.00", valor_servico="300.00",
        )
        self.client.force_login(self.usuario)

    def test_selo_aparece_so_para_lote_importado_em_massa(self):
        importado, _ri = _criar_escola_elegivel_lote("22000001")
        criar_lote_mip("GO", "Abadiânia", None, None, self.usuario)
        Lote.objects.update(importado_em_massa=True)
        comum, _ri = _criar_escola_elegivel_lote("22000002", municipio="Anápolis")
        criar_lote_mip("GO", "Anápolis", None, None, self.usuario)

        self.assertTrue(importado.em_lote_importado_em_massa)
        self.assertFalse(comum.em_lote_importado_em_massa)
        grid = self.client.get(reverse("grid_inep")).content.decode()
        self.assertEqual(grid.count(">Importação em massa<"), 1)
        self.assertContains(self.client.get(reverse("ri_detail", args=[importado.inep])), ">Importação em massa<")
        self.assertNotContains(self.client.get(reverse("ri_detail", args=[comum.inep])), ">Importação em massa<")


class HistoricoImportacaoEmMassaTests(TestCase):
    """Pedido do usuário (2026-09-30): cada INEP importado ganha 1 entrada
    "Importação em massa" no histórico, com data, usuário e o que mudou."""

    def setUp(self):
        self.usuario = User.objects.create_user(
            username="analista-historico", password="senha-teste-123", first_name="Ana", last_name="Lima",
        )
        KitPadrao.objects.create(
            descricao="Kit Cobertura Wi-Fi - 2 Access Points", lote=9,
            valor_equipamento="1000.00", valor_servico="300.00",
        )
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def _aplicar(self, linhas):
        workbook = openpyxl.Workbook()
        workbook.active.append(CABECALHO)
        for linha in linhas:
            workbook.active.append(linha)
        caminho = self.tmp / "BASE.xlsx"
        workbook.save(caminho)
        call_command("importar_lotes_mip_em_massa", str(caminho), "--usuario", self.usuario.username, "--aplicar",
                     stdout=io.StringIO())

    def _mensagem(self, ri):
        return RiHistorico.objects.get(ri=ri, tipo=RiHistorico.IMPORTACAO_MASSA).mensagem

    def test_registra_quando_quem_e_o_que_mudou(self):
        _escola, ri_so_status = _criar_escola_elegivel_lote("23000001", estado="SP", municipio="São Paulo")
        com_lado3, ri_com_lado3 = _criar_escola_elegivel_lote("23000002", estado="SP", municipio="SAO PAULO")
        com_lado3.itens_relatorio_eace_mip.all().delete()
        fora, ri_fora = _criar_escola_elegivel_lote("23000003")
        self._aplicar([
            _linha("23000001", uf="SP"), _linha("23000002", uf="SP"), _linha("23000003", "999.00"),
        ])
        lote = Lote.objects.get()

        mensagem = self._mensagem(ri_so_status)
        self.assertIn("Processo importado em massa em ", mensagem)
        self.assertIn("por Ana Lima (arquivo BASE.xlsx)", mensagem)
        self.assertIn(f"- LOTE: {lote} (São Paulo/SP)", mensagem)
        self.assertIn("- Status (MIP): Aguardando Validação EACE → Processo Concluído", mensagem)
        self.assertIn("- Status do RI: sem alteração (Faturamento RI Concluído)", mensagem)
        self.assertIn("- Equipamentos: sem alteração", mensagem)
        self.assertIn("- Município: sem alteração (São Paulo)", mensagem)
        self.assertEqual(RiHistorico.objects.get(ri=ri_so_status, tipo=RiHistorico.IMPORTACAO_MASSA).autor, self.usuario)

        mensagem = self._mensagem(ri_com_lado3)
        self.assertIn(
            "- Equipamentos (Relatório EACE MIP): incluídos pela planilha — Kit Cobertura Wi-Fi - 2 Access Points (1 un.)",
            mensagem,
        )
        self.assertIn("- Município: SAO PAULO → São Paulo", mensagem)
        self.assertFalse(RiHistorico.objects.filter(ri=ri_fora, tipo=RiHistorico.IMPORTACAO_MASSA).exists())

    def test_painel_do_historico_mostra_a_entrada(self):
        escola, _ri = _criar_escola_elegivel_lote("23000010")
        self._aplicar([_linha("23000010")])
        self.client.force_login(self.usuario)
        resposta = self.client.get(reverse("ri_detail", args=[escola.inep]))
        self.assertContains(resposta, "Processo importado em massa em ")
        self.assertContains(resposta, "Status (MIP): Aguardando Validação EACE → Processo Concluído")


class MipSoComInepDoRelatorioEaceTests(TestCase):
    """Pedido do usuário (2026-09-30): o grid do MIP e o "Criar LOTE" só
    consideram INEP que veio no Relatório EACE (MIP) na última
    sincronização (`encontrado_relatorio_eace_mip=True`)."""

    def setUp(self):
        self.usuario = User.objects.create_user(username="analista-mip-planilha", password="senha-teste-123")
        KitPadrao.objects.create(
            descricao="Kit Cobertura Wi-Fi - 2 Access Points", lote=9,
            valor_equipamento="1000.00", valor_servico="300.00",
        )
        self.na_planilha, _ = _criar_escola_elegivel_lote("24000001")
        self.fora, _ = _criar_escola_elegivel_lote("24000002")
        self.nunca_sincronizado, _ = _criar_escola_elegivel_lote("24000003")
        Escola.objects.filter(pk=self.fora.pk).update(encontrado_relatorio_eace_mip=False)
        Escola.objects.filter(pk=self.nunca_sincronizado.pk).update(encontrado_relatorio_eace_mip=None)

    def test_grid_mostra_so_inep_da_planilha(self):
        self.client.force_login(self.usuario)
        resposta = self.client.get(reverse("mip_inep"))
        ineps = {linha["escola"].inep for linha in resposta.context["page_obj"]}
        self.assertEqual(ineps, {"24000001"})

    def test_criar_lote_so_com_inep_da_planilha(self):
        lote = criar_lote_mip("GO", "Abadiânia", None, None, self.usuario)
        self.assertEqual(list(lote.escolas.all()), [self.na_planilha])

    def test_importacao_em_massa_nao_depende_da_sincronizacao(self):
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
        workbook = openpyxl.Workbook()
        workbook.active.append(CABECALHO)
        workbook.active.append(_linha("24000003"))
        caminho = tmp / "base.xlsx"
        workbook.save(caminho)
        call_command("importar_lotes_mip_em_massa", str(caminho), "--usuario", self.usuario.username, "--aplicar",
                     stdout=io.StringIO())
        self.assertEqual(list(Lote.objects.get().escolas.all()), [self.nunca_sincronizado])


class RetirarDoGridMipTests(TestCase):
    """Pedido do usuário (2026-09-30): botão para retirar do grid do MIP os
    INEPs que sobraram — só some da lista, volta numa nova sincronização."""

    def setUp(self):
        self.usuario = User.objects.create_user(username="analista-retirar", password="senha-teste-123")
        KitPadrao.objects.create(
            descricao="Kit Cobertura Wi-Fi - 2 Access Points", lote=9,
            valor_equipamento="1000.00", valor_servico="300.00",
        )
        self.escola1, self.ri1 = _criar_escola_elegivel_lote("25000001")
        self.escola2, _ = _criar_escola_elegivel_lote("25000002")
        self.escola3, _ = _criar_escola_elegivel_lote("25000003")
        self.client.force_login(self.usuario)

    def _ineps_no_grid(self):
        resposta = self.client.get(reverse("mip_inep"))
        return {linha["escola"].inep for linha in resposta.context["page_obj"]}

    def test_retira_um_inep_so_do_grid_sem_mexer_no_resto(self):
        resposta = self.client.post(reverse("mip_retirar_grid"), {"escola_ids": [self.escola1.pk]})
        self.assertRedirects(resposta, reverse("mip_inep"), fetch_redirect_response=False)
        self.assertEqual(self._ineps_no_grid(), {"25000002", "25000003"})
        self.escola1.refresh_from_db()
        self.assertFalse(self.escola1.encontrado_relatorio_eace_mip)
        self.assertEqual(self.escola1.status_mip, Escola.AGUARDANDO_VALIDACAO_EACE)
        self.assertEqual(self.escola1.itens_relatorio_eace_mip.count(), 1)
        historico = RiHistorico.objects.get(ri=self.ri1, campo="Grid do MIP")
        self.assertEqual(historico.autor, self.usuario)
        self.assertIn("Retirado da lista", historico.valor_novo)

    def test_retira_varios_selecionados(self):
        self.client.post(reverse("mip_retirar_grid"), {"escola_ids": [self.escola1.pk, self.escola2.pk]})
        self.assertEqual(self._ineps_no_grid(), {"25000003"})

    def test_inep_retirado_nao_entra_no_lote(self):
        self.client.post(reverse("mip_retirar_grid"), {"escola_ids": [self.escola1.pk]})
        lote = criar_lote_mip("GO", "Abadiânia", None, None, self.usuario)
        self.assertNotIn(self.escola1, lote.escolas.all())

    def test_inep_fora_do_grid_nao_e_alterado(self):
        Escola.objects.filter(pk=self.escola3.pk).update(status_mip=Escola.EM_ANDAMENTO)
        self.client.post(reverse("mip_retirar_grid"), {"escola_ids": [self.escola3.pk, "abc"]})
        self.escola3.refresh_from_db()
        self.assertTrue(self.escola3.encontrado_relatorio_eace_mip)
        self.assertFalse(RiHistorico.objects.filter(campo="Grid do MIP").exists())

    def test_get_nao_altera_nada(self):
        self.client.get(reverse("mip_retirar_grid"))
        self.escola1.refresh_from_db()
        self.assertTrue(self.escola1.encontrado_relatorio_eace_mip)

    def test_visualizador_nao_pode_retirar_nem_ve_o_botao(self):
        visualizador = User.objects.create_user(
            username="visualizador-retirar", password="senha-teste-123", perfil=User.PERFIL_VISUALIZADOR,
        )
        self.client.force_login(visualizador)
        self.assertNotContains(self.client.get(reverse("mip_inep")), "form-retirar-grid-mip")
        self.client.post(reverse("mip_retirar_grid"), {"escola_ids": [self.escola1.pk]})
        self.escola1.refresh_from_db()
        self.assertTrue(self.escola1.encontrado_relatorio_eace_mip)

    def test_tela_mostra_selecao_e_icone_por_linha(self):
        html = self.client.get(reverse("mip_inep")).content.decode()
        self.assertIn('id="form-retirar-grid-mip"', html)
        self.assertEqual(html.count('data-retirar-checkbox'), 4)  # 3 linhas + seletor no JS
        self.assertEqual(html.count('title="Retirar da lista"'), 3)
        # Confirmação pelo modal padrão do sistema (`core/_modal_confirmar.html`), não pelo `confirm()`.
        self.assertEqual(html.count('data-confirmar-texto-botao="Retirar da lista"'), 4)
        self.assertNotIn("return confirm(", html)


@override_settings(MIP_CONFERENCIA_PLANILHA=True)
class DashboardFaturamentoMipTests(TestCase):
    """Pedido do usuário (2026-10-07): aba "Faturamento MIP" do Dashboard —
    valor do MIP por Status (MIP) comparado com a planilha BASE CONSOLIDADA MIP."""

    setUp_base = ImportarLotesMipEmMassaTests.setUp
    _planilha = ImportarLotesMipEmMassaTests._planilha
    _rodar = ImportarLotesMipEmMassaTests._rodar

    def setUp(self):
        self.setUp_base()
        self.usuario.is_superuser = True
        self.usuario.save()
        self.client.force_login(self.usuario)

    def _usar_planilha(self, caminho):
        patcher = patch("apps.escolas.services.CAMINHO_BASE_CONSOLIDADA_MIP", Path(caminho))
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_processo_concluido_bate_com_a_planilha(self):
        _criar_escola_elegivel_lote("20200001")
        _criar_escola_elegivel_lote("20200002")
        _criar_escola_elegivel_lote("20200003", municipio="Goiânia")  # fica fora da planilha
        caminho = self._planilha([_linha("20200001"), _linha("20200002", "999.00")])
        self._rodar(caminho, "--aplicar", "--incluir-divergentes")
        self._usar_planilha(caminho)

        resposta = self.client.get(reverse("dashboard_faturamento_mip"))
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta.context["total_planilha"], Decimal("1299.00"))
        self.assertEqual(resposta.context["total_concluido"], Decimal("1299.00"))
        self.assertEqual(resposta.context["diferenca"], Decimal("0.00"))
        self.assertContains(resposta, "bate ✓")
        cards = {card["status"]: card for card in resposta.context["cards"]}
        self.assertEqual(cards[Escola.FATURAMENTO_CONCLUIDO]["quantidade_ineps"], 2)
        self.assertEqual(cards[Escola.AGUARDANDO_VALIDACAO_EACE]["valor"], Decimal("300.00"))
        self.assertEqual(cards[Escola.AGUARDANDO_VALIDACAO_EACE]["valor_planilha"], Decimal("0.00"))

    def test_diferenca_aparece_e_ineps_por_uf(self):
        _criar_escola_elegivel_lote("20200010")
        caminho_importacao = self._planilha([_linha("20200010")])
        self._rodar(caminho_importacao, "--aplicar")
        caminho_dashboard = self.tmp / "planilha_dashboard.xlsx"
        Path(self._planilha([_linha("20200010"), _linha("29999990", "500.00")])).rename(caminho_dashboard)
        self._usar_planilha(caminho_dashboard)

        resposta = self.client.get(
            reverse("dashboard_faturamento_mip"), {"status_mip": Escola.FATURAMENTO_CONCLUIDO, "status_uf": "GO"},
        )
        self.assertEqual(resposta.context["diferenca"], Decimal("-500.00"))
        self.assertContains(resposta, "faltam R$ 500")
        self.assertEqual(resposta.context["fora_do_mip"], 1)
        self.assertEqual([item["inep"] for item in resposta.context["por_estado"][0]["ineps"]], ["20200010"])
        self.assertContains(resposta, "Planilha ✓")

    def test_filtro_por_estado(self):
        _criar_escola_elegivel_lote("20200020")
        _criar_escola_elegivel_lote("20200021", estado="DF", municipio="Brasília")
        caminho = self._planilha([_linha("20200020"), _linha("20200021", cidade="Brasília", uf="DF")])
        self._rodar(caminho, "--aplicar")
        self._usar_planilha(caminho)
        resposta = self.client.get(reverse("dashboard_faturamento_mip"), {"estado": "DF"})
        self.assertEqual(resposta.context["total_planilha"], Decimal("300.00"))
        self.assertEqual([linha["rotulo"] for linha in resposta.context["grafico_municipio"]], ["Brasília"])

    def test_valor_total_do_projeto_usa_valor_de_servico_da_lpu(self):
        KitPadrao.objects.create(
            descricao="Nobreak (serviço, material, equipamento)", lote=9,
            valor_equipamento="900.00", valor_servico="50.00",
        )
        escola, _ = _criar_escola_elegivel_lote("20200030")
        escola.kit_inicial, escola.nobreak_inicial = "2", "Nobreak"
        escola.save()
        Escola.objects.create(inep="20200031", nome="Sem kit", estado="GO", municipio="Abadiânia", lote=9)
        caminho = self._planilha([_linha("20200030")])
        self._rodar(caminho, "--aplicar")
        self._usar_planilha(caminho)

        resposta = self.client.get(reverse("dashboard_faturamento_mip"))
        # Kit (serviço 300) + Nobreak (serviço 50) da 1ª escola + Nobreak
        # padrão da 2ª (sem kit) — valor de serviço, nunca o de equipamento.
        self.assertEqual(resposta.context["total_projeto"], Decimal("400.00"))
        self.assertEqual(resposta.context["total_concluido"], Decimal("300.00"))
        self.assertEqual(resposta.context["falta_projeto"], Decimal("100.00"))
        self.assertContains(resposta, "Falta R$ 100")
        linha_go = resposta.context["grafico_estado"][0]
        self.assertEqual((linha_go["rotulo"], linha_go["meta"], linha_go["valor"]), ("GO", Decimal("400.00"), Decimal("300.00")))

    @override_settings(MIP_CONFERENCIA_PLANILHA=False)
    def test_producao_nao_mostra_nem_le_a_planilha(self):
        _criar_escola_elegivel_lote("20200040")
        caminho = self._planilha([_linha("20200040")])
        self._rodar(caminho, "--aplicar")
        self._usar_planilha(caminho)
        with patch("apps.escolas.services.valores_planilha_base_consolidada_mip") as ler_planilha:
            resposta = self.client.get(
                reverse("dashboard_faturamento_mip"), {"status_mip": Escola.FATURAMENTO_CONCLUIDO, "status_uf": "GO"},
            )
        ler_planilha.assert_not_called()
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta.context["total_concluido"], Decimal("300.00"))
        for texto in ("BASE CONSOLIDADA", "Planilha:", "Planilha ✓", "Fora da planilha", "bate ✓", "não encontrada"):
            self.assertNotContains(resposta, texto)
        self.assertContains(resposta, "20200040")

    def test_sem_planilha_avisa(self):
        self._usar_planilha(self.tmp / "nao-existe.xlsb")
        resposta = self.client.get(reverse("dashboard_faturamento_mip"))
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "não encontrada")
