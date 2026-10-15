"""Importação da planilha FR_XXXX do BDO (pedido do usuário, 2026-10-14).

Cada linha da planilha é um município com os valores do município (Valor
serviços, créditos, descontos, descontos proporcionais e total) e a lista
de INEPs separados por vírgula. Cada Boletim de desempenho é independente
(pedido do usuário, 2026-10-14): nada já importado é sobreposto, e um
Boletim já importado é recusado."""

import datetime
import unicodedata
from decimal import Decimal, InvalidOperation

import openpyxl
from django.db import transaction
from django.utils import timezone

from .models import BoletimBdo, InepBdo, LogStatusDivergenciaBdo, MunicipioBdo

COLUNA_BOLETIM = "BOLETIM DE DESEMPENHO"
COLUNA_MUNICIPIO = "MUNICIPIO"
COLUNA_UF = "UF"
COLUNA_INEPS = "INEPS"
# campo do `MunicipioBdo` -> (coluna normalizada, nome para o usuário),
# na ordem de `MunicipioBdo.CAMPOS_VALORES`
COLUNAS_VALORES = {
    "valor_servicos": ("VALOR SERVICOS", "Valor serviços"),
    "valor_creditos": ("VALOR CREDITOS", "Valor créditos"),
    "valor_descontos": ("VALOR DESCONTOS", "Valor descontos"),
    "valor_descontos_proporcionais": ("VALOR DESCONTOS PROPORCIONAIS", "Valor descontos proporcionais"),
    "valor_total": ("VALOR TOTAL", "Valor total"),
}
NOMES_COLUNAS = {
    COLUNA_BOLETIM: "Boletim de desempenho",
    COLUNA_MUNICIPIO: "Município",
    COLUNA_UF: "UF",
    **{coluna: nome for coluna, nome in COLUNAS_VALORES.values()},
    COLUNA_INEPS: "INEPs",
}
COLUNAS_OBRIGATORIAS = tuple(NOMES_COLUNAS)


class PlanilhaBdoError(Exception):
    pass


def _normalizar(valor):
    texto = " ".join((str(valor) if valor is not None else "").strip().upper().split())
    return "".join(c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c))


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _valor_decimal(valor):
    """Aceita número do Excel ou texto no formato brasileiro ("1.094,45")."""
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor)).quantize(Decimal("0.01"))
    texto = _texto(valor).replace("R$", "").replace(" ", "")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    return Decimal(texto).quantize(Decimal("0.01"))


def ler_planilha_bdo(arquivo):
    """Lê a 1ª aba que tiver as colunas obrigatórias no cabeçalho. Devolve
    `{numero_boletim: [{municipio, uf, ineps, <campos de COLUNAS_VALORES>}, ...]}` ou
    levanta `PlanilhaBdoError` com a mensagem para o usuário."""
    try:
        workbook = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    except Exception:
        raise PlanilhaBdoError("Não foi possível ler o arquivo — verifique se é um .xlsx válido.")

    try:
        for aba in workbook.worksheets:
            linhas = aba.iter_rows(values_only=True)
            cabecalho = [_normalizar(valor) for valor in next(linhas, None) or ()]
            if set(COLUNAS_OBRIGATORIAS) <= set(cabecalho):
                indice = {coluna: cabecalho.index(coluna) for coluna in COLUNAS_OBRIGATORIAS}
                return _ler_linhas(linhas, indice)
    finally:
        workbook.close()

    nomes = ", ".join(NOMES_COLUNAS[coluna] for coluna in COLUNAS_OBRIGATORIAS)
    raise PlanilhaBdoError(f"Nenhuma aba do arquivo tem as colunas obrigatórias: {nomes}.")


def _ler_linhas(linhas, indice):
    boletins = {}
    erros = []
    for numero_linha, linha in enumerate(linhas, start=2):
        def campo(coluna):
            posicao = indice[coluna]
            return linha[posicao] if posicao < len(linha) else None

        if all(_texto(valor) == "" for valor in linha):
            continue
        boletim = _texto(campo(COLUNA_BOLETIM))
        municipio = _texto(campo(COLUNA_MUNICIPIO))
        uf = _texto(campo(COLUNA_UF)).upper()
        ineps = [inep.strip() for inep in _texto(campo(COLUNA_INEPS)).split(",") if inep.strip()]
        valores = {}
        for campo_modelo, (coluna, _nome) in COLUNAS_VALORES.items():
            try:
                valores[campo_modelo] = _valor_decimal(campo(coluna))
            except (InvalidOperation, ValueError):
                valores[campo_modelo] = None

        faltando = [
            nome for nome, ok in (
                ("Boletim de desempenho", boletim),
                ("Município", municipio),
                ("UF", len(uf) == 2),
                *((nome, valores[campo_modelo] is not None) for campo_modelo, (_coluna, nome) in COLUNAS_VALORES.items()),
                ("INEPs", ineps),
            ) if not ok
        ]
        if faltando:
            erros.append(f"linha {numero_linha} ({', '.join(faltando)})")
            continue
        boletins.setdefault(boletim, []).append({"municipio": municipio, "uf": uf, "ineps": ineps, **valores})

    if erros:
        exemplos = "; ".join(erros[:5]) + ("; ..." if len(erros) > 5 else "")
        raise PlanilhaBdoError(f"{len(erros)} linha(s) com dado ausente ou inválido: {exemplos}.")
    if not boletins:
        raise PlanilhaBdoError("A planilha não tem nenhuma linha de dados.")
    return boletins


@transaction.atomic
def importar_planilha_bdo(boletins, tipo, arquivo_nome, usuario, mes_referencia=None):
    """Grava o que `ler_planilha_bdo` leu, cada Boletim independente dos
    que já existem, todos com o `tipo` (RE/RI) escolhido na tela. Levanta `PlanilhaBdoError` sem gravar nada se algum Boletim da
    planilha já foi importado. `mes_referencia` vai no dia 1 do mês. Devolve
    os Boletins gravados."""
    repetidos = sorted(BoletimBdo.objects.filter(numero__in=boletins).values_list("numero", flat=True))
    if repetidos:
        raise PlanilhaBdoError(
            f"Boletim de desempenho já importado: {', '.join(repetidos)}. "
            "Os dados já gravados não são sobrepostos."
        )
    gravados = []
    for numero, linhas in boletins.items():
        boletim = BoletimBdo.objects.create(
            numero=numero, tipo=tipo, arquivo_nome=arquivo_nome, importado_por=usuario,
            mes_referencia=mes_referencia.replace(day=1) if mes_referencia else None,
        )
        ineps = []
        for linha in linhas:
            # `create` um a um: no MySQL o `bulk_create` não devolve o pk,
            # que os INEPs precisam.
            municipio = MunicipioBdo.objects.create(
                boletim=boletim, municipio=linha["municipio"], uf=linha["uf"],
                **{campo_modelo: linha[campo_modelo] for campo_modelo in COLUNAS_VALORES},
            )
            ineps.extend(InepBdo(municipio_bdo=municipio, inep=inep) for inep in linha["ineps"])
        InepBdo.objects.bulk_create(ineps)
        gravados.append(boletim)
    return gravados


COLUNAS_PRO_RATA = {
    "INEP": "INEP",
    "ATIVADO EM": "Ativado em",
    "DIAS DESCONTADOS": "Dias Descontados",
    "VALOR": "Valor",
}


def _data(valor):
    """Data do Excel ou texto "dd/mm/aaaa"; `None` se vazia ou inválida."""
    if isinstance(valor, datetime.datetime):
        return valor.date()
    if isinstance(valor, datetime.date):
        return valor
    try:
        return datetime.datetime.strptime(_texto(valor).split(" ")[0], "%d/%m/%Y").date()
    except ValueError:
        return None


def _inteiro(valor):
    try:
        numero = Decimal(_texto(valor).replace(",", "."))
    except InvalidOperation:
        return None
    return int(numero) if numero == numero.to_integral_value() and numero >= 0 else None


def ler_planilha_pro_rata(arquivo):
    """Pedido do usuário (2026-10-15): planilha Pro rata do cliente de um
    Boletim, com o "Ativado em", os "Dias Descontados" e o "Valor" (o
    desconto) de cada INEP. Lê a 1ª aba com as 4 colunas; devolve
    `{inep: {"ativado_em", "dias_descontados", "valor_desconto"}}` ou levanta
    `PlanilhaBdoError`. INEP repetido com dados diferentes é erro: não dá
    para saber qual vale."""
    try:
        workbook = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    except Exception:
        raise PlanilhaBdoError("Não foi possível ler o arquivo — verifique se é um .xlsx válido.")

    try:
        for aba in workbook.worksheets:
            linhas = aba.iter_rows(values_only=True)
            cabecalho = [_normalizar(valor) for valor in next(linhas, None) or ()]
            if set(COLUNAS_PRO_RATA) <= set(cabecalho):
                return _ler_linhas_pro_rata(linhas, {coluna: cabecalho.index(coluna) for coluna in COLUNAS_PRO_RATA})
    finally:
        workbook.close()
    raise PlanilhaBdoError(f"Nenhuma aba do arquivo tem as colunas obrigatórias: {', '.join(COLUNAS_PRO_RATA.values())}.")


def _ler_linhas_pro_rata(linhas, indice):
    dados = {}
    erros = []
    for numero_linha, linha in enumerate(linhas, start=2):
        def campo(coluna):
            posicao = indice[coluna]
            return linha[posicao] if posicao < len(linha) else None

        if all(_texto(valor) == "" for valor in linha):
            continue
        inep = _texto(campo("INEP"))
        try:
            valor_desconto = _valor_decimal(campo("VALOR"))
        except (InvalidOperation, ValueError):
            valor_desconto = None
        item = {
            "ativado_em": _data(campo("ATIVADO EM")),
            "dias_descontados": _inteiro(campo("DIAS DESCONTADOS")),
            "valor_desconto": valor_desconto,
        }
        faltando = [
            nome for nome, ok in (
                ("INEP", inep), ("Ativado em", item["ativado_em"]),
                ("Dias Descontados", item["dias_descontados"] is not None), ("Valor", valor_desconto is not None),
            ) if not ok
        ]
        if faltando:
            erros.append(f"linha {numero_linha} ({', '.join(faltando)})")
        elif dados.setdefault(inep, item) != item:
            erros.append(f"linha {numero_linha} (INEP {inep} repetido com outros dados)")
    if erros:
        exemplos = "; ".join(erros[:5]) + ("; ..." if len(erros) > 5 else "")
        raise PlanilhaBdoError(f"{len(erros)} linha(s) com dado ausente ou inválido: {exemplos}.")
    if not dados:
        raise PlanilhaBdoError("A planilha não tem nenhuma linha de dados.")
    return dados


@transaction.atomic
def importar_pro_rata(boletim, dados, arquivo_nome, usuario=None):
    """Grava o que `ler_planilha_pro_rata` leu nos INEPs do Boletim.
    Reenviar substitui a anterior: INEP do Boletim fora da planilha fica
    vazio. INEP com status de divergência cujo dado do cliente mudou volta a
    pendente, com log (pedido do usuário, 2026-10-15). Devolve
    `(atualizados, ineps_fora_do_boletim)` - os INEPs da planilha que não
    estão no Boletim são só contados."""
    campos = {
        "ativado_em_cliente": "ativado_em", "dias_descontados_cliente": "dias_descontados",
        "valor_desconto_cliente": "valor_desconto",
    }
    ineps = list(InepBdo.objects.filter(municipio_bdo__boletim=boletim))
    logs = []
    for inep in ineps:
        item = dados.get(inep.inep, {})
        novos = {campo: item.get(chave) for campo, chave in campos.items()}
        if inep.status_divergencia and any(getattr(inep, campo) != valor for campo, valor in novos.items()):
            logs.append(LogStatusDivergenciaBdo(
                inep_bdo=inep, boletim=boletim.numero, inep=inep.inep,
                acao=LogStatusDivergenciaBdo.ACAO_RETIRADO_PRO_RATA, status_anterior=inep.status_divergencia,
                ativado_em_cliente=inep.ativado_em_cliente, valor_desconto_cliente=inep.valor_desconto_cliente,
                usuario=usuario,
            ))
            inep.status_divergencia, inep.status_divergencia_em, inep.status_divergencia_por = "", timezone.now(), usuario
        for campo, valor in novos.items():
            setattr(inep, campo, valor)
    InepBdo.objects.bulk_update(
        ineps, [*campos, "status_divergencia", "status_divergencia_em", "status_divergencia_por"], batch_size=500,
    )
    LogStatusDivergenciaBdo.objects.bulk_create(logs)
    boletim.pro_rata_arquivo = arquivo_nome
    boletim.save(update_fields=["pro_rata_arquivo"])
    atualizados = sum(1 for inep in ineps if inep.ativado_em_cliente)
    return atualizados, sorted(set(dados) - {inep.inep for inep in ineps})


def dias_proporcionais(data_ativacao):
    """Pedido do usuário (2026-10-14): dias do valor proporcional, da Data
    de Ativação até o dia 30 do mês (mês comercial de 30 dias, inclusive
    fevereiro), sem contar o dia da ativação: 15/09 = 15 dias. Ativação no
    dia 31 vai para o mês seguinte, que é cobrado cheio: 30 dias."""
    if data_ativacao is None:
        return None
    return 30 if data_ativacao.day == 31 else 30 - data_ativacao.day


def dias_no_mes_referencia(data_ativacao, mes_referencia):
    """Pedido do usuário (2026-10-14): dias cobrados do INEP no mês de
    referência do Boletim. Ativado antes do mês: mês cheio (30). Ativado no
    mês: `dias_proporcionais` - no dia 31 vai para o mês seguinte (0).
    Ativado depois do mês ou sem data: vazio. Boletim sem mês de
    referência: `dias_proporcionais`, como antes do campo existir."""
    if data_ativacao is None:
        return None
    if mes_referencia is None:
        return dias_proporcionais(data_ativacao)
    mes_ativacao = (data_ativacao.year, data_ativacao.month)
    mes = (mes_referencia.year, mes_referencia.month)
    if mes_ativacao < mes:
        return 30
    if mes_ativacao > mes:
        return None
    return 0 if data_ativacao.day == 31 else dias_proporcionais(data_ativacao)


def valor_proporcional(valor_mensal, dias):
    """Pedido do usuário (2026-10-14): valor mensal do INEP (assinatura RE
    ou manutenção RI) proporcional aos dias, no mês comercial de 30 dias,
    arredondado em centavos por INEP. Vazio sem valor mensal ou sem dias."""
    if valor_mensal is None or dias is None:
        return None
    return (valor_mensal * dias / 30).quantize(Decimal("0.01"))
