import io
from unittest.mock import patch

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from apps.core.models import User

from .models import ExecucaoAutomacaoIxc, LinhaExecucaoIxc


class AutomacoesIxcPermissaoTests(TestCase):
    """Automações IXC > Login (Endereços) / Atendimentos: rotas exigem
    login e perfil Administrador — mesmo critério das demais telas
    administrativas (RN-004), decisão explícita do usuário (2026-09-22)."""

    def setUp(self):
        self.administrador = User.objects.create_user(
            username="admin-ixc-perm", password="senha-teste-123", perfil=User.PERFIL_ADMINISTRADOR,
        )
        self.analista = User.objects.create_user(
            username="analista-ixc-perm", password="senha-teste-123",
        )

    def test_login_enderecos_exige_login(self):
        resp = self.client.get(reverse("ixc_login_enderecos"))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse("login"), resp.url)

    def test_atendimentos_exige_login(self):
        resp = self.client.get(reverse("ixc_atendimentos"))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse("login"), resp.url)

    def test_login_enderecos_bloqueado_para_analista(self):
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("ixc_login_enderecos"))
        self.assertEqual(resp.status_code, 403)

    def test_atendimentos_bloqueado_para_analista(self):
        self.client.force_login(self.analista)
        resp = self.client.get(reverse("ixc_atendimentos"))
        self.assertEqual(resp.status_code, 403)

    def test_administrador_ve_login_enderecos(self):
        self.client.force_login(self.administrador)
        resp = self.client.get(reverse("ixc_login_enderecos"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Login (Endereços)")

    def test_administrador_ve_atendimentos(self):
        self.client.force_login(self.administrador)
        resp = self.client.get(reverse("ixc_atendimentos"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Atendimentos")

    def test_upload_bloqueado_para_analista(self):
        self.client.force_login(self.analista)
        resp = self.client.post(reverse("ixc_upload", kwargs={
            "tipo": ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, "slot": 1,
        }))
        self.assertEqual(resp.status_code, 403)


def _planilha_xlsx(colunas, linhas):
    """Monta um .xlsx em memória (mesmo formato que o usuário sobe na
    tela) — `colunas` é o cabeçalho, `linhas` é uma lista de listas na
    mesma ordem das colunas."""
    workbook = openpyxl.Workbook()
    aba = workbook.active
    aba.append(list(colunas))
    for linha in linhas:
        aba.append(linha)
    saida = io.BytesIO()
    workbook.save(saida)
    saida.seek(0)
    return SimpleUploadedFile(
        "planilha.xlsx", saida.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


COLUNAS_LOGIN_TESTE = (
    "Cliente_ID", "Login_Contrato_ID", "Plano_ID", "Login_Login", "Login_Senha_Cliente",
    "End_CEP", "End_Bairro", "End_Cidade_ID_IXC", "End_Logradouro", "End_Numero",
)
COLUNAS_ATENDIMENTO_TESTE = (
    "Cliente_ID", "Login_ID", "Contrato_ID", "Filial_ID", "Assunto_ID", "Departamento_ID",
    "Tipo_Processo", "Workflow_ID", "Assunto_Descricao", "Descricao", "Endereco",
)


class AutomacoesIxcUploadTests(TestCase):
    """Upload da planilha (FEAT-053) — valida extensão/colunas antes de
    aceitar (mesmo padrão de `apps.escolas.forms`) e, para Atendimentos, a
    regra "tudo ou nada" (README trazido em 2026-09-23)."""

    def setUp(self):
        self.administrador = User.objects.create_user(
            username="admin-ixc-upload", password="senha-teste-123", perfil=User.PERFIL_ADMINISTRADOR,
        )
        self.client.force_login(self.administrador)

    def test_upload_login_enderecos_colunas_ausentes_e_rejeitado(self):
        arquivo = _planilha_xlsx(["Cliente_ID"], [[1]])
        resp = self.client.post(
            reverse("ixc_upload", kwargs={"tipo": ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, "slot": 1}),
            {"arquivo": arquivo}, follow=True,
        )
        self.assertContains(resp, "Colunas obrigatórias ausentes")
        self.assertEqual(ExecucaoAutomacaoIxc.objects.count(), 0)

    def test_upload_login_enderecos_valido_cria_execucao_e_linhas(self):
        arquivo = _planilha_xlsx(COLUNAS_LOGIN_TESTE, [
            [55, 1151, 4, "joao_silva", "senha123", "60346165", "Centro", 1, "Rua A", 10],
            [56, 1152, 4, "maria_souza", "senha456", "60346166", "Centro", 1, "Rua B", 20],
        ])
        resp = self.client.post(
            reverse("ixc_upload", kwargs={"tipo": ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, "slot": 1}),
            {"arquivo": arquivo}, follow=True,
        )
        self.assertContains(resp, "Planilha enviada")
        execucao = ExecucaoAutomacaoIxc.objects.get()
        self.assertEqual(execucao.tipo, ExecucaoAutomacaoIxc.LOGIN_ENDERECOS)
        self.assertEqual(execucao.slot, 1)
        self.assertEqual(execucao.status, ExecucaoAutomacaoIxc.PENDENTE)
        self.assertEqual(execucao.total_linhas, 2)
        self.assertEqual(execucao.linhas.first().dados_entrada["Login_Login"], "joao_silva")

    def test_upload_atendimentos_tudo_ou_nada_rejeita_planilha_inteira(self):
        linhas = [
            [1, 10, 100, 5, 20, 3, "Avulso", None, "Assunto", "Descrição", "Rua X"],
            # 2ª linha inválida: "Avulso" exige Workflow_ID vazio.
            [2, 11, 101, 5, 20, 3, "Avulso", 99, "Assunto 2", "Descrição 2", "Rua Y"],
        ]
        arquivo = _planilha_xlsx(COLUNAS_ATENDIMENTO_TESTE, linhas)
        resp = self.client.post(
            reverse("ixc_upload", kwargs={"tipo": ExecucaoAutomacaoIxc.ATENDIMENTOS, "slot": 1}),
            {"arquivo": arquivo}, follow=True,
        )
        self.assertContains(resp, "nenhuma linha foi enviada ao IXC")
        self.assertEqual(ExecucaoAutomacaoIxc.objects.count(), 0)

    def test_upload_atendimentos_valido_cria_execucao(self):
        linhas = [
            [1, 10, 100, 5, 20, 3, "Avulso", None, "Assunto", "Descrição", "Rua X"],
        ]
        arquivo = _planilha_xlsx(COLUNAS_ATENDIMENTO_TESTE, linhas)
        resp = self.client.post(
            reverse("ixc_upload", kwargs={"tipo": ExecucaoAutomacaoIxc.ATENDIMENTOS, "slot": 2}),
            {"arquivo": arquivo}, follow=True,
        )
        self.assertContains(resp, "Planilha enviada")
        execucao = ExecucaoAutomacaoIxc.objects.get()
        self.assertEqual(execucao.slot, 2)
        self.assertEqual(execucao.total_linhas, 1)

    def test_baixar_modelo_login_enderecos_serve_arquivo_real(self):
        resp = self.client.get(
            reverse("ixc_baixar_modelo", kwargs={"tipo": ExecucaoAutomacaoIxc.LOGIN_ENDERECOS})
        )
        self.assertEqual(resp.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertIn("Preencher_Aqui", workbook.sheetnames)
        self.assertIn("Instrucoes_Ajuda", workbook.sheetnames)

    def test_baixar_modelo_atendimentos_serve_arquivo_real(self):
        resp = self.client.get(
            reverse("ixc_baixar_modelo", kwargs={"tipo": ExecucaoAutomacaoIxc.ATENDIMENTOS})
        )
        self.assertEqual(resp.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertIn("Modelo_OS", workbook.sheetnames)
        self.assertIn("Instrucoes_Ajuda", workbook.sheetnames)
        cabecalho = next(workbook["Modelo_OS"].iter_rows(min_row=1, max_row=1, values_only=True))
        # RN a formalizar: o modelo real marca campo obrigatório com "*"
        # no próprio cabeçalho (normalizado na leitura, ver
        # `services._normalizar_cabecalho`).
        self.assertIn("Assunto_Descricao*", cabecalho)

    def test_upload_atendimentos_aceita_cabecalho_com_asterisco(self):
        """`doc/Modelo Atendimento IXC.xlsx` marca campo obrigatório com um
        "*" (ex.: "Cliente_ID*") — precisa continuar sendo aceito como
        "Cliente_ID" normal (`services._normalizar_cabecalho`)."""
        colunas_com_asterisco = (
            "Cliente_ID*", "Login_ID*", "Contrato_ID*", "Filial_ID*", "Assunto_ID*",
            "Departamento_ID*", "Tipo_Processo*", "Workflow_ID", "Assunto_Descricao*", "Descricao*",
        )
        arquivo = _planilha_xlsx(colunas_com_asterisco, [
            [1, 10, 100, 5, 20, 3, "Avulso", None, "Assunto", "Descrição"],
        ])
        resp = self.client.post(
            reverse("ixc_upload", kwargs={"tipo": ExecucaoAutomacaoIxc.ATENDIMENTOS, "slot": 1}),
            {"arquivo": arquivo}, follow=True,
        )
        self.assertContains(resp, "Planilha enviada")
        execucao = ExecucaoAutomacaoIxc.objects.get()
        self.assertEqual(execucao.linhas.first().dados_entrada["Cliente_ID"], 1)


class AutomacoesIxcProcessamentoTests(TestCase):
    """Start/chunk/Stop (FEAT-053) — chamadas reais ao IXC sempre
    mockadas: nenhum teste deste projeto pode bater na API de produção."""

    def setUp(self):
        self.administrador = User.objects.create_user(
            username="admin-ixc-proc", password="senha-teste-123", perfil=User.PERFIL_ADMINISTRADOR,
        )
        self.client.force_login(self.administrador)
        arquivo = _planilha_xlsx(COLUNAS_LOGIN_TESTE, [
            [55, 1151, 4, "joao_silva", "senha123", "60346165", "Centro", 1, "Rua A", 10],
            [56, 1152, 4, "maria_souza", "senha456", "60346166", "Centro", 1, "Rua B", 20],
        ])
        self.client.post(
            reverse("ixc_upload", kwargs={"tipo": ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, "slot": 1}),
            {"arquivo": arquivo},
        )
        self.execucao = ExecucaoAutomacaoIxc.objects.get()

    def test_iniciar_muda_status_para_processando(self):
        resp = self.client.post(reverse("ixc_iniciar", kwargs={"pk": self.execucao.pk}), follow=True)
        self.execucao.refresh_from_db()
        self.assertEqual(self.execucao.status, ExecucaoAutomacaoIxc.PROCESSANDO)
        self.assertEqual(resp.status_code, 200)

    @patch("apps.ixc.services.executar_cadastro_ixc")
    def test_chunk_processa_linhas_e_conclui(self, mock_executar):
        mock_executar.side_effect = [
            (True, "Criado com sucesso!", "111"),
            (False, "IXC Negou: login duplicado", None),
        ]
        self.execucao.status = ExecucaoAutomacaoIxc.PROCESSANDO
        self.execucao.save(update_fields=["status"])

        resp = self.client.post(
            reverse("ixc_processar_chunk", kwargs={"pk": self.execucao.pk}),
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(resp.status_code, 200)
        self.execucao.refresh_from_db()
        self.assertEqual(self.execucao.status, ExecucaoAutomacaoIxc.CONCLUIDO)
        self.assertEqual(self.execucao.linhas_sucesso, 1)
        self.assertEqual(self.execucao.linhas_erro, 1)
        self.assertEqual(mock_executar.call_count, 2)

        linha_sucesso = self.execucao.linhas.get(status=LinhaExecucaoIxc.SUCESSO)
        self.assertEqual(linha_sucesso.id_ixc, "111")

    @patch("apps.ixc.services.executar_cadastro_ixc")
    def test_parar_cancela_antes_do_proximo_chunk(self, mock_executar):
        self.execucao.status = ExecucaoAutomacaoIxc.PROCESSANDO
        self.execucao.save(update_fields=["status"])

        self.client.post(reverse("ixc_parar", kwargs={"pk": self.execucao.pk}))
        self.client.post(
            reverse("ixc_processar_chunk", kwargs={"pk": self.execucao.pk}),
            HTTP_HX_REQUEST="true",
        )

        self.execucao.refresh_from_db()
        self.assertEqual(self.execucao.status, ExecucaoAutomacaoIxc.CANCELADO)
        mock_executar.assert_not_called()

    def test_baixar_saida_bloqueado_antes_de_concluir(self):
        resp = self.client.get(reverse("ixc_baixar_saida", kwargs={"pk": self.execucao.pk}))
        self.assertEqual(resp.status_code, 403)

    @patch("apps.ixc.services.executar_cadastro_ixc")
    def test_baixar_saida_depois_de_concluido(self, mock_executar):
        mock_executar.return_value = (True, "Criado com sucesso!", "111")
        self.execucao.status = ExecucaoAutomacaoIxc.PROCESSANDO
        self.execucao.save(update_fields=["status"])
        self.client.post(
            reverse("ixc_processar_chunk", kwargs={"pk": self.execucao.pk}),
            HTTP_HX_REQUEST="true",
        )

        resp = self.client.get(reverse("ixc_baixar_saida", kwargs={"pk": self.execucao.pk}))
        self.assertEqual(resp.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(resp.content))
        cabecalho = next(workbook.active.iter_rows(min_row=1, max_row=1, values_only=True))
        self.assertIn("Status", cabecalho)
        self.assertIn("ID no IXC", cabecalho)


class AutomacoesIxcHistoricoTests(TestCase):
    """Histórico de execuções (pedido do usuário, 2026-09-23): quem
    rodou, quando e a planilha de saída de qualquer execução já criada —
    não só a que está aparecendo no grid do slot no momento."""

    def setUp(self):
        self.administrador = User.objects.create_user(
            username="admin-ixc-historico", password="senha-teste-123", perfil=User.PERFIL_ADMINISTRADOR,
        )
        self.client.force_login(self.administrador)

    def test_historico_lista_execucoes_com_usuario_e_data(self):
        execucao = ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, slot=1,
            nome_arquivo_original="planilha_antiga.xlsx",
            status=ExecucaoAutomacaoIxc.CONCLUIDO, criado_por=self.administrador,
        )
        resp = self.client.get(reverse("ixc_login_enderecos"))
        self.assertContains(resp, "planilha_antiga.xlsx")
        self.assertContains(resp, "admin-ixc-historico")
        self.assertContains(resp, execucao.criado_em.strftime("%d/%m/%Y"))

    def test_historico_baixar_so_aparece_quando_concluido(self):
        concluida = ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.ATENDIMENTOS, slot=1, nome_arquivo_original="a.xlsx",
            status=ExecucaoAutomacaoIxc.CONCLUIDO, criado_por=self.administrador,
        )
        pendente = ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.ATENDIMENTOS, slot=2, nome_arquivo_original="b.xlsx",
            status=ExecucaoAutomacaoIxc.PENDENTE, criado_por=self.administrador,
        )
        resp = self.client.get(reverse("ixc_atendimentos"))
        self.assertContains(resp, reverse("ixc_baixar_saida", kwargs={"pk": concluida.pk}))
        self.assertNotContains(resp, reverse("ixc_baixar_saida", kwargs={"pk": pendente.pk}))

    def test_historico_nao_mistura_os_2_tipos(self):
        ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, slot=1, nome_arquivo_original="login.xlsx",
            status=ExecucaoAutomacaoIxc.CONCLUIDO, criado_por=self.administrador,
        )
        ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.ATENDIMENTOS, slot=1, nome_arquivo_original="atendimento.xlsx",
            status=ExecucaoAutomacaoIxc.CONCLUIDO, criado_por=self.administrador,
        )
        resp = self.client.get(reverse("ixc_login_enderecos"))
        self.assertContains(resp, "login.xlsx")
        self.assertNotContains(resp, "atendimento.xlsx")

    def test_historico_paginacao(self):
        for indice in range(11):
            ExecucaoAutomacaoIxc.objects.create(
                tipo=ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, slot=1,
                nome_arquivo_original=f"planilha_{indice}.xlsx",
                status=ExecucaoAutomacaoIxc.CONCLUIDO, criado_por=self.administrador,
            )
        resp_pagina_1 = self.client.get(reverse("ixc_login_enderecos"))
        self.assertContains(resp_pagina_1, "planilha_10.xlsx")  # mais recente primeiro
        self.assertNotContains(resp_pagina_1, "planilha_0.xlsx")

        resp_pagina_2 = self.client.get(reverse("ixc_login_enderecos"), {"historico_page": 2})
        self.assertContains(resp_pagina_2, "planilha_0.xlsx")


class AutomacoesIxcRegrasNegocioTests(TestCase):
    """Regras pedidas pelo usuário em 2026-09-23: progresso trava em 99%
    até `Concluído`; upload trava enquanto o slot está `Processando`;
    grid volta a ficar vazio (slot_atual) depois de `Concluído`/
    `Cancelado`; só quem criou a execução (ou o superadmin) pode parar."""

    def setUp(self):
        self.administrador = User.objects.create_user(
            username="admin-ixc-regras", password="senha-teste-123", perfil=User.PERFIL_ADMINISTRADOR,
        )
        self.outro_administrador = User.objects.create_user(
            username="outro-admin-ixc-regras", password="senha-teste-123", perfil=User.PERFIL_ADMINISTRADOR,
        )
        self.superadmin = User.objects.create_user(
            username="superadmin-ixc-regras", password="senha-teste-123",
            perfil=User.PERFIL_ADMINISTRADOR, is_superuser=True,
        )

    def test_progresso_trava_em_99_ate_concluido(self):
        execucao = ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, slot=1, nome_arquivo_original="a.xlsx",
            status=ExecucaoAutomacaoIxc.PROCESSANDO, criado_por=self.administrador,
        )
        LinhaExecucaoIxc.objects.create(
            execucao=execucao, numero_linha=2, dados_entrada={}, status=LinhaExecucaoIxc.SUCESSO,
        )
        # 1/1 linha processada (100% pela conta simples), mas status ainda
        # é "processando" — trava em 99%, nunca mostra 100% fora de
        # "Concluído".
        self.assertEqual(execucao.progresso_pct, 99)

        execucao.status = ExecucaoAutomacaoIxc.CONCLUIDO
        self.assertEqual(execucao.progresso_pct, 100)

    def test_slot_atual_nao_devolve_execucao_concluida_ou_cancelada(self):
        ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, slot=1, nome_arquivo_original="concluida.xlsx",
            status=ExecucaoAutomacaoIxc.CONCLUIDO, criado_por=self.administrador,
        )
        self.assertIsNone(ExecucaoAutomacaoIxc.slot_atual(ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, 1))

        ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, slot=1, nome_arquivo_original="cancelada.xlsx",
            status=ExecucaoAutomacaoIxc.CANCELADO, criado_por=self.administrador,
        )
        self.assertIsNone(ExecucaoAutomacaoIxc.slot_atual(ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, 1))

    def test_grid_volta_a_ficar_vazio_apos_concluir_ao_recarregar(self):
        self.client.force_login(self.administrador)
        execucao = ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, slot=1, nome_arquivo_original="concluida_unica.xlsx",
            status=ExecucaoAutomacaoIxc.CONCLUIDO, criado_por=self.administrador,
        )
        resp = self.client.get(reverse("ixc_login_enderecos"))
        conteudo = resp.content.decode()
        # Aparece só 1 vez (no Histórico) — o grid (Processamento 1) não
        # mostra mais essa execução depois de concluída.
        self.assertEqual(conteudo.count("concluida_unica.xlsx"), 1)
        self.assertContains(resp, "Nenhum arquivo selecionado")

    def test_upload_bloqueado_enquanto_slot_esta_processando(self):
        self.client.force_login(self.administrador)
        ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, slot=1, nome_arquivo_original="em_andamento.xlsx",
            status=ExecucaoAutomacaoIxc.PROCESSANDO, criado_por=self.administrador,
        )
        arquivo = _planilha_xlsx(COLUNAS_LOGIN_TESTE, [
            [55, 1151, 4, "novo_login", "senha123", "60346165", "Centro", 1, "Rua A", 10],
        ])
        resp = self.client.post(
            reverse("ixc_upload", kwargs={"tipo": ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, "slot": 1}),
            {"arquivo": arquivo}, follow=True,
        )
        self.assertContains(resp, "está em andamento")
        self.assertEqual(ExecucaoAutomacaoIxc.objects.filter(tipo=ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, slot=1).count(), 1)

    def test_parar_bloqueado_para_outro_administrador(self):
        execucao = ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, slot=1, nome_arquivo_original="a.xlsx",
            status=ExecucaoAutomacaoIxc.PROCESSANDO, criado_por=self.administrador,
        )
        self.client.force_login(self.outro_administrador)
        resp = self.client.post(reverse("ixc_parar", kwargs={"pk": execucao.pk}))
        self.assertEqual(resp.status_code, 403)
        execucao.refresh_from_db()
        self.assertFalse(execucao.cancelar_solicitado)

    def test_parar_permitido_para_quem_criou(self):
        execucao = ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, slot=1, nome_arquivo_original="a.xlsx",
            status=ExecucaoAutomacaoIxc.PROCESSANDO, criado_por=self.administrador,
        )
        self.client.force_login(self.administrador)
        resp = self.client.post(reverse("ixc_parar", kwargs={"pk": execucao.pk}), follow=True)
        self.assertEqual(resp.status_code, 200)
        execucao.refresh_from_db()
        self.assertTrue(execucao.cancelar_solicitado)

    def test_parar_permitido_para_superadmin(self):
        execucao = ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, slot=1, nome_arquivo_original="a.xlsx",
            status=ExecucaoAutomacaoIxc.PROCESSANDO, criado_por=self.administrador,
        )
        self.client.force_login(self.superadmin)
        resp = self.client.post(reverse("ixc_parar", kwargs={"pk": execucao.pk}), follow=True)
        self.assertEqual(resp.status_code, 200)
        execucao.refresh_from_db()
        self.assertTrue(execucao.cancelar_solicitado)

    def test_grid_processando_trava_upload_e_nao_oferece_planilha_de_saida(self):
        self.client.force_login(self.administrador)
        ExecucaoAutomacaoIxc.objects.create(
            tipo=ExecucaoAutomacaoIxc.LOGIN_ENDERECOS, slot=1, nome_arquivo_original="a.xlsx",
            status=ExecucaoAutomacaoIxc.PROCESSANDO, criado_por=self.administrador,
        )
        resp = self.client.get(reverse("ixc_login_enderecos"))
        # A linha viva do grid (Processamento 1) não tem mais o link de
        # download — só o Histórico (mais abaixo) oferece "Baixar".
        self.assertContains(resp, "Travado")
        self.assertNotContains(resp, 'id="arquivo-ixc-login_enderecos-1"')
