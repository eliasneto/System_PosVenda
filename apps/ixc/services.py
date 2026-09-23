"""Orquestração Django das Automações IXC (FEAT-053, a formalizar) — lê a
planilha enviada (`openpyxl`, mesmo padrão de `apps/escolas/services.py`),
grava 1 `LinhaExecucaoIxc` por linha e processa em chunks chamando as
funções de negócio de `apps/integracoes/ixc/services/` (client HTTP real
do IXC, RN a formalizar)."""

import datetime
import io
from decimal import Decimal

import openpyxl
from django.conf import settings
from django.db import transaction
from django.utils import timezone

from apps.integracoes.ixc.services.criacao_atendimento_ixc import (
    executar_abertura_atendimento,
    validar_tipo_processo,
)
from apps.integracoes.ixc.services.criacao_login_ixc import executar_cadastro_ixc

from .models import ExecucaoAutomacaoIxc, LinhaExecucaoIxc

# Mesma ordem de `doc/Modelo Login Enderecos IXC.xlsx` (planilha modelo
# real, trazida de sgpspeed em 2026-09-23 — aba "Preencher_Aqui").
COLUNAS_LOGIN_ENDERECOS = (
    "Cliente_ID", "Login_Contrato_ID", "Plano_ID", "Login_Login", "Login_Senha_Cliente",
    "End_CEP", "End_Bairro", "End_Cidade_ID_IXC", "End_Logradouro", "End_Numero",
    "End_Referencia", "Obs_Cliente", "VELOCIDADE", "TIPO DE ACESSO", "BLOCO IP",
    "DUPLA ABORDAGEM", "ENTREGA RB",
)
# Colunas sem as quais a linha não tem como ser processada (o resto é
# enriquecimento opcional, dobrado dentro de "obs" — ver
# `executar_cadastro_ixc`).
COLUNAS_OBRIGATORIAS_LOGIN_ENDERECOS = (
    "Cliente_ID", "Login_Contrato_ID", "Plano_ID", "Login_Login", "Login_Senha_Cliente",
    "End_CEP", "End_Logradouro", "End_Numero", "End_Bairro", "End_Cidade_ID_IXC",
)

COLUNAS_ATENDIMENTOS = (
    "Cliente_ID", "Login_ID", "Contrato_ID", "Filial_ID", "Assunto_ID", "Departamento_ID",
    "Tipo_Processo", "Workflow_ID", "Assunto_Descricao", "Descricao", "Endereco",
)
# Mesmas colunas de `doc/Modelo Atendimento IXC.xlsx` (aba "Modelo_OS") —
# lá cada uma delas vem marcada com "*" (RN a formalizar, normalizado por
# `_normalizar_cabecalho`). "Endereco" não existe nesse modelo real (segue
# opcional, só usado quando a planilha do usuário a incluir).
COLUNAS_OBRIGATORIAS_ATENDIMENTOS = (
    "Cliente_ID", "Login_ID", "Contrato_ID", "Filial_ID", "Assunto_ID", "Departamento_ID",
    "Tipo_Processo", "Assunto_Descricao", "Descricao",
)

COLUNAS_POR_TIPO = {
    ExecucaoAutomacaoIxc.LOGIN_ENDERECOS: COLUNAS_LOGIN_ENDERECOS,
    ExecucaoAutomacaoIxc.ATENDIMENTOS: COLUNAS_ATENDIMENTOS,
}
COLUNAS_OBRIGATORIAS_POR_TIPO = {
    ExecucaoAutomacaoIxc.LOGIN_ENDERECOS: COLUNAS_OBRIGATORIAS_LOGIN_ENDERECOS,
    ExecucaoAutomacaoIxc.ATENDIMENTOS: COLUNAS_OBRIGATORIAS_ATENDIMENTOS,
}

TAMANHO_CHUNK = 5


def _normalizar_cabecalho(valor):
    """`doc/Modelo Atendimento IXC.xlsx` marca campo obrigatório com um "*"
    no próprio cabeçalho (ex.: "Cliente_ID*") — removido aqui para casar
    com o nome de coluna esperado pelas automações (`apps/integracoes/
    ixc/services/`), que não conhecem esse "*"."""
    texto = (str(valor) if valor is not None else "").strip()
    return texto[:-1].strip() if texto.endswith("*") else texto


def colunas_presentes(workbook):
    """Colunas do cabeçalho (1ª linha, 1ª aba) — usado pelos forms para
    validar se as colunas obrigatórias estão todas lá antes de aceitar o
    arquivo."""
    aba = workbook[workbook.sheetnames[0]]
    cabecalho = next(aba.iter_rows(min_row=1, max_row=1, values_only=True), None)
    return {_normalizar_cabecalho(v) for v in (cabecalho or ()) if _normalizar_cabecalho(v)}


def ler_linhas_planilha(workbook):
    """Gera `(numero_linha, dict)` para cada linha de dados da 1ª aba —
    `numero_linha` é a linha real no arquivo (2 = 1ª linha de dados, já que
    a linha 1 é cabeçalho), usada nas mensagens de erro e na planilha de
    saída. Linhas totalmente vazias são ignoradas."""
    aba = workbook[workbook.sheetnames[0]]
    linhas = aba.iter_rows(min_row=1, values_only=True)
    cabecalho = [_normalizar_cabecalho(v) for v in next(linhas)]
    for numero_linha, valores in enumerate(linhas, start=2):
        if all(v in (None, "") for v in valores):
            continue
        dados = {chave: valor for chave, valor in zip(cabecalho, valores) if chave}
        yield numero_linha, dados


def _valor_serializavel(valor):
    """`JSONField` exige tipos serializáveis — data/Decimal do openpyxl
    viram texto/número simples; o resto passa direto."""
    if isinstance(valor, (datetime.datetime, datetime.date)):
        return valor.isoformat()
    if isinstance(valor, Decimal):
        return float(valor)
    return valor


@transaction.atomic
def criar_execucao(tipo, slot, arquivo, usuario):
    """Lê o arquivo inteiro e grava 1 `LinhaExecucaoIxc` por linha de dados
    — chamado só depois de `PlanilhaIxc*UploadForm.clean_arquivo` já ter
    validado extensão/colunas (e, no caso de Atendimentos, a regra "tudo ou
    nada" de `validar_tipo_processo`)."""
    workbook = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    try:
        linhas = list(ler_linhas_planilha(workbook))
    finally:
        workbook.close()

    execucao = ExecucaoAutomacaoIxc.objects.create(
        tipo=tipo, slot=slot, nome_arquivo_original=arquivo.name, criado_por=usuario,
    )
    LinhaExecucaoIxc.objects.bulk_create([
        LinhaExecucaoIxc(
            execucao=execucao,
            numero_linha=numero_linha,
            dados_entrada={chave: _valor_serializavel(valor) for chave, valor in dados.items()},
        )
        for numero_linha, dados in linhas
    ])
    return execucao


def validar_tudo_ou_nada_atendimentos(workbook):
    """RN a formalizar (README trazido em 2026-09-23): TODAS as linhas
    precisam ter `Tipo_Processo`/`Workflow_ID` válidos antes de aceitar o
    arquivo — 1 linha inválida rejeita a planilha inteira. Retorna a lista
    de mensagens de erro (vazia = tudo válido)."""
    erros = []
    for numero_linha, dados in ler_linhas_planilha(workbook):
        valido, erro = validar_tipo_processo(dados)
        if not valido:
            erros.append(f"Linha {numero_linha}: {erro}")
    return erros


def _executar_linha(execucao, linha, usuario):
    if execucao.tipo == ExecucaoAutomacaoIxc.LOGIN_ENDERECOS:
        sucesso, mensagem, id_ixc = executar_cadastro_ixc(linha.dados_entrada)
        return sucesso, mensagem, id_ixc or ""
    sucesso, mensagem = executar_abertura_atendimento(linha.dados_entrada, usuario_sistema=usuario)
    return sucesso, mensagem, ""


def processar_proximo_chunk(execucao, usuario, tamanho=TAMANHO_CHUNK):
    """1 chunk = até `tamanho` linhas chamadas ao IXC nesta mesma
    requisição (sem fila/worker dedicado — decisão registrada com o
    usuário em 2026-09-23). Chamado repetidamente pelo HTMX
    (`hx-trigger=load` encadeado na própria linha do grid) até não sobrar
    linha pendente. Relê `cancelar_solicitado` a cada chamada — Stop pode
    chegar entre 2 chunks."""
    execucao.refresh_from_db()

    if execucao.cancelar_solicitado:
        execucao.status = ExecucaoAutomacaoIxc.CANCELADO
        execucao.concluido_em = timezone.now()
        execucao.save(update_fields=["status", "concluido_em"])
        return execucao

    if execucao.status != ExecucaoAutomacaoIxc.PROCESSANDO:
        return execucao

    pendentes = list(execucao.linhas.filter(status=LinhaExecucaoIxc.PENDENTE)[:tamanho])
    for linha in pendentes:
        try:
            sucesso, mensagem, id_ixc = _executar_linha(execucao, linha, usuario)
        except Exception as exc:
            sucesso, mensagem, id_ixc = False, f"Erro técnico: {exc}", ""
        linha.status = LinhaExecucaoIxc.SUCESSO if sucesso else LinhaExecucaoIxc.ERRO
        linha.mensagem = str(mensagem)[:500]
        linha.id_ixc = str(id_ixc or "")[:50]
        linha.save(update_fields=["status", "mensagem", "id_ixc"])

    if not execucao.linhas.filter(status=LinhaExecucaoIxc.PENDENTE).exists():
        execucao.status = ExecucaoAutomacaoIxc.CONCLUIDO
        execucao.concluido_em = timezone.now()
        execucao.save(update_fields=["status", "concluido_em"])

    return execucao


# Planilhas modelo reais (trazidas de sgpspeed em 2026-09-23, pedido do
# usuário) — cada uma com 1 aba de cabeçalho (sem dados) + 1 aba
# "Instrucoes_Ajuda" explicando cada campo. Servidas como vieram, sem
# gerar nada por código.
CAMINHO_PLANILHA_MODELO_POR_TIPO = {
    ExecucaoAutomacaoIxc.LOGIN_ENDERECOS: settings.BASE_DIR / "doc" / "Modelo Login Enderecos IXC.xlsx",
    ExecucaoAutomacaoIxc.ATENDIMENTOS: settings.BASE_DIR / "doc" / "Modelo Atendimento IXC.xlsx",
}


def gerar_planilha_modelo(tipo):
    return CAMINHO_PLANILHA_MODELO_POR_TIPO[tipo].read_bytes()


def gerar_planilha_saida(execucao):
    """Planilha de saída: mesmas colunas de entrada + Status/Mensagem/ID no
    IXC de cada linha — gerada na hora do download, nada fica salvo em
    disco (mesmo padrão de `apps.ri.services.gerar_planilha_faturamento`)."""
    colunas_entrada = COLUNAS_POR_TIPO[execucao.tipo]
    cabecalho = list(colunas_entrada) + ["Status", "Mensagem", "ID no IXC"]

    workbook = openpyxl.Workbook()
    aba = workbook.active
    aba.title = "Resultado"
    aba.append(cabecalho)
    for linha in execucao.linhas.order_by("numero_linha"):
        aba.append(
            [linha.dados_entrada.get(coluna, "") for coluna in colunas_entrada]
            + [linha.get_status_display(), linha.mensagem, linha.id_ixc]
        )

    saida = io.BytesIO()
    workbook.save(saida)
    return saida.getvalue()
